"""Verify the normally imported cross-batch workbook, its original rows and results.

No business writes. Set MOTOR_SQLITE_PATH to a completed public Excel replay.
"""
import hashlib
import json
import os
import sys
from decimal import Decimal, ROUND_CEILING
from pathlib import Path
from collections import defaultdict

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
import django
django.setup()
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert
from app import joint_schedule_data, joint_material_evidence
from validate_joint_material_http import oracle

BOOK = '44_跨批次共享物料试排_模拟.xlsx'
KEYS = ('MP-261009-001', 'MP-261009-002')
JOBS = ('SP-261009-001-B001', 'SP-261009-001-B002')
MATERIAL = '01.01.0002'


def verify():
    book_path = ROOT/'demo/workbooks'/BOOK
    workbook = load_workbook(book_path, read_only=True, data_only=False)
    schemas = {s['label']: s for s in SCHEMAS.values()}
    expected = {}
    try:
        for sheet in workbook:
            if sheet.title == '导入说明':
                continue
            schema = schemas[sheet.title]
            rows = list(sheet.iter_rows(values_only=True))
            assert list(rows[0]) == [f['label'] for f in schema['fields']]
            for number, row in enumerate(rows[1:], 2):
                values = {f['name']: convert(v, f) for v, f in zip(row, schema['fields'])}
                identity = schema['key'], values['id']
                assert identity not in expected
                expected[identity] = values
                record = Record.objects.select_related('source_row__batch').get(dataset=identity[0], business_key=identity[1])
                source = record.source_row
                assert record.values == values and source.normalized == values
                assert (source.batch.filename, source.sheet, source.row_number) == (BOOK, sheet.title, number)
                assert source.batch.file_hash == hashlib.sha256(book_path.read_bytes()).hexdigest()
    finally:
        workbook.close()
    assert len(expected) == 741
    rows = list(Record.objects.order_by('dataset', 'business_key').values_list('dataset', 'business_key', 'record_hash'))
    old = [r for r in rows if (r[0], r[1]) not in expected]
    def digest(values):
        h = hashlib.sha256()
        for value in values:
            h.update((json.dumps(value, ensure_ascii=False)+'\n').encode())
        return h.hexdigest()
    assert len(old) == 232200 and digest(old) == '0206fed03b9befe314bdd457ae0ee38f97e70f0300ee95d74ba54a68cfd704b3'
    assert len(rows) == 232941
    profiles = []
    for key in KEYS:
        for policy in ('due', 'priority'):
            loaded = joint_schedule_data.load(key, policy)
            result = loaded['result']
            assert result['state'] == 'trial' and not result['issues']
            assert result['summary']['tasks'] == 52 and result['summary']['jobs'] == 2 and result['summary']['qty'] == 12
            winner = JOBS[0 if policy == 'due' else 1]
            assert {j['id'] for j in result['jobs'] if j['state'] == 'scheduled'} == ({winner} if key == KEYS[0] else set(JOBS))
            assert all(j['lateness_minutes'] == (None if j['state'] == 'blocked' else 0) for j in result['jobs'])
            assert all(j['finished'] is None for j in result['jobs'] if j['state'] == 'blocked')
            material, _ = joint_material_evidence.build(result, MATERIAL, 'kg')
            assert material['summary']['required_qty'] == '171.398'
            amount = '85.699' if key == KEYS[0] else '171.398'
            assert material['summary']['reserved_qty'] == material['summary']['total_supply_qty'] == amount
            assert len(material['jobs']) == 2 and len(material['demands']) == 4 and len(material['lots']) == 1
            for b in result['balances']:
                evidence, _ = joint_material_evidence.build(result, b['material_id'], b['unit'])
                oracle(result, evidence)
            # Independently recalculate every imported BOM demand, using exact decimals.
            bom = {b['id']: b for b in loaded['references']['bom']}
            jobs = {j['id']: j for j in result['jobs']}
            for d in result['demands']:
                b = bom[d['bom_id']]
                step = Decimal(str(d['quantum']))
                qty = (Decimal(str(jobs[d['job_id']]['qty']))*Decimal(str(b['qty']))*(1+Decimal(str(b['scrap_allowance'])))/step).to_integral_value(rounding=ROUND_CEILING)*step
                assert qty == Decimal(d['required_qty'])
            # One device and one worker per successful interval, no double-booking.
            for field in ('resource_id', 'employee_id'):
                intervals = defaultdict(list)
                for task in result['tasks']:
                    if task['state'] == 'scheduled':
                        intervals[task[field]].append((task['started'], task['finished']))
                for slots in intervals.values():
                    ordered = sorted(slots)
                    assert all(a[1] <= b[0] for a, b in zip(ordered, ordered[1:]))
            profiles.append(dict(study=key, policy=policy, complete_jobs=result['summary']['complete_jobs'],
                                 scheduled_jobs=[j['id'] for j in result['jobs'] if j['state'] == 'scheduled'],
                                 reserved_kg=amount, tasks=result['summary']['scheduled_tasks']))
    return dict(success=True, workbook=BOOK, source_rows=741, records=len(rows), record_hash_digest=digest(rows),
                prior_business_facts_preserved=len(old), source_rows_exact=True, independent_decimal_checks=72,
                material_ledger_checks=32, no_resource_or_worker_overlap=True, profiles=profiles)


if __name__ == '__main__':
    proof = verify()
    (ROOT/'data/shared_material_validation.json').write_text(json.dumps(proof, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(proof, ensure_ascii=False))
