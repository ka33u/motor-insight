"""Typed XLSX round trip and normal import/replay in a disposable DB only."""
import os,sys,json,sqlite3,tempfile,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
workspace=Path(tempfile.mkdtemp(prefix='motor-wip-import-'))
with sqlite3.connect(ROOT/'data/platform.sqlite3') as source,sqlite3.connect(workspace/'platform.sqlite3') as target:source.backup(target)
os.environ['MOTOR_SQLITE_PATH']=str(workspace/'platform.sqlite3');os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from django.test import override_settings
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert,stage_file,commit_batch
from app import wip_flow as eng
file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/29_在制流转_模拟.xlsx'
expected=json.loads((ROOT/'data/wip_scenario.json').read_text());build=json.loads((ROOT/'data/wip_scenario_build.json').read_text());parsed={};cells=dates=0
book=load_workbook(file,read_only=True,data_only=False)
for ds in eng.TABLES:
    schema=SCHEMAS[ds];rows=book[schema['label']].iter_rows(values_only=True)
    assert list(next(rows))==[f['label'] for f in schema['fields']];parsed[ds]=[]
    for row in rows:
        assert len(row)==len(schema['fields'])
        assert not any(isinstance(v,str) and v.startswith('=') for v in row)
        parsed[ds].append({f['name']:convert(v,f) for f,v in zip(schema['fields'],row)})
        dates+=sum(f['type'] in ['date','datetime'] and v is not None and not isinstance(v,str) for f,v in zip(schema['fields'],row))
    assert parsed[ds]==expected['tables'][ds],ds
    cells+=len(parsed[ds])*len(schema['fields'])
book.close();before=Record.objects.count();assert not Record.objects.filter(dataset__in=eng.TABLES).exists()
old_rows=list(Record.objects.order_by('id').values());count=sum(map(len,parsed.values()))
with override_settings(BASE_DIR=workspace):
    batch,repeated=stage_file(file)
    assert not repeated and batch.summary['valid']==count and not batch.summary.get('invalid') and not batch.summary.get('conflict'),batch.summary
    commit_batch(batch.pk);assert Record.objects.count()==before+count
    again,repeated=stage_file(file);assert repeated;commit_batch(again.pk);assert Record.objects.count()==before+count
    for ds,rows in parsed.items():assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values',flat=True))==sorted(rows,key=lambda r:r['id'])
    assert list(Record.objects.exclude(dataset__in=eng.TABLES).order_by('id').values())==old_rows
    w=eng.Wip();h=eng.Wip(cutoff='2026-09-20T18:00:00')
    assert eng.summary(w.rows)==build['summary'];assert eng.summary(h.rows)==build['historical_summary']
    assert not w.global_issues and not h.global_issues
    for key,case in build['cases'].items():assert w.index[key]['state']==case['state'],key
report={'workbook_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':count,'cells':cells,'native_date_cells':dates,
        'tables':{k:len(v) for k,v in parsed.items()},'isolated_import':True,'idempotent':True,
        'all_old_records_unchanged':True,'records_before':before,'records_after':before+count,'summary':build['summary'],'historical_summary':build['historical_summary']}
(ROOT/'data/wip_import_rehearsal.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
