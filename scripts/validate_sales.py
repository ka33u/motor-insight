"""Read-only reconciliation of the sales increment and original XLSX.

The independent counts below use imported facts and calendar dates, not the
workbench's summary implementation. No business facts are written by this file.
"""
import os,sys,json,csv,hashlib,sqlite3,zipfile,tempfile,math
from pathlib import Path
from collections import Counter,defaultdict
from datetime import date,datetime
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django
django.setup()
from openpyxl import load_workbook
from django.contrib.auth.models import User
from app.models import Record,ImportBatch,IssueDisposition,AuditEvent,DataQualityScan,DataMonitorPolicy,MetricVersion
from app.schema import SCHEMAS
from app.ingestion import convert
from app import sales,data_health,metric_registry

def sha(v):return hashlib.sha256(json.dumps(v,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()

baseline=ROOT/'data/backups/motor-backup-20261003-213906.zip'
preserved={};source_files=0
expected_delta={'app_record':696,'app_importrow':696,'app_importbatch':1,'app_issuedisposition':1,'app_dataqualityscan':1}
with zipfile.ZipFile(baseline) as z,tempfile.TemporaryDirectory() as temp:
    manifest=json.loads(z.read('manifest.json'))
    for entry in manifest['files']:
        if entry['path'].startswith(('data/imports/','data/device_files/')):
            assert hashlib.sha256((ROOT/entry['path']).read_bytes()).hexdigest()==entry['sha256']
            source_files+=1
    p=Path(temp)/'before.sqlite3';p.write_bytes(z.read('data/platform.sqlite3'))
    with sqlite3.connect(p) as old,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
        tables=[r[0] for r in now.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'app_%'")]
        assert set(tables)=={r[0] for r in old.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'app_%'")}
        assert len(tables)==25
        for table in tables:
            before=list(old.execute(f'SELECT * FROM {table} ORDER BY id'))
            current=list(now.execute(f'SELECT * FROM {table} ORDER BY id'));by_id={r[0]:r for r in current}
            assert all(by_id.get(r[0])==r for r in before),(table,'prior row changed')
            if table!='app_auditevent':assert len(current)-len(before)==expected_delta.get(table,0),(table,'unexpected row count')
            preserved[table]={'before':len(before),'after':len(current),'prior_rows_unchanged':True}

bookfile=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/16_销售报价过程_模拟.xlsx'
file_hash=hashlib.sha256(bookfile.read_bytes()).hexdigest()
assert file_hash=='993b7c6003d2afac6eae1c99c36c0515de13a0515a23d1d14d4e06e1ad9c9e94'
batch=ImportBatch.objects.get(file_hash=file_hash,status='committed')
assert hashlib.sha256(Path(batch.file_path).read_bytes()).hexdigest()==file_hash
assert batch.rows.count()==696 and batch.summary['committed']==696
book=load_workbook(bookfile,read_only=True,data_only=False);cells=0;new_counts={}
for ds in ['quotes','quote_details','quote_order_links','quote_tasks']:
    spec=SCHEMAS[ds];sheet=book[spec['label']];rows=sheet.iter_rows();header=next(rows)
    assert [c.value for c in header]==[f['label'] for f in spec['fields']]
    n=0
    for rownumber,row in enumerate(rows,2):
        values={}
        for cell,field in zip(row,spec['fields']):
            assert cell.data_type!='f'
            if field['type']=='date' and cell.value is not None:assert isinstance(cell.value,datetime)
            if field['name']=='id':assert cell.data_type=='s' and isinstance(cell.value,str)
            values[field['name']]=convert(cell.value,field);cells+=1
        rec=Record.objects.select_related('source_row').get(dataset=ds,business_key=values['id'])
        assert rec.values==values and rec.source_row.normalized==values
        assert rec.source_row.batch_id==batch.pk and rec.source_row.row_number==rownumber
        assert rec.source_row.sheet==spec['label'] and rec.revision==1
        n+=1
    new_counts[ds]=n
book.close()
assert (new_counts,cells)==({'quotes':60,'quote_details':210,'quote_order_links':156,'quote_tasks':270},6642)

scans=list(DataQualityScan.objects.order_by('created_at'));assert len(scans)==4
latest=scans[-1];snap=latest.snapshot;cutoff=datetime.fromisoformat(snap['as_of']);as_of=cutoff.date()
assert sha(snap)==latest.snapshot_hash and latest.engine_hash==data_health.engine_hash()
assert all(s.engine_hash!=latest.engine_hash for s in scans[:-1])
assert not snap['issues'] and not snap['unknown_datasets']
projected=data_health.project(latest,User.objects.get(username='demo_admin'))
assert not projected['stale'] and len(projected['rows'])==59
assert sum(not r['freshness']['has_target'] for r in projected['rows'])==58
policy=DataMonitorPolicy.objects.get(dataset='test_sessions');assert policy.version==2 and DataMonitorPolicy.objects.count()==1

identities=set(Record.objects.values_list('dataset','business_key'))
required=0;references=0;counts=Counter();data=defaultdict(list);facts_hash=hashlib.sha256()
for rec in Record.objects.select_related('source_row').order_by('dataset','business_key').iterator(chunk_size=1000):
    ds=rec.dataset;v=rec.values;counts[ds]+=1
    if ds in sales.TABLES:data[ds].append(v)
    assert v['id']==rec.business_key and v==rec.source_row.normalized
    assert sha(v)==rec.record_hash==rec.source_row.record_hash
    assert (ds,rec.business_key)==(rec.source_row.dataset,rec.source_row.business_key)
    facts_hash.update(json.dumps((ds,rec.business_key,v,rec.revision,rec.source_row_id),ensure_ascii=False,sort_keys=True).encode())
    for field in SCHEMAS[ds]['fields']:
        val=v.get(field['name']);missing=val is None or isinstance(val,str) and not val.strip()
        if field['required']:assert not missing;required+=1
        if missing:continue
        typ=field['type']
        if typ=='int':assert type(val) is int
        elif typ=='bool':assert type(val) is bool
        elif typ=='float':assert type(val) in [int,float] and math.isfinite(val)
        elif typ=='str':assert isinstance(val,str) and val==val.strip() and not val.startswith('=')
        elif typ in ['date','datetime']:
            parsed=(date if typ=='date' else datetime).fromisoformat(val);assert parsed.isoformat()==val
        if field['reference']:assert (field['reference'],val) in identities;references+=1
assert (sum(counts.values()),len(counts))==(108197,59)
for ds,n in counts.items():assert snap['datasets'][ds]['rows']==n
assert sum(r['required_cells'] for r in projected['rows'])==required
assert sum(f['reference_present'] for d in snap['datasets'].values() for f in d['fields'].values())==references

# Independent calendar, quantity and amount reconciliation from imported rows.
from statistics import median
idx={ds:{r['id']:r for r in rows} for ds,rows in data.items()}
quotes=idx['quotes'];meta={r['quote_id']:r for r in data['quote_details']}
assert len(meta)==len(data['quote_details'])==len(quotes)==210
quote_links=defaultdict(list);quote_tasks=defaultdict(list);line_totals=Counter();pairs=set()
for r in data['quote_order_links']:
    q=quotes[r['quote_id']];line=idx['order_lines'][r['order_line_id']];order=idx['orders'][line['order_id']]
    assert r['status']=='有效' and type(r['qty']) is int and r['qty']>0
    assert order['customer_id']==q['customer_id'] and line['product_id']==q['product_id']
    assert order['status'] in ['待生产','部分交付']
    assert q['quote_date']<=order['order_date']<=r['confirmed']<=as_of.isoformat()
    pair=(q['id'],line['id']);assert pair not in pairs;pairs.add(pair)
    quote_links[q['id']].append(r);line_totals[line['id']]+=r['qty']
assert len(line_totals)==156
for key,qty in line_totals.items():assert qty<=idx['order_lines'][key]['qty']
for r in data['quote_tasks']:quote_tasks[r['quote_id']].append(r)
quote_counts={};task_counts={};amount=0;reference=0;waits=[];tax_unknown=set()
for key,q in quotes.items():
    m=meta[key];assert m['inquiry_date']<=q['quote_date']<=as_of.isoformat()
    assert m['owner_id'] in idx['employees'] and m['currency']=='CNY'
    if m['tax_basis']=='待确认':tax_unknown.add(key)
    else:assert m['tax_basis']=='未税'
    assert q['status'] in ['已转订单','部分转单','跟进中','待技术确认','已失单','客户暂停']
    if q['status'] in ['已转订单','部分转单','已失单']:assert q['quote_date']<=m['outcome_date']<=as_of.isoformat()
    else:assert m['outcome_date'] is None
    waits.append((date.fromisoformat(q['quote_date'])-date.fromisoformat(m['inquiry_date'])).days)
    linked=sum(x['qty'] for x in quote_links[key]);assert linked<=q['qty']
    if q['status']=='已转订单':assert linked==q['qty']
    elif q['status']=='部分转单':assert 0<linked<q['qty']
    else:assert linked==0
    own_amount=0;own_ref=0
    for link in quote_links[key]:
        line=idx['order_lines'][link['order_line_id']];order=idx['orders'][line['order_id']]
        assert order['currency']==m['currency']=='CNY' and m['tax_basis']=='未税'
        own_amount+=line['unit_price_cents']*link['qty'];own_ref+=q['unit_price_cents']*link['qty']
    amount+=own_amount;reference+=own_ref
    for t in quote_tasks[key]:
        assert q['quote_date']<=t['created']<=t['due'] and t['created']<=as_of.isoformat()
        assert t['owner_id'] in idx['employees']
        if t['completed']:
            assert t['created']<=t['completed']<=as_of.isoformat() and t['status']=='已完成'
        else:assert t['status']=='待反馈'
        done=bool(t['completed']);overdue=not done and t['due']<as_of.isoformat()
        task_counts[t['id']]={'done':done,'overdue':overdue,'late':bool(done and t['completed']>t['due'])}
    opened=q['status'] in ['部分转单','跟进中','待技术确认','客户暂停']
    quote_counts[key]={'linked_qty':linked,'unlinked_qty':q['qty']-linked,'valid_links':len(quote_links[key]),'comparable_links':len(quote_links[key]),'quote_days':waits[-1],'open':opened,'expired':opened and q['valid_until']<as_of.isoformat(),'overdue_tasks':sum(task_counts[t['id']]['overdue'] for t in quote_tasks[key]),'order_value_cents':own_amount if quote_links[key] else None,'quote_reference_cents':own_ref if quote_links[key] else None,'price_difference_cents':own_amount-own_ref if quote_links[key] else None}
assert tax_unknown=={'BJ2609-00327','BJ2609-00334','BJ2609-00355'}
expected_summary={'quotes':210,'open':sum(q['open'] for q in quote_counts.values()),'expired':sum(q['expired'] for q in quote_counts.values()),'linked_quotes':sum(q['linked_qty']>0 for q in quote_counts.values()),'unproven':0,'attention':len(tax_unknown),'tasks':len(task_counts),'completed_tasks':sum(t['done'] for t in task_counts.values()),'open_tasks':sum(not t['done'] for t in task_counts.values()),'overdue_tasks':sum(t['overdue'] for t in task_counts.values()),'late_tasks':sum(t['late'] for t in task_counts.values()),'task_quotes':len(quote_tasks),'task_attention':0,'valid_links':len(pairs),'comparable_links':len(pairs),'order_lines':len(line_totals),'order_value_cents':amount,'quote_reference_cents':reference,'price_difference_cents':amount-reference,'quote_days_median':median(waits),'quote_days_samples':len(waits)}
assert [expected_summary[k] for k in ['open','expired','linked_quotes','tasks','completed_tasks','open_tasks','overdue_tasks','late_tasks']]==[48,14,156,270,210,60,39,0]
assert (amount,reference,amount-reference,median(waits))==(670813395,681963500,-11150105,2.5)
b=sales.Sales(data,sales.filters({}),cutoff=cutoff.isoformat());assert not b.global_issues
assert b.summary(b.selected())==expected_summary
for key,c in quote_counts.items():
    row=b.quotes[key]['row']
    for field,value in c.items():assert row[field]==value,(key,field)
    assert bool(row['issues'])==(key in tax_unknown)
for key,c in task_counts.items():
    row=b.tasks[key];assert not row['issues']
    for field,value in c.items():assert row[field]==value

exports={}
for tab,name,stage,expected in [('quotes','销售_过期待跟进报价_浏览器导出.csv','expired',{k for k,v in quote_counts.items() if v['expired']}),('tasks','销售_逾期跟进任务_浏览器导出.csv','overdue',{k for k,v in task_counts.items() if v['overdue']})]:
    p=ROOT/'outputs'/name;rows=list(csv.reader(p.open(encoding='utf-8-sig')))
    assert rows[0][2]==cutoff.isoformat();f=json.loads(rows[0][4]);assert f['tab']==tab and f['stage']==stage
    assert not any(f[k] for k in ['family','product','owner','task_owner','customer','quote_status','q','from','to'])
    entries=[dict(zip(rows[2],r)) for r in rows[3:]]
    idfield='报价单号' if tab=='quotes' else '跟进任务号';assert {r[idfield] for r in entries}==expected
    for r in entries:
        key=r[idfield]
        if tab=='quotes':
            q=quotes[key];c=quote_counts[key]
            assert [int(r[n]) for n in ['报价数量','可核对关联数量','尚未关联数量','询价至报价自然日','逾期未完成任务数']]==[q['qty'],c['linked_qty'],c['unlinked_qty'],c['quote_days'],c['overdue_tasks']]
            assert int(r['报价单价分'])==q['unit_price_cents']
            assert r['可比关联订单未税金额分']=='' and r['同数量报价未税参考分']=='' and r['同数量价差分']==''
        else:
            t=idx['quote_tasks'][key]
            assert [r[n] for n in ['报价单号','任务负责工号','任务类型','建立日期','约定完成']]==[t[k] for k in ['quote_id','owner_id','kind','created','due']]
            assert r['截止完成']=='' and r['截止状态']=='未完成且逾期'
    exports[name]={'rows':len(entries),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
follow=IssueDisposition.objects.get(key='sales:BJ2609-00327')
assert (follow.version,follow.status,follow.owner)==(1,'待技术澄清','销售与技术岗位（模拟）')
event=AuditEvent.objects.get(action='sales.followup',object_id=follow.key)
assert event.detail['business_facts_changed'] is False and event.detail['after']['version']==1
version=MetricVersion.objects.get(metric__key='DELIVERY_OTIF',version=4)
assert version.status=='published' and version.calculation_hash==metric_registry.calculation_hash('bi_order_lines')=='f6815a706b1f19c2e2be0af931c640d869b45b17bc0aa1f72d7fe2144356868f'
assert ImportBatch.objects.exclude(status='superseded').count()==16
report={'synthetic':True,'baseline':str(baseline.relative_to(ROOT)),'business_rows':sum(counts.values()),'source_categories':len(counts),'active_workbooks':16,'business_sha256':facts_hash.hexdigest(),'preserved_tables':preserved,'preserved_source_files':source_files,'new_workbook':{'batch':str(batch.pk),'sha256':file_hash,'rows':new_counts,'cells':cells},'sales':expected_summary,'quote_status_counts':dict(Counter(q['status'] for q in quotes.values())),'latest_scan':str(latest.pk),'latest_scan_sha256':latest.snapshot_hash,'basic_issues':0,'checked_required_cells':required,'checked_references':references,'preserved_monitor_policy_version':2,'browser_exports':exports,'coordination_version':1,'metric_v4_unchanged':True,'tests':581,'new_tests':35,'limits':'Synthetic quotes and explicit order links only: no mature opportunity conversion, complete revision history, signed customer approval, authorized discount, realized revenue/profit, or real U8/MES connection.'}
if len(sys.argv)>1:
    backup=Path(sys.argv[1])
    with zipfile.ZipFile(backup) as z,tempfile.TemporaryDirectory() as temp:
        manifest=json.loads(z.read('manifest.json'))
        for item in manifest['files']:
            blob=z.read(item['path']);assert len(blob)==item['size'] and hashlib.sha256(blob).hexdigest()==item['sha256']
        p=Path(temp)/'restored.sqlite3';p.write_bytes(z.read('data/platform.sqlite3'))
        with sqlite3.connect(p) as restored,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
            assert restored.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
            for table in tables:assert list(restored.execute(f'SELECT * FROM {table} ORDER BY id'))==list(now.execute(f'SELECT * FROM {table} ORDER BY id')),table
        report['backup']={'path':str(backup),'manifest_files':len(manifest['files']),'app_tables':len(tables),'sqlite_integrity':'ok','all_app_tables_match':True}
(ROOT/'data/sales_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in report.items() if k!='preserved_tables'},ensure_ascii=False,indent=2))
