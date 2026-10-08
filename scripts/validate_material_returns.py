"""Check normally imported Excel inputs and independently reconcile issue/return quantities."""
import hashlib,json,os,sys
from datetime import date,datetime
from decimal import Decimal,ROUND_CEILING
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert
from app import analytics,material_planning

BOOK='46_生产领退料核对_模拟.xlsx'
WORKS=[f'MO-260928-R{n:03d}' for n in range(1,6)]

def verify():
    path=ROOT/'demo/workbooks'/BOOK;workbook=load_workbook(path,read_only=True,data_only=False);schemas={s['label']:s for s in SCHEMAS.values()};expected={}
    try:
        assert len(workbook.sheetnames)==9
        for sheet in workbook:
            if sheet.title=='导入说明':continue
            schema=schemas[sheet.title];rows=list(sheet.iter_rows(values_only=True));assert list(rows[0])==[f['label'] for f in schema['fields']]
            for number,row in enumerate(rows[1:],2):
                for value,field in zip(row,schema['fields']):
                    if value is None:continue
                    if field['type'] in ('int','float'):assert type(value) in (int,float)
                    elif field['type'] in ('date','datetime'):assert isinstance(value,(date,datetime))
                    elif field['type']=='bool':assert type(value) is bool
                    else:assert isinstance(value,str)
                values={f['name']:convert(v,f) for v,f in zip(row,schema['fields'])};identity=schema['key'],values['id'];assert identity not in expected;expected[identity]=values
                record=Record.objects.select_related('source_row__batch').get(dataset=identity[0],business_key=identity[1]);source=record.source_row
                assert record.values==values==source.normalized
                assert (source.batch.filename,source.sheet,source.row_number)==(BOOK,sheet.title,number)
                assert source.batch.file_hash==hashlib.sha256(path.read_bytes()).hexdigest()
    finally:workbook.close()
    assert len(expected)==114
    rows=list(Record.objects.order_by('dataset','business_key').values_list('dataset','business_key','record_hash'));old=[r for r in rows if (r[0],r[1]) not in expected];digest=hashlib.sha256()
    for row in old:digest.update((json.dumps(row,ensure_ascii=False)+'\n').encode())
    assert len(rows)==233902 and len(old)==233788 and digest.hexdigest()=='01565772db588d470cb47485e0a033013a14751bf20f396f220e363b23868b5d'
    source=analytics._tables(analytics.revision());d=material_planning.MaterialPlanning(source,material_planning.filters({'q':'MO-260928-R','stock_policy':'all_usable'}))
    profiles=[];checks=0
    for n,work in enumerate(WORKS,1):
        row=d.orders[work];assert row['state']==('attention' if n==5 else 'covered')
        for item in row['items']:
            bom=[r for r in source['bom'] if r['product_id']==row['product_id'] and r['version']==row['bom_version'] and r['material_id']==item['material_id']]
            gross=sum((Decimal(str(r['qty']))*row['planned_qty']*(1+Decimal(str(r['scrap_allowance']))) for r in bom),Decimal(0))
            if item['unit']=='件':gross=gross.to_integral_value(rounding=ROUND_CEILING)
            assert item['gross_required']==gross
            events=[r for r in source['inventory_movements'] if r.get('work_order_id')==work and r['material_id']==item['material_id'] and r['occurred']<=analytics.AS_OF]
            issued=-sum((Decimal(str(r['qty_signed'])) for r in events if r['movement']=='生产领料'),Decimal(0))
            returned=sum((Decimal(str(r['qty_signed'])) for r in events if r['movement']=='生产退料'),Decimal(0))
            if n==5 and item['material_id']=='01.01.9002':assert item['issued_qty'] is None and item['remaining_required'] is None
            else:
                assert item['gross_issued_qty']==issued and item['returned_qty']==returned
                assert item['issued_qty']==issued-returned and item['remaining_required']==max(Decimal(0),gross-issued+returned)
            checks+=1
        item=next(x for x in row['items'] if x['material_id']==('01.01.9002' if n==5 else '01.01.9001'))
        profiles.append(dict(id=work,state=row['state'],net_issued=material_planning.clean(item['issued_qty']),remaining=material_planning.clean(item['remaining_required'])))
    assert [p['net_issued'] for p in profiles]==[6.5,0,25,40,None]
    stock=d.stock.stock_index['01.01.9001'];assert (stock['balance_qty'],stock['usable_state_qty'])==(928.5,923.5)
    qc=next(r for r in d.stock.lots if r['material_id']=='01.01.9001' and r['location']=='DJ-RT-01');assert qc['balance_qty']==5 and qc['usable_state_qty']==0 and not qc['issues']
    assert d.stock.stock_index['01.01.9002']['usable_state_qty'] is None
    proof=dict(success=True,source_rows=114,records=233902,protected_old_facts=233788,source_rows_exact=True,independent_material_checks=checks,profiles=profiles,
               raw_material_balance=928.5,usable_material_balance=923.5,pending_return_qty=5,zero_net_is_not_no_history=True,browser_acceptance=False)
    (ROOT/'data/material_returns_validation.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n');return proof

if __name__=='__main__':print(json.dumps(verify(),ensure_ascii=False))
