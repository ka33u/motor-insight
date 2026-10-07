"""Validate typed XLSX cells and commit twice in an isolated database copy."""
import os,sys,json,sqlite3,tempfile,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));workspace=Path(tempfile.mkdtemp(prefix='motor-commitments-import-'))
with sqlite3.connect(ROOT/'data/platform.sqlite3') as source,sqlite3.connect(workspace/'platform.sqlite3') as target:source.backup(target)
os.environ['MOTOR_SQLITE_PATH']=str(workspace/'platform.sqlite3');os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from django.test import override_settings
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert,stage_file,commit_batch
from app import purchase_commitments as eng,supply,analytics
file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/25_采购承诺_模拟.xlsx'
expected=json.loads((ROOT/'data/purchase_commitments_scenario.json').read_text());parsed={};cells=0;dates=0
book=load_workbook(file,read_only=True,data_only=False)
for ds in expected['tables']:
 s=SCHEMAS[ds];rows=book[s['label']].iter_rows(values_only=True);assert list(next(rows))==[f['label'] for f in s['fields']];parsed[ds]=[]
 for row in rows:
  parsed[ds].append({f['name']:convert(v,f) for f,v in zip(s['fields'],row)})
  dates+=sum(f['type'] in ['date','datetime'] and v is not None and not isinstance(v,str) for f,v in zip(s['fields'],row))
 assert parsed[ds]==expected['tables'][ds],ds;cells+=len(parsed[ds])*len(s['fields'])
book.close();before=Record.objects.count();added=sum(map(len,parsed.values()));old=supply.SupplyData(analytics.tables());baseline={'purchase':[supply.clean(r) for r in old.po_rows],'stock':[supply.clean(r) for r in old.lots]}
with override_settings(BASE_DIR=workspace):
 batch,repeated=stage_file(file);assert not repeated;assert batch.summary.get('valid')==added and not batch.summary.get('invalid') and not batch.summary.get('conflict'),batch.summary
 commit_batch(batch.pk);assert Record.objects.count()==before+added
 repeat,is_repeat=stage_file(file);assert is_repeat;commit_batch(repeat.pk);assert Record.objects.count()==before+added
 for ds,rows in parsed.items():assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values',flat=True))==sorted(rows,key=lambda r:r['id'])
 d=eng.Commitments();summary=eng.summary(d.rows)
 assert summary['objects']==144 and summary['valid']==135 and summary['attention']==9 and summary['changed']==52 and summary['late_change']==25
 assert (summary['original_on_time'],summary['original_due'])==(114,144)
 assert (summary['segment_on_time'],summary['segment_due'])==(155,177)
 assert {r['id'] for r in d.rows if r['issues']}==set(expected['cases']),expected['cases']
 assert not d.global_issues
 after=supply.SupplyData(analytics.tables());assert baseline=={'purchase':[supply.clean(r) for r in after.po_rows],'stock':[supply.clean(r) for r in after.lots]}
report={'workbook_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':added,'cells':cells,'native_date_cells':dates,'tables':{k:len(v) for k,v in parsed.items()},'isolated_import':True,'idempotent':True,'original_purchase_and_stock_unchanged':True,'records_before':before,'records_after':before+added,'summary':summary,'cases':expected['cases']}
(ROOT/'data/purchase_commitments_import_rehearsal.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
