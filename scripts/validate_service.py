"""Independent XLSX, service chronology, browser exports and preservation checks."""
import os,sys,json,csv,hashlib,sqlite3,zipfile,tempfile
from pathlib import Path
from datetime import datetime
from collections import Counter
from statistics import mean
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from app import service_board as s,service_views,metric_registry
from app.models import Record,IssueDisposition,MetricVersion
from app.schema import SCHEMAS
from app.ingestion import convert
from app.trace_cases import capture_sources
baseline=ROOT/'data/backups/motor-backup-20261003-193942.zip';preserved={}
with zipfile.ZipFile(baseline) as z,tempfile.TemporaryDirectory() as temp:
    manifest=json.loads(z.read('manifest.json'))
    for item in manifest['files']:
        if item['path'].startswith(('data/imports/','data/device_files/')):assert hashlib.sha256((ROOT/item['path']).read_bytes()).hexdigest()==item['sha256']
    p=Path(temp)/'old.sqlite3';p.write_bytes(z.read('data/platform.sqlite3'))
    with sqlite3.connect(p) as old,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
        tables=[r[0] for r in old.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'app_%'")]
        changes={'app_record':449,'app_importrow':449,'app_importbatch':1,'app_issuedisposition':1}
        for t in tables:
            previous=list(old.execute(f'SELECT * FROM {t} ORDER BY id'));current=list(now.execute(f'SELECT * FROM {t} ORDER BY id'));current_by_id={r[0]:r for r in current}
            assert all(current_by_id.get(r[0])==r for r in previous),(t,'old row changed')
            if t!='app_auditevent':assert len(current)-len(previous)==changes.get(t,0),(t,len(previous),len(current))
            preserved[t]={'before':len(previous),'after':len(current),'prior_rows_unchanged':True}
workbook=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/14_售后过程_模拟.xlsx';wb=load_workbook(workbook,read_only=True,data_only=True);cells=0;newrows=0
for ds in ['service_events','service_tasks','service_conditions']:
    spec=SCHEMAS[ds];sheet=wb[spec['label']];original=list(sheet.iter_rows(values_only=True));headers=original[0];mapping={f['label']:f for f in spec['fields']}
    for row in original[1:]:
        expected={mapping[label]['name']:convert(row[i],mapping[label]) for i,label in enumerate(headers)}
        record=Record.objects.select_related('source_row__batch').get(dataset=ds,business_key=expected['id']);assert expected==record.values;assert record.source_row.normalized==expected;assert record.record_hash==record.source_row.record_hash
        assert record.source_row.batch.file_hash==hashlib.sha256(workbook.read_bytes()).hexdigest();cells+=len(expected);newrows+=1
assert (newrows,cells)==(449,4042)
raw=list(Record.objects.filter(dataset='service').values_list('values',flat=True));tasks=list(Record.objects.filter(dataset='service_tasks').values_list('values',flat=True));conditions=list(Record.objects.filter(dataset='service_conditions').values_list('values',flat=True));events=list(Record.objects.filter(dataset='service_events').values_list('values',flat=True));cutoff=datetime.fromisoformat('2026-10-01T18:00:00')
d=s.current(s.filters({}));rows=d.selected();summary=d.summary(rows);tb=s.current(s.filters({'tab':'tasks'}));ts=tb.summary(tb.selected());assert not d.global_issues
assert (summary['objects'],summary['units'],summary['open'],summary['closed'],summary['attention'])==(45,45,15,30,0)
response=[];closed=[];cost=0
for r in raw:
    reported=datetime.fromisoformat(r['reported']);responded=datetime.fromisoformat(r['response']);finish=datetime.fromisoformat(r['closed']) if r['closed'] else None;assert reported<=responded<=cutoff
    response.append((responded-reported).total_seconds()/3600);cost+=r['cost_cents'];row=d.cases[r['id']]['row'];assert row['response_hours']==response[-1] and row['shipment_customer_match'] and not row['issues']
    if finish:
        assert responded<=finish<=cutoff;closed.append((finish-reported).total_seconds()/3600);assert row['closed_hours']==closed[-1]
    else:assert row['open_hours']==(cutoff-reported).total_seconds()/3600
assert summary['response_hours']==mean(response)==.75 and summary['closed_hours']==mean(closed)==8 and summary['cost_cents']==cost==1257000
assert sum(x['count'] for x in d.breakdown(rows)['failures'])==len(raw)
assert sum(x['count'] for x in d.breakdown(rows)['age'])==15
parent={r['id']:r for r in raw};done=[];opened=[];overdue=[];late=[]
for r in tasks:
    created=datetime.fromisoformat(r['created']);due=datetime.fromisoformat(r['due']);completed=datetime.fromisoformat(r['completed']) if r['completed'] else None;reported=datetime.fromisoformat(parent[r['service_id']]['reported']);assert reported<=created<=cutoff and due>=created
    if completed and completed<=cutoff:
        assert completed>=created;done.append(r['id'])
        if completed>due:late.append(r['id'])
    else:
        opened.append(r['id'])
        if due<cutoff:overdue.append(r['id'])
assert (len(tasks),len(done),len(opened),len(overdue),len(late))==(90,70,20,11,6)
assert (ts['objects'],ts['done'],ts['open'],ts['overdue'],ts['late'],ts['attention'])==(90,70,20,11,6,0)
assert len(events)==224 and len(conditions)==135 and sum(r['status']=='未提供' for r in conditions)==20
assert summary['with_events']==summary['with_tasks']==summary['with_conditions']==45
for r in events:
    assert datetime.fromisoformat(parent[r['service_id']]['reported'])<=datetime.fromisoformat(r['occurred'])<=cutoff
sources=[]
for r in rows:sources+=capture_sources(service_views.full_detail(d,r['id']))
sources={(r['dataset'],r['key']):r for r in sources};files={}
for (ds,key),src in sources.items():
    assert not src.get('missing');r=Record.objects.select_related('source_row__batch').get(dataset=ds,business_key=key);assert r.values==r.source_row.normalized and r.record_hash==r.source_row.record_hash;b=r.source_row.batch;files[str(b.pk)]=(b.file_path,b.file_hash)
for path,sha in files.values():assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==sha
exports={}
for filename in ['售后_逾期任务_浏览器导出.csv','售后_单日客户队列_浏览器导出.csv']:
    p=ROOT/'outputs'/filename;csvrows=list(csv.reader(p.open(encoding='utf-8-sig')));f=json.loads(csvrows[0][4]);actual=csvrows[3:]
    if f['tab']=='tasks':assert f['stage']=='overdue' and len(actual)==11 and {r[0] for r in actual}==set(overdue)
    else:assert f['customer']=='KH00010' and f['from']==f['to']=='2026-09-30' and len(actual)==1 and actual[0][0]=='SH2609-0045' and float(actual[0][6])==.75
    exports[filename]={'rows':len(actual),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
note=IssueDisposition.objects.get(key='service:SH2609-0045');assert note.version==1 and note.status=='待资料核对';assert parent['SH2609-0045']['status']=='跟进中' and parent['SH2609-0045']['closed'] is None
v=MetricVersion.objects.get(metric__key='DELIVERY_OTIF',version=4);assert v.status=='published' and v.calculation_hash==metric_registry.calculation_hash('bi_order_lines')
h=hashlib.sha256()
for r in Record.objects.order_by('dataset','business_key').values_list('dataset','business_key','values','revision','source_row_id').iterator():h.update(json.dumps(r,ensure_ascii=False,sort_keys=True).encode())
report={'synthetic':True,'baseline':str(baseline.relative_to(ROOT)),'business_rows':Record.objects.count(),'source_categories':Record.objects.values('dataset').distinct().count(),'business_sha256':h.hexdigest(),'preserved_tables':preserved,'new_workbook':{'rows':newrows,'cells':cells,'sha256':hashlib.sha256(workbook.read_bytes()).hexdigest()},'cases':summary,'tasks':ts,'events':len(events),'conditions':len(conditions),'sources':len(sources),'source_files':len(files),'browser_exports':exports,'metric_v4_unchanged':True,'tests':477,'new_tests':30,'limits':'Synthetic case cohort only; customer reported conditions, not verified measurements. No installed-population warranty rate, contractual SLA, root-cause approval, reopening, spare-part or acceptance workflow.'}
assert report['business_rows']==107389 and report['source_categories']==54
if len(sys.argv)>1:
    backup=Path(sys.argv[1])
    with zipfile.ZipFile(backup) as z,tempfile.TemporaryDirectory() as temp:
        manifest=json.loads(z.read('manifest.json'))
        for item in manifest['files']:
            blob=z.read(item['path']);assert len(blob)==item['size'] and hashlib.sha256(blob).hexdigest()==item['sha256']
        p=Path(temp)/'restored.sqlite3';p.write_bytes(z.read('data/platform.sqlite3'))
        with sqlite3.connect(p) as restored,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
            assert restored.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
            for t in tables:assert list(restored.execute(f'SELECT * FROM {t} ORDER BY id'))==list(now.execute(f'SELECT * FROM {t} ORDER BY id')),t
        report['backup']={'path':str(backup),'manifest_files':len(manifest['files']),'app_tables':len(tables),'sqlite_integrity':'ok','all_app_tables_match':True}
(ROOT/'data/service_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in report.items() if k!='preserved_tables'},ensure_ascii=False))
