"""Normal import, six-role read/export checks and exact preservation of old facts."""
import argparse,hashlib,json,os,shutil,sqlite3,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];BEFORE=ROOT/'data/oee-before';BOOK=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/38_班次设备效率_模拟.xlsx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def preserve(db):
 from app.schema import SCHEMAS
 manifest=json.loads((BEFORE/'manifest.json').read_text());assert sha(BEFORE/'platform.sqlite3')==manifest['database_sha256']
 for name,digest in manifest['physical'].items():assert sha(ROOT/name)==digest,name
 # oee.py was a new authoring stub already present when this checkpoint was frozen.
 changed={'app/oee.py','app/schema.py','app/manufacturing_rules.py','config/urls.py','app/views.py','templates/index.html','static/app.js','scripts/refine_bi_catalog.py','scripts/bi_reader_guide.py','scripts/export_bi_design.py','scripts/build_bi_business_guide.py'}
 for name,digest in manifest['files'].items():
  if name not in changed:assert sha(ROOT/name)==digest,name
 for key,schema in manifest['schemas'].items():assert SCHEMAS[key]==schema,key
 assert len(SCHEMAS)==135
 suffix='\nfrom .oee_schema import register_oee\nregister_oee(register)\n';s=(ROOT/'app/schema.py').read_text();assert s.endswith(suffix) and s[:-len(suffix)]==(BEFORE/'app/schema.py').read_text()
 branch="    if dataset in {'oee_studies','oee_windows','oee_events','oee_cycles','oee_outputs'}:\n        from .oee_contract import issues as oee_issues\n        return oee_issues(dataset,row)\n";s=(ROOT/'app/manufacturing_rules.py').read_text();assert s.count(branch)==1 and s.replace(branch,'')==(BEFORE/'app/manufacturing_rules.py').read_text()
 with sqlite3.connect(db) as conn:
  conn.execute('ATTACH DATABASE ? AS prior',(str(BEFORE/'platform.sqlite3'),));names=[r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")];assert set(names)==set(manifest['tables'])
  additions={}
  for name in names:
   if name in ('sqlite_sequence','app_importtemplate','app_importtemplateversion'):continue
   deleted=conn.execute('SELECT COUNT(*) FROM (SELECT * FROM prior."'+name+'" EXCEPT SELECT * FROM main."'+name+'")').fetchone()[0];assert deleted==0,(name,deleted)
   additions[name]=conn.execute('SELECT COUNT(*) FROM main."'+name+'"').fetchone()[0]-manifest['tables'][name]
  assert {k for k,v in additions.items() if v}<= {'app_record','app_importrow','app_importbatch','app_auditevent'},additions
  assert additions['app_record']==123 and additions['app_importrow']==123 and additions['app_importbatch']==1,additions
  template=conn.execute('SELECT code,revision,current_version FROM app_importtemplate').fetchall();assert len(template)==1 and template[0]==('TPL-SIM-DEPTS-001',14,7)
  cols=[r[1] for r in conn.execute('PRAGMA table_info(app_importtemplateversion)') if r[1]!='state'];columnlist=','.join('"'+c+'"' for c in cols);assert not conn.execute('SELECT '+columnlist+' FROM prior.app_importtemplateversion EXCEPT SELECT '+columnlist+' FROM main.app_importtemplateversion').fetchall()
 return dict(old_sql_rows_exact=True,old_contracts_exact=130,old_source_files_exact=458,protected_files_exact=len(manifest['files'])-len(changed),additions={k:v for k,v in additions.items() if v},sql_tables=len(names)-1)
def main():
 p=argparse.ArgumentParser();p.add_argument('--db',type=Path,required=True);p.add_argument('--copy',action='store_true');p.add_argument('--exports',action='store_true');p.add_argument('--output',type=Path,required=True);args=p.parse_args()
 if args.copy:assert not args.db.exists();shutil.copyfile(BEFORE/'platform.sqlite3',args.db)
 os.environ['MOTOR_SQLITE_PATH']=str(args.db.resolve());os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings');sys.path.insert(0,str(ROOT))
 import django;django.setup()
 from openpyxl import load_workbook
 from django.test import RequestFactory
 from django.urls import resolve
 from django.contrib.auth import get_user_model
 from app.ingestion import stage_file,commit_batch,convert
 from app.schema import SCHEMAS
 from app.models import Record,ImportBatch,ImportRow,AuditEvent
 from refresh_oee_template import refresh
 expected=json.loads((ROOT/'data/oee_scenario.json').read_text());book=load_workbook(BOOK,read_only=True,data_only=False)
 try:
  assert book.sheetnames==['导入说明']+[SCHEMAS[ds]['label'] for ds in expected['tables']]
  for ds,rows in expected['tables'].items():
   fields=SCHEMAS[ds]['fields'];matrix=list(book[SCHEMAS[ds]['label']].iter_rows(values_only=True));assert list(matrix[0])==[f['label'] for f in fields];converted=[{f['name']:convert(row[i],f) for i,f in enumerate(fields)} for row in matrix[1:]];assert converted==rows,ds
 finally:book.close()
 counts=(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count());batch,replayed=stage_file(BOOK);assert not replayed and batch.summary['total']==batch.summary['valid']==123;assert not batch.summary['unknown_sheets'];commit_batch(batch.pk);batch.refresh_from_db();assert batch.status=='committed';after=(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count());assert after==tuple(n+d for n,d in zip(counts,(123,1,123,2)));again,replayed=stage_file(BOOK);assert replayed and again.pk==batch.pk;commit_batch(again.pk);assert after==(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count())
 user=get_user_model().objects.get(username='demo_admin');template=refresh(user,ROOT);factory=RequestFactory();verified=[]
 def get(path,user,params=None):
  req=factory.get(path,params or {});req.user=user;match=resolve(path);response=match.func(req,**match.kwargs);assert response.status_code==200,(path,response.status_code,response.content[:100]);return response
 for role in ('admin','analyst','quality','operations','finance','viewer'):
  user=get_user_model().objects.get(username='demo_'+role)
  for profile in expected['profiles']:
   key=profile['id'];board=json.loads(get('/api/oee/'+key,user).content);assert board['state']==profile['state'] and board['summary']==profile['summary'];assert 'employee_id' not in board['resource'];assert all('price_cents' not in p for p in board['products'].values());q=dict(receipt=board['receipt']);get('/api/oee/'+key+'/sources',user,q)
   if board['windows']:get('/api/oee/'+key+'/windows/'+board['windows'][0]['id'],user,q)
   if args.exports and key in ('OE-260925-001','OE-260925-006'):
    export=get('/api/oee/'+key+'/export',user,dict(**q,format='json'));doc=json.loads(export.content);assert doc['result']['summary']==board['summary'] and 'receipt' not in doc
  verified.append(role)
 proof=dict(success=True,synthetic=True,rows=123,studies=10,typed_roundtrip_exact=True,normal_import_and_replay=True,template=template,roles=verified,cases_checked=60,browser_acceptance=False,mobile_acceptance=False,download_acceptance=False,preservation=preserve(args.db),workbook_sha256=sha(BOOK),database_sha256=sha(args.db))
 args.output.write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n');print(json.dumps(proof,ensure_ascii=False))
if __name__=='__main__':main()
