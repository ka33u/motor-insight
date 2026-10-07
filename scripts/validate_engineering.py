"""Read-only reconciliation of the engineering increment and original XLSX.

The independent counts below use imported facts and calendar dates, not the
workbench's summary implementation. No business facts are written by this file.
"""
import os,sys,json,csv,hashlib,sqlite3,zipfile,tempfile,math
from pathlib import Path
from collections import Counter,defaultdict
from datetime import date,datetime
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django
django.setup()
from openpyxl import load_workbook
from django.contrib.auth.models import User
from app.models import Record,ImportBatch,IssueDisposition,AuditEvent,DataQualityScan,DataMonitorPolicy,MetricVersion
from app.schema import SCHEMAS
from app.ingestion import convert
from app import engineering,data_health,metric_registry

def sha(v):return hashlib.sha256(json.dumps(v,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()

baseline=ROOT/'data/backups/motor-backup-20261003-205413.zip'
preserved={};source_files=0
expected_delta={'app_record':112,'app_importrow':112,'app_importbatch':1,'app_issuedisposition':1,'app_dataqualityscan':1}
with zipfile.ZipFile(baseline) as z,tempfile.TemporaryDirectory() as temp:
    manifest=json.loads(z.read('manifest.json'))
    for entry in manifest['files']:
        if entry['path'].startswith(('data/imports/','data/device_files/')):
            assert hashlib.sha256((ROOT/entry['path']).read_bytes()).hexdigest()==entry['sha256']
            source_files+=1
    p=Path(temp)/'before.sqlite3';p.write_bytes(z.read('data/platform.sqlite3'))
    with sqlite3.connect(p) as old,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
        tables=[r[0] for r in now.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'app_%'")]
        assert set(tables)=={r[0] for r in old.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'app_%'")}
        assert len(tables)==25
        for table in tables:
            before=list(old.execute(f'SELECT * FROM {table} ORDER BY id'))
            current=list(now.execute(f'SELECT * FROM {table} ORDER BY id'));by_id={r[0]:r for r in current}
            assert all(by_id.get(r[0])==r for r in before),(table,'prior row changed')
            if table!='app_auditevent':assert len(current)-len(before)==expected_delta.get(table,0),(table,'unexpected row count')
            preserved[table]={'before':len(before),'after':len(current),'prior_rows_unchanged':True}

bookfile=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/15_研发工艺过程_模拟.xlsx'
file_hash=hashlib.sha256(bookfile.read_bytes()).hexdigest()
assert file_hash=='acc1ecdb5e0bf997c3606c4c527528a11d38abc344cccfce62ca463dfdbb1504'
batch=ImportBatch.objects.get(file_hash=file_hash,status='committed')
assert hashlib.sha256(Path(batch.file_path).read_bytes()).hexdigest()==file_hash
assert batch.rows.count()==112 and batch.summary['committed']==112
book=load_workbook(bookfile,read_only=True,data_only=False);cells=0;new_counts={}
for ds in ['project_milestones','change_actions']:
    spec=SCHEMAS[ds];sheet=book[spec['label']];rows=sheet.iter_rows();header=next(rows)
    assert [c.value for c in header]==[f['label'] for f in spec['fields']]
    n=0
    for rownumber,row in enumerate(rows,2):
        values={}
        for cell,field in zip(row,spec['fields']):
            assert cell.data_type!='f'
            if field['type']=='date' and cell.value is not None:assert isinstance(cell.value,datetime)
            if field['name']=='id':assert cell.data_type=='s' and isinstance(cell.value,str)
            values[field['name']]=convert(cell.value,field);cells+=1
        rec=Record.objects.select_related('source_row').get(dataset=ds,business_key=values['id'])
        assert rec.values==values and rec.source_row.normalized==values
        assert rec.source_row.batch_id==batch.pk and rec.source_row.row_number==rownumber
        assert rec.source_row.sheet==spec['label'] and rec.revision==1
        n+=1
    new_counts[ds]=n
book.close()
assert (new_counts,cells)==({'project_milestones':84,'change_actions':28},1316)

scans=list(DataQualityScan.objects.order_by('created_at'));assert len(scans)==3
latest=scans[-1];snap=latest.snapshot;cutoff=datetime.fromisoformat(snap['as_of']);as_of=cutoff.date()
assert sha(snap)==latest.snapshot_hash and latest.engine_hash==data_health.engine_hash()
assert all(s.engine_hash!=latest.engine_hash for s in scans[:-1])
assert not snap['issues'] and not snap['unknown_datasets']
projected=data_health.project(latest,User.objects.get(username='demo_admin'))
assert not projected['stale'] and len(projected['rows'])==56
assert sum(not r['freshness']['has_target'] for r in projected['rows'])==55
policy=DataMonitorPolicy.objects.get(dataset='test_sessions');assert policy.version==2 and DataMonitorPolicy.objects.count()==1

identities=set(Record.objects.values_list('dataset','business_key'))
required=0;references=0;counts=Counter();data=defaultdict(list);facts_hash=hashlib.sha256()
for rec in Record.objects.select_related('source_row').order_by('dataset','business_key').iterator(chunk_size=1000):
    ds=rec.dataset;v=rec.values;counts[ds]+=1
    if ds in engineering.TABLES:data[ds].append(v)
    assert v['id']==rec.business_key and v==rec.source_row.normalized
    assert sha(v)==rec.record_hash==rec.source_row.record_hash
    assert (ds,rec.business_key)==(rec.source_row.dataset,rec.source_row.business_key)
    facts_hash.update(json.dumps((ds,rec.business_key,v,rec.revision,rec.source_row_id),ensure_ascii=False,sort_keys=True).encode())
    for field in SCHEMAS[ds]['fields']:
        val=v.get(field['name']);missing=val is None or isinstance(val,str) and not val.strip()
        if field['required']:assert not missing;required+=1
        if missing:continue
        typ=field['type']
        if typ=='int':assert type(val) is int
        elif typ=='bool':assert type(val) is bool
        elif typ=='float':assert type(val) in [int,float] and math.isfinite(val)
        elif typ=='str':assert isinstance(val,str) and val==val.strip() and not val.startswith('=')
        elif typ in ['date','datetime']:
            parsed=(date if typ=='date' else datetime).fromisoformat(val);assert parsed.isoformat()==val
        if field['reference']:assert (field['reference'],val) in identities;references+=1
assert (sum(counts.values()),len(counts))==(107501,56)
for ds,n in counts.items():assert snap['datasets'][ds]['rows']==n
assert sum(r['required_cells'] for r in projected['rows'])==required
assert sum(f['reference_present'] for d in snap['datasets'].values() for f in d['fields'].values())==references

def completed(r,key):return bool(r.get(key) and date.fromisoformat(r[key])<=as_of)
projects={r['id']:r for r in data['projects']};changes={r['id']:r for r in data['engineering_changes']}
milestones=defaultdict(list);actions=defaultdict(list)
for r in data['project_milestones']:milestones[r['project_id']].append(r)
for r in data['change_actions']:actions[r['change_id']].append(r)
project_counts={}
for key,r in projects.items():
    ms=milestones[key];done=[x for x in ms if completed(x,'actual_end')]
    overdue=[x for x in ms if not completed(x,'actual_end') and date.fromisoformat(x['planned_end'])<as_of]
    assert len(ms)==6 and len({x['sequence'] for x in ms})==6
    project_counts[key]={'milestones':len(ms),'completed_milestones':len(done),'overdue_milestones':len(overdue),'closed':completed(r,'actual_end'),'overdue':not completed(r,'actual_end') and date.fromisoformat(r['planned_end'])<as_of}
change_counts={}
for key,r in changes.items():
    tasks=[x for x in actions[key] if date.fromisoformat(x['created'])<=as_of]
    done=[x for x in tasks if completed(x,'completed')]
    overdue=[x for x in tasks if not completed(x,'completed') and date.fromisoformat(x['due'])<as_of]
    assert len(tasks)==4
    change_counts[key]={'actions':len(tasks),'completed_actions':len(done),'pending_actions':len(tasks)-len(done),'overdue_actions':len(overdue)}
b=engineering.Engineering(data,engineering.filters({}),cutoff=cutoff.isoformat())
assert not b.global_issues
for key,count in project_counts.items():
    r=b.projects[key]['row'];assert not r['issues']
    for field in ['milestones','completed_milestones','overdue_milestones']:assert r[field]==count[field]
    assert ('closed' in r['flags'])==count['closed'] and ('overdue' in r['flags'])==count['overdue']
for key,count in change_counts.items():
    r=b.changes[key]['row'];assert not r['issues'] and r['before_bom_rows']==0 and r['after_bom_rows']==9
    for field,n in count.items():assert r[field]==n
    assert r['status']=='已批准'
for key,obj in b.products.items():
    assert not obj['row']['issues'] and obj['graph']['valid']
    routes=obj['routes'];edges=obj['graph']['edges'];assert (len(routes),len(edges),len(obj['bom']))==(11,10,9)
    assert Counter(r['branch'] for r in routes)==Counter({'定子':4,'转子':4,'整机':3})
    ranks={n['id']:n['rank'] for n in obj['graph']['nodes']};assert all(ranks[e['from_route_id']]<ranks[e['to_route_id']] for e in edges)
    assembly=next(r['id'] for r in routes if r['process']=='装配')
    assert len([e for e in edges if e['to_route_id']==assembly])==2
    assert all(w['version_rows_present'] and not w['version_difference'] for w in obj['work_orders'])
assert (len(b.products),sum(len(p['bom']) for p in b.products.values()),sum(len(p['routes']) for p in b.products.values()))==(48,432,528)
assert (sum(x['closed'] for x in project_counts.values()),sum(x['overdue'] for x in project_counts.values()),sum(x['completed_milestones'] for x in project_counts.values()),sum(x['overdue_milestones'] for x in project_counts.values()))==(9,5,71,13)
assert (sum(x['completed_actions'] for x in change_counts.values()),sum(x['pending_actions'] for x in change_counts.values()),sum(x['overdue_actions'] for x in change_counts.values()))==(20,8,8)

exports={}
for tab,name,expected in [('projects','研发_逾期项目_浏览器导出.csv',{k for k,v in project_counts.items() if v['overdue']}),('changes','工艺_变更执行待办_浏览器导出.csv',{k for k,v in change_counts.items() if v['overdue_actions']})]:
    p=ROOT/'outputs'/name;rows=list(csv.reader(p.open(encoding='utf-8-sig')))
    assert rows[0][2]==cutoff.isoformat();f=json.loads(rows[0][4]);assert f['tab']==tab and f['stage']=='overdue'
    assert not any(f[k] for k in ['family','product','owner','q','from','to'])
    entries=[dict(zip(rows[2],r)) for r in rows[3:]]
    idfield='项目编号' if tab=='projects' else '变更单号';assert {r[idfield] for r in entries}==expected
    for r in entries:
        if tab=='projects':
            key=r[idfield];c=project_counts[key]
            assert [int(r[n]) for n in ['阶段数','有完成记录阶段数','未完成逾期阶段数']]==[c[k] for k in ['milestones','completed_milestones','overdue_milestones']]
            assert int(r['项目台账预算分'])==projects[key]['budget_cents']
            assert int(r['项目逾期自然日'])==(as_of-date.fromisoformat(projects[key]['planned_end'])).days
        else:
            c=change_counts[r[idfield]];assert [int(r[n]) for n in ['执行任务数','有完成记录任务数','未完成任务数','逾期未完成任务数']]==[c[k] for k in ['actions','completed_actions','pending_actions','overdue_actions']]
            assert r['原单批准标记']=='已批准' and r['原号BOM明细数']=='0' and r['目标号BOM明细数']=='9'
    exports[name]={'rows':len(entries),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
follow=IssueDisposition.objects.get(key='engineering:changes:ECN2609-022')
assert (follow.version,follow.status,follow.owner)==(1,'待执行反馈','研发工艺岗位（模拟）')
event=AuditEvent.objects.get(action='engineering.followup',object_id=follow.key)
assert event.detail['business_facts_changed'] is False and event.detail['after']['version']==1
version=MetricVersion.objects.get(metric__key='DELIVERY_OTIF',version=4)
assert version.status=='published' and version.calculation_hash==metric_registry.calculation_hash('bi_order_lines')=='f6815a706b1f19c2e2be0af931c640d869b45b17bc0aa1f72d7fe2144356868f'
assert ImportBatch.objects.exclude(status='superseded').count()==15
report={'synthetic':True,'baseline':str(baseline.relative_to(ROOT)),'business_rows':sum(counts.values()),'source_categories':len(counts),'active_workbooks':15,'business_sha256':facts_hash.hexdigest(),'preserved_tables':preserved,'preserved_source_files':source_files,'new_workbook':{'batch':str(batch.pk),'sha256':file_hash,'rows':new_counts,'cells':cells},'engineering':{'configurations':48,'bom_rows':432,'routes':528,'dependencies':480,'work_orders':300,'projects':14,'closed_projects':9,'overdue_projects':5,'milestones':84,'completed_milestones':71,'overdue_milestones':13,'changes':7,'changes_with_overdue_actions':4,'actions':28,'completed_actions':20,'overdue_actions':8,'missing_historical_bom':7},'latest_scan':str(latest.pk),'latest_scan_sha256':latest.snapshot_hash,'basic_issues':0,'checked_required_cells':required,'checked_references':references,'preserved_monitor_policy_version':2,'browser_exports':exports,'coordination_version':1,'metric_v4_unchanged':True,'tests':546,'new_tests':33,'limits':'Synthetic coordination only: no PLM approval, historical BOM differences, confirmed change applicability, physical implementation verification, project actual spend, or real U8/MES connection.'}
if len(sys.argv)>1:
    backup=Path(sys.argv[1])
    with zipfile.ZipFile(backup) as z,tempfile.TemporaryDirectory() as temp:
        manifest=json.loads(z.read('manifest.json'))
        for item in manifest['files']:
            blob=z.read(item['path']);assert len(blob)==item['size'] and hashlib.sha256(blob).hexdigest()==item['sha256']
        p=Path(temp)/'restored.sqlite3';p.write_bytes(z.read('data/platform.sqlite3'))
        with sqlite3.connect(p) as restored,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
            assert restored.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
            for table in tables:assert list(restored.execute(f'SELECT * FROM {table} ORDER BY id'))==list(now.execute(f'SELECT * FROM {table} ORDER BY id')),table
        report['backup']={'path':str(backup),'manifest_files':len(manifest['files']),'app_tables':len(tables),'sqlite_integrity':'ok','all_app_tables_match':True}
(ROOT/'data/engineering_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in report.items() if k!='preserved_tables'},ensure_ascii=False,indent=2))
