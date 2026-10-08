"""Validate the normally imported mixed-product XLSX and independently reconcile results."""
import hashlib
import json
import os
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_CEILING
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
import django
django.setup()
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert
from app import joint_schedule_data, joint_material_evidence, joint_comparison
from validate_joint_material_http import oracle

BOOK = '45_三品种换型与交期_模拟.xlsx'
KEYS = ('MP-261016-001',)
JOBS = tuple(f'SP-261016-001-B{n:03}' for n in (1, 2, 3))
MATERIAL = '01.01.0002'


def verify():
    book_path = ROOT/'demo/workbooks'/BOOK
    workbook = load_workbook(book_path, read_only=True, data_only=False)
    schemas = {s['label']: s for s in SCHEMAS.values()}
    expected = {}
    try:
        assert len(workbook.sheetnames) == 17
        for sheet in workbook:
            if sheet.title == '导入说明':
                continue
            schema = schemas[sheet.title]
            rows = list(sheet.iter_rows(values_only=True))
            assert list(rows[0]) == [f['label'] for f in schema['fields']]
            for number, row in enumerate(rows[1:], 2):
                values = {f['name']: convert(v, f) for v, f in zip(row, schema['fields'])}
                identity = schema['key'], values['id']
                assert identity not in expected
                expected[identity] = values
                record = Record.objects.select_related('source_row__batch').get(dataset=identity[0], business_key=identity[1])
                source = record.source_row
                assert record.values == values == source.normalized
                assert (source.batch.filename, source.sheet, source.row_number) == (BOOK, sheet.title, number)
                assert source.batch.file_hash == hashlib.sha256(book_path.read_bytes()).hexdigest()
    finally:
        workbook.close()
    assert len(expected) == 847
    all_rows = list(Record.objects.order_by('dataset', 'business_key').values_list('dataset', 'business_key', 'record_hash'))
    old = [r for r in all_rows if (r[0], r[1]) not in expected]
    h = hashlib.sha256()
    for row in old:
        h.update((json.dumps(row, ensure_ascii=False)+'\n').encode())
    assert len(old) == 232941 and h.hexdigest() == '2e0a72b20acd7e72393fa0433a170f3349dbf7698d583b6e20aec5060b44d473'
    assert len(all_rows) == 233788
    outputs, profiles = {}, []
    finish = {'due': ['2026-10-17T11:44:42', '2026-10-17T17:42:32', '2026-10-18T09:43:00'],
              'priority': ['2026-10-17T11:44:42', '2026-10-18T09:49:36', '2026-10-17T14:31:54']}
    for policy in ('due', 'priority'):
        loaded = joint_schedule_data.load(KEYS[0], policy)
        r = loaded['result']
        outputs[policy] = r
        assert r['state'] == 'trial' and not r['issues']
        assert r['summary'] == dict(tasks=96, scheduled_tasks=96, blocked_tasks=0, jobs=3, qty=24,
            complete_jobs=3, late_jobs=int(policy == 'priority'), blocked_jobs=0, demands=27, reserved_demands=27,
            direct_shortage_tasks=0, material_delayed_tasks=0)
        assert [j['id'] for j in r['jobs']] == list(JOBS)
        assert [j['finished'] for j in r['jobs']] == finish[policy]
        assert [j['qty'] for j in r['jobs']] == [6, 8, 10]
        assert abs(r['jobs'][1]['lateness_minutes']-(949.6 if policy == 'priority' else 0)) < 1e-9
        bt = loaded['parent']['base']['tables']
        jobs = {j['id']: j for j in bt['schedule_jobs']}
        tasks = {t['id']: t for t in bt['schedule_tasks']}
        options = {o['id']: o for o in bt['schedule_options']}
        assert [jobs[k]['family'] for k in JOBS] == ['FRAME-A', 'FRAME-B', 'FRAME-A']
        assert len(options) == len(tasks) == 96
        groups = defaultdict(list)
        for task in r['tasks']:
            groups[task['resource_id']].append(task)
            seconds = (Decimal(str(tasks[task['id']]['unit_minutes']))*tasks[task['id']]['lot_qty']*60).to_integral_value(rounding=ROUND_CEILING)
            assert abs(task['process_minutes']*60-float(seconds)) < 1e-8
        initial, changes, setup = 0, 0, Decimal(0)
        for rows in groups.values():
            family = None
            for task in sorted(rows, key=lambda t: t['started']):
                current = jobs[task['job_id']]['family']
                change = current != family
                initial += int(family is None)
                changes += int(family is not None and change)
                minutes = Decimal(str(options[task['option_id']]['setup_minutes'])) if change else Decimal(0)
                assert Decimal(str(task['setup_minutes'])) == minutes
                assert (datetime.fromisoformat(task['process_started'])-datetime.fromisoformat(task['started'])).total_seconds() == float(minutes*60)
                setup += minutes
                family = current
        assert initial == len(groups) == 10
        assert changes == (20 if policy == 'due' else 10)
        assert setup == (360 if policy == 'due' else 240)
        for field, aggregate in [('resource_id', 'resources'), ('employee_id', 'workers')]:
            by_id = defaultdict(list)
            for task in r['tasks']:
                by_id[task[field]].append(task)
            for identity, rows in by_id.items():
                ordered = sorted(rows, key=lambda t: t['started'])
                assert all(a['finished'] <= b['started'] for a, b in zip(ordered, ordered[1:]))
                busy = sum((datetime.fromisoformat(t['finished'])-datetime.fromisoformat(t['started'])).total_seconds()/60 for t in rows)
                assert abs(next(x for x in r[aggregate] if x['id'] == identity)['busy_minutes']-busy) < 1e-8
        bom = {b['id']: b for b in loaded['references']['bom']}
        for d in r['demands']:
            b, step = bom[d['bom_id']], Decimal(str(d['quantum']))
            qty = (Decimal(str(jobs[d['job_id']]['qty']))*Decimal(str(b['qty']))*(1+Decimal(str(b['scrap_allowance'])))/step).to_integral_value(rounding=ROUND_CEILING)*step
            assert Decimal(str(d['required_qty'])) == qty
        for balance in r['balances']:
            evidence, _ = joint_material_evidence.build(r, balance['material_id'], balance['unit'])
            oracle(r, evidence)
            assert Decimal(balance['required_qty']) == Decimal(balance['reserved_qty'])
            assert Decimal(balance['total_remaining_qty']) == 0
        profiles.append(dict(policy=policy, initial_setups=initial, cross_family_changes=changes,
            setup_minutes=float(setup), summary=r['summary'], jobs=r['jobs']))
    comparison = joint_comparison.compare(outputs['due'], outputs['priority'])
    assert all(comparison['summary'][k] == 1 for k in ('earlier', 'later', 'same'))
    assert max(t['finished'] for t in outputs['priority']['tasks']) > max(t['finished'] for t in outputs['due']['tasks'])
    proof = dict(success=True, source_rows=847, records=233788, protected_old_facts=232941,
        profiles=profiles, independent_bom_checks=54, independent_material_checks=48,
        source_rows_exact=True, original_ids_preserved=True, setup_reconciled_by_resource=True,
        resource_worker_nonoverlap=True, browser_acceptance=False)
    (ROOT/'data/changeover_validation.json').write_text(json.dumps(proof, ensure_ascii=False, indent=2)+'\n')
    return proof


if __name__ == '__main__':
    print(json.dumps(verify(), ensure_ascii=False))
