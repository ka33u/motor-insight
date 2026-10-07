"""All-cell typed parity, native ingestion and scoped APIs on an isolated copy."""
import csv,hashlib,io,json,os,sqlite3,sys,tempfile,time
from collections import Counter
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def rehearse():
 started=time.monotonic();main=ROOT/'data/platform.sqlite3';original=sha(main)
 with tempfile.TemporaryDirectory(prefix='motor-device-intake-') as temp:
  copy=Path(temp)/'platform.sqlite3'
  with sqlite3.connect('file:'+str(main)+'?mode=ro',uri=True) as src,sqlite3.connect(copy) as dst:src.backup(dst)
  os.environ['MOTOR_SQLITE_PATH']=str(copy);os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
  import django;django.setup()
  from django.test import RequestFactory,override_settings
  from django.contrib.auth.models import User
  from django.db import connection
  from openpyxl import load_workbook
  from app import device_intake as eng,device_intake_views as views,metric_registry,import_templates
  from app.models import Record,MetricVersion,DeviceFileReview,FileReadGrant,ImportTemplate
  from app.ingestion import convert,stage_file,commit_batch
  from app.schema import SCHEMAS
  file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/33_设备目录与待采集_模拟.xlsx';expected=json.loads((ROOT/'data/device_intake_scenario.json').read_text())['tables'];parsed={};cells=dates=0
  book=load_workbook(file,read_only=True,data_only=False)
  for ds in eng.DATASETS:
   schema=SCHEMAS[ds];iterator=book[schema['label']].iter_rows(values_only=True);assert list(next(iterator))==[f['label'] for f in schema['fields']];parsed[ds]=[]
   for row in iterator:
    assert len(row)==len(schema['fields']);assert not any(isinstance(v,str) and v.startswith('=') for v in row)
    parsed[ds].append({f['name']:convert(v,f) for f,v in zip(schema['fields'],row)})
    dates+=sum(f['type']=='datetime' and v is not None and not isinstance(v,str) for f,v in zip(schema['fields'],row));cells+=len(row)
   assert parsed[ds]==expected[ds],ds
  book.close();count=sum(map(len,parsed.values()));old=list(Record.objects.order_by('id').values());reviews=DeviceFileReview.objects.count();grants=FileReadGrant.objects.count();admin=User.objects.get(username='demo_admin');factory=RequestFactory()
  assert not Record.objects.filter(dataset__in=eng.DATASETS).exists()
  with override_settings(BASE_DIR=Path(temp),DEVICE_FILE_ROOT=ROOT/'data/device_files'):
   batch,repeated=stage_file(file);assert not repeated and batch.summary==dict(valid=count,total=count,unknown_sheets=[]),batch.summary
   commit_batch(batch.pk);again,repeated=stage_file(file);assert repeated and again.pk==batch.pk;commit_batch(again.pk)
   assert list(Record.objects.exclude(dataset__in=eng.DATASETS).order_by('id').values())==old
   for ds,rows in parsed.items():assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values',flat=True))==sorted(rows,key=lambda r:r['id'])
   def request(fn,args=None,user=admin,key=None):
    r=factory.get('/api/device-intake',args or {});r.user=user;return fn(r,**({'key':key} if key else {}))
   d=eng.Workspace(admin);response=request(views.board);assert response.status_code==200;board=json.loads(response.content);keys=[]
   for page in range(1,(board['total']+24)//25+1):
    r=json.loads(request(views.board,dict(page=str(page))).content);assert r['receipt']==d.receipt;keys.extend(x['key'] for x in r['rows'])
   assert keys==[r['key'] for r in d.selected()] and len(keys)==len(set(keys))
   out=request(views.export,dict(receipt=d.receipt));assert out.status_code==200;csvrows=list(csv.reader(io.StringIO(out.content.decode('utf-8-sig'))));assert [r[0] for r in csvrows[4:]]==keys
   # Independent source population and fingerprint union; never infer file existence.
   scanmap={r['id']:r for r in parsed['device_scan_runs']};latest={}
   for r in scanmap.values():
    if r['started']<=eng.AS_OF and (r['source_id'] not in latest or r['started']>latest[r['source_id']]['started']):latest[r['source_id']]=r
   independent=[r for r in parsed['device_file_observations'] if r['scan_id'] in {r['id'] for r in latest.values()} and r['discovered']<=eng.AS_OF]
   assert set(keys)=={r['id'] for r in independent}
   assert d.summary(d.selected())['known_contents']==len({(r['sha256'].lower(),r['size_bytes']) for r in independent if eng.fingerprint(r)})
   allrows=eng.Workspace(admin,{'mode':'all'}).selected();assert len(allrows)==196
   states={s['id']:s['state'] for s in d.source_rows};assert [states[f'DS-PC-2026-{i:03}'] for i in (8,9,10,11,12)]==['partial','empty','not_executed','failed','not_scanned']
   details=[]
   for row in d.selected():
    response=request(views.detail,dict(receipt=d.receipt),key=row['key']);assert response.status_code==200,(row['key'],response.content[:200]);value=json.loads(response.content);assert value['row']['provenance']['row']>=2
    details.append(dict(key=row['key'],archives=len(value['row']['archives']),checks=len(value['checks'])))
   quality=User.objects.get(username='demo_quality');q=eng.Workspace(quality);assert sum(len(r['archives']) for r in q.rows)==1;assert sum(len(r['archives']) for r in d.rows)==3
   assert request(views.export,dict(receipt=d.receipt),quality).status_code==400
   profiles=[]
   for role in ('admin','analyst','quality','operations','finance','viewer'):
    u=User.objects.filter(groups__name=role,is_active=True).first();status=request(views.board,user=u).status_code;assert status==(200 if role in ('admin','quality') else 403);profiles.append(dict(role=role,status=status))
   for metric in MetricVersion.objects.filter(status='published',metric__key='DELIVERY_OTIF',version=8):assert metric.calculation_hash==metric_registry.calculation_hash(metric.metric.dataset)=='9d9a5af3d45172eca16790d628a36bc6acbe18d1e2a811a66a5b5fbff7037fe3'
   assert DeviceFileReview.objects.count()==reviews and FileReadGrant.objects.count()==grants
   # Templates bind the complete global schema catalog. An additive source
   # intentionally requires a new header review; never mutate active v1.
   template=import_templates.info(ImportTemplate.objects.get(code='TPL-SIM-DEPTS-001'),admin,True)
   assert template['versions'][0]['rules_current'] is False and template['versions'][0]['can_use'] is False
   from refresh_device_intake_template import refresh
   refreshed=refresh(admin,ROOT);assert refreshed['fresh_example_can_use'] and refreshed['old_content_hash']==template['versions'][0]['content_hash']
   assert len(old)==212466 and Record.objects.count()==len(old)+count
  connection.close();assert sha(main)==original
  proof=dict(synthetic=True,live_collection=False,workbook_sha256=sha(file),rows=count,tables={k:len(v) for k,v in parsed.items()},cells=cells,native_date_cells=dates,all_cell_typed_parity=True,native_stage_commit_and_replay=True,all_old_records_unchanged=True,main_database_sha256=original,main_database_unchanged=True,all_page_and_csv_keys_match=True,independent_current_scan_keys_match=True,current_summary=d.summary(d.selected()),all_scans_observations=len(allrows),all_current_details=len(details),role_profiles=profiles,private_archive_matches=dict(admin=3,quality=1),no_associations_or_grants_added=True,published_v8_hash_unchanged=True,existing_template_requires_fresh_header_review=True,elapsed_seconds=round(time.monotonic()-started,3),browser_mobile_and_actual_download_accepted=False)
  proof['existing_template_native_v2_refresh']=refreshed
  (ROOT/'data/device_intake_import_rehearsal.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in proof.items() if k!='existing_template_native_v2_refresh'},ensure_ascii=False,indent=2));return proof
if __name__=='__main__':rehearse()
