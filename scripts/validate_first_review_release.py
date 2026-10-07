"""Normal Excel import into rehearsal/main, preserving all pre-existing facts."""
import argparse
import hashlib
import json
import os
import shutil
import sqlite3
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/42_首件复核登记_模拟.xlsx'
REHEARSAL = Path('/private/tmp/motorinsight-first-review-rehearsal.sqlite3')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['rehearse', 'release'])
    args = parser.parse_args()
    before = json.loads((ROOT/'data/first-review-before/manifest.json').read_text())
    main_db = ROOT/'data/platform.sqlite3'
    assert sha(main_db) == before['source_database_sha'], 'Inspect main database changes before importing.'
    if args.mode == 'rehearse':
        assert not REHEARSAL.exists(), 'Do not overwrite a rehearsal.'
        shutil.copy2(main_db, REHEARSAL)
        db = REHEARSAL
    else:
        for name in ('first_review_rehearsal.json', 'first_review_bootstrap.json', 'first_review_http.json', 'first_review_ui.json'):
            assert json.loads((ROOT/'data'/name).read_text())['success'], name
        assert '\nOK\n' in Path('/private/tmp/motorinsight-first-review-full-tests.log').read_text()
        db = main_db
    os.environ['DJANGO_SETTINGS_MODULE'] = 'config.settings'
    os.environ['MOTOR_SQLITE_PATH'] = str(db)
    sys.path.insert(0, str(ROOT))
    import django
    django.setup()
    from openpyxl import load_workbook
    from django.contrib.auth.models import User
    from app.ingestion import stage_file, commit_batch, convert
    from app.models import Record, ImportBatch, ImportRow, AuditEvent, MetricVersion
    from app.schema import SCHEMAS
    from app import first_piece_data, first_review, finite_schedule, launch_data
    from app.metric_registry import calculation_hash
    from validate_joint_schedule_release import frozen_tables
    from refresh_first_review_template import refresh
    scenario = json.loads((ROOT/'data/first_review_scenario.json').read_text())
    expected = scenario['tables'][first_review.DATASET]
    fields = SCHEMAS[first_review.DATASET]['fields']
    wb = load_workbook(BOOK, read_only=True, data_only=False)
    try:
        assert wb.sheetnames == ['首件复核登记', '导入说明']
        raw = list(wb['首件复核登记'].iter_rows(values_only=True))
        assert list(raw[0]) == [f['label'] for f in fields]
        actual = [{f['name']: convert(row[i], f) for i, f in enumerate(fields)} for row in raw[1:]]
        assert actual == expected and len(actual) == 293
    finally:
        wb.close()
    frozen = frozen_tables(db)
    counts = (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    batch, repeated = stage_file(BOOK)
    assert not repeated and batch.summary['valid'] == batch.summary['total'] == 293, batch.summary
    batch = commit_batch(batch.pk)
    assert batch.status == 'committed'
    after = (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    assert after == tuple(n+d for n, d in zip(counts, (293, 1, 293, 2)))
    same, repeated = stage_file(BOOK)
    assert repeated and same.pk == batch.pk
    commit_batch(same.pk)
    assert after == (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    assert list(Record.objects.filter(dataset=first_review.DATASET).order_by('business_key').values_list('values', flat=True)) == sorted(expected, key=lambda r: r['id'])
    d = first_piece_data.load()
    assert d['reviews']['summary'] == scenario['summary']
    assert {k: v['state'] for k, v in d['reviews']['rows'].items()} == scenario['expected']
    assert d['result'] == json.loads((ROOT/'data/first-review-before/first_piece.json').read_text())
    old_launch = json.loads((ROOT/'data/first-review-before/launch.json').read_text())
    launch = launch_data.load('LR-261001-001')
    assert launch['result'] == old_launch['result'] and launch['source_hash'] == old_launch['source_hash']
    template = refresh(User.objects.get(username='demo_admin'), ROOT)
    old = list(Record.objects.exclude(dataset=first_review.DATASET).order_by('dataset', 'business_key').values('dataset', 'business_key', 'record_hash'))
    assert len(old) == before['record_count'] and finite_schedule.digest(old) == before['record_digest']
    assert {k: SCHEMAS[k] for k in before['schemas']} == before['schemas']
    assert set(SCHEMAS)-set(before['schemas']) == {first_review.DATASET}
    for path, digest in before['originals'].items():
        assert sha(ROOT/path) == digest
    assert calculation_hash('bi_order_lines') == MetricVersion.objects.get(pk=11).calculation_hash == '9d9a5af3d45172eca16790d628a36bc6acbe18d1e2a811a66a5b5fbff7037fe3'
    assert frozen_tables(db) == frozen
    with sqlite3.connect(db) as c:
        assert c.execute('PRAGMA integrity_check').fetchall() == [('ok',)] and c.execute('PRAGMA foreign_key_check').fetchall() == []
    if args.mode == 'rehearse':
        assert sha(main_db) == before['source_database_sha']
    proof = dict(success=True, mode=args.mode, records=Record.objects.count(), new_rows=293, old_facts_preserved=len(old),
                 old_schemas_preserved=148, original_files_preserved=456, workbook_sha256=sha(BOOK), all_excel_fields_exact=True,
                 first_piece_and_launch_unchanged=True, published_metric_unchanged=True, review_summary=d['reviews']['summary'],
                 template=template, batch_id=str(batch.pk), browser_acceptance=False)
    (ROOT/('data/first_review_'+('rehearsal' if args.mode == 'rehearse' else 'release')+'.json')).write_text(json.dumps(proof, ensure_ascii=False, indent=2)+'\n')
    (ROOT/'data/first_review_example.json').write_text(json.dumps(d, ensure_ascii=False)+'\n')
    print(json.dumps(proof, ensure_ascii=False))


if __name__ == '__main__':
    main()
