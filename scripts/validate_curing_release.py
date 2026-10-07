"""Normal XLSX import and exact pre-existing-row preservation."""
import argparse,hashlib,json,os,shutil,sqlite3,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
BOOK=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/43_固化温度曲线_模拟.xlsx'
REHEARSAL=Path('/private/tmp/motorinsight-curing-rehearsal-20261007.sqlite3')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def row_digest(c,table,maxrow):
    h=hashlib.sha256()
    for r in c.execute('SELECT * FROM "'+table+'" WHERE rowid<=? ORDER BY rowid',(maxrow,)):h.update((json.dumps(r,ensure_ascii=False,default=str)+'\n').encode())
    return h.hexdigest()
def main():
    parser=argparse.ArgumentParser();parser.add_argument('mode',choices=['rehearse','release']);args=parser.parse_args()
    before=json.loads((ROOT/'data/curing-before/manifest.json').read_text());frozen=json.loads((ROOT/'data/curing-before/tables.json').read_text());main_db=ROOT/'data/platform.sqlite3'
    assert sha(main_db)==before['source_database_sha'],'Inspect intervening main database changes.'
    if args.mode=='rehearse':
        assert not REHEARSAL.exists();shutil.copy2(main_db,REHEARSAL);db=REHEARSAL
    else:
        for name in ('curing_rehearsal','curing_bootstrap','curing_http','curing_ui','curing_design'):
            assert json.loads((ROOT/'data'/f'{name}.json').read_text())['success'],name
        assert '\nOK\n' in Path('/private/tmp/motorinsight-curing-full-tests.log').read_text();db=main_db
    os.environ['MOTOR_SQLITE_PATH']=str(db);os.environ['DJANGO_SETTINGS_MODULE']='config.settings';sys.path.insert(0,str(ROOT))
    import django;django.setup()
    from openpyxl import load_workbook
    from app.ingestion import convert,stage_file,commit_batch
    from app.models import Record,ImportBatch,MetricVersion
    from app.schema import SCHEMAS
    from app import curing_data,first_piece_data,launch_data,finite_schedule
    from app.metric_registry import calculation_hash
    from django.contrib.auth.models import User
    from refresh_curing_template import refresh
    scenario=json.loads((ROOT/'data/curing_scenario.json').read_text());datasets=set(scenario['tables']);total=sum(scenario['counts'].values())
    wb=load_workbook(BOOK,read_only=True,data_only=False)
    try:
        assert wb.sheetnames==[SCHEMAS[ds]['label'] for ds in scenario['tables']]+['导入说明']
        for ds,expected in scenario['tables'].items():
            raw=list(wb[SCHEMAS[ds]['label']].iter_rows(values_only=True));fields=SCHEMAS[ds]['fields'];assert list(raw[0])==[f['label'] for f in fields]
            actual=[{f['name']:convert(r[i],f) for i,f in enumerate(fields)} for r in raw[1:]];assert actual==expected,ds
    finally:wb.close()
    n=Record.objects.count();imports=ImportBatch.objects.count();batch,repeated=stage_file(BOOK)
    assert not repeated and batch.summary['valid']==batch.summary['total']==total,batch.summary
    batch=commit_batch(batch.pk);assert batch.status=='committed';assert Record.objects.count()==n+total and ImportBatch.objects.count()==imports+1
    same,repeated=stage_file(BOOK);assert repeated and same.pk==batch.pk;commit_batch(same.pk);assert Record.objects.count()==n+total and ImportBatch.objects.count()==imports+1
    d=curing_data.load();assert {r['run']['id']:r['state'] for r in d['result']['rows']}==scenario['expected']
    for ds,expected in scenario['tables'].items():assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values',flat=True))==sorted(expected,key=lambda r:r['id'])
    assert first_piece_data.load()==json.loads((ROOT/'data/curing-before/first_piece.json').read_text())
    assert launch_data.load('LR-261001-001')==json.loads((ROOT/'data/curing-before/launch.json').read_text())
    template=refresh(User.objects.get(username='demo_admin'),ROOT)
    old=list(Record.objects.exclude(dataset__in=datasets).order_by('dataset','business_key').values('dataset','business_key','record_hash'))
    assert len(old)==before['record_count'] and finite_schedule.digest(old)==before['record_digest']
    assert {k:SCHEMAS[k] for k in before['schemas']}==before['schemas'] and set(SCHEMAS)-set(before['schemas'])==datasets
    for p,h in before['originals'].items():assert sha(ROOT/p)==h
    assert calculation_hash('bi_order_lines')==MetricVersion.objects.get(pk=11).calculation_hash=='9d9a5af3d45172eca16790d628a36bc6acbe18d1e2a811a66a5b5fbff7037fe3'
    with sqlite3.connect(db) as c:
        for table,info in frozen.items():
            if table in ('sqlite_sequence','app_importtemplate','app_importtemplateversion'):continue
            assert row_digest(c,table,info['max_rowid'])==info['digest'],table
            count=c.execute('SELECT count(*) FROM "'+table+'"').fetchone()[0]
            if table not in ('app_record','app_importbatch','app_importrow','app_auditevent'):assert count==info['count'],table
        assert c.execute('PRAGMA integrity_check').fetchall()==[('ok',)] and c.execute('PRAGMA foreign_key_check').fetchall()==[]
    if args.mode=='rehearse':assert sha(main_db)==before['source_database_sha']
    proof=dict(success=True,mode=args.mode,records=Record.objects.count(),new_rows=total,old_facts_preserved=len(old),old_schemas_preserved=len(before['schemas']),original_files_preserved=len(before['originals']),
               workbook_sha256=sha(BOOK),normal_xlsx_import=True,all_excel_fields_exact=True,idempotent_reimport=True,first_piece_and_launch_unchanged=True,published_metric_unchanged=True,old_metadata_rows_preserved=True,
               template=template,batch_id=str(batch.pk),browser_acceptance=False)
    (ROOT/'data'/('curing_'+('rehearsal' if args.mode=='rehearse' else 'release')+'.json')).write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n')
    (ROOT/'data/curing_example.json').write_text(json.dumps(d,ensure_ascii=False)+'\n');print(json.dumps(proof,ensure_ascii=False))
if __name__=='__main__':main()
