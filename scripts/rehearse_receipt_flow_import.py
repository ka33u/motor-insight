"""Typed-cell comparison and repeated import in a disposable SQLite copy."""
import os,sys,json,sqlite3,tempfile,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));workspace=Path(tempfile.mkdtemp(prefix='motor-receipt-flow-import-'))
with sqlite3.connect(ROOT/'data/platform.sqlite3') as source,sqlite3.connect(workspace/'platform.sqlite3') as target:source.backup(target)
os.environ['MOTOR_SQLITE_PATH']=str(workspace/'platform.sqlite3');os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from django.test import override_settings
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert,stage_file,commit_batch
from app import receipt_flow,supply,analytics,purchase_commitments
file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/26_来料流程_模拟.xlsx';expected=json.loads((ROOT/'data/receipt_flow_scenario.json').read_text());parsed={};cells=dates=0
book=load_workbook(file,read_only=True,data_only=False)
for ds in expected['tables']:
 s=SCHEMAS[ds];rows=book[s['label']].iter_rows(values_only=True);assert list(next(rows))==[f['label'] for f in s['fields']];parsed[ds]=[]
 for row in rows:
  parsed[ds].append({f['name']:convert(v,f) for f,v in zip(s['fields'],row)});dates+=sum(f['type'] in ['date','datetime'] and v is not None and not isinstance(v,str) for f,v in zip(s['fields'],row))
 assert parsed[ds]==expected['tables'][ds];cells+=len(parsed[ds])*len(s['fields'])
book.close();before=Record.objects.count();old=supply.SupplyData(analytics.tables());baseline={'po':[supply.clean(r) for r in old.po_rows],'lots':[supply.clean(r) for r in old.lots],'commitments':purchase_commitments.summary(purchase_commitments.Commitments().rows)}
with override_settings(BASE_DIR=workspace):
 batch,repeated=stage_file(file);assert not repeated and batch.summary['valid']==502 and not batch.summary.get('invalid') and not batch.summary.get('conflict'),batch.summary
 commit_batch(batch.pk);assert Record.objects.count()==before+502
 again,repeated=stage_file(file);assert repeated;commit_batch(again.pk);assert Record.objects.count()==before+502
 for ds,rows in parsed.items():assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values',flat=True))==sorted(rows,key=lambda r:r['id'])
 d=receipt_flow.ReceiptFlow();s=receipt_flow.summary(d.rows);assert (s['objects'],s['completed'],s['disposition'],s['attention'],s['work_attention'])==(133,125,8,0,36)
 assert [x['n'] for x in s['decompositions']]==[111,111];assert {j['id'] for r in d.rows for j in r['jobs'] if j['issues']}==set(expected['cases']);assert not d.global_issues
 now=supply.SupplyData(analytics.tables());assert baseline=={'po':[supply.clean(r) for r in now.po_rows],'lots':[supply.clean(r) for r in now.lots],'commitments':purchase_commitments.summary(purchase_commitments.Commitments().rows)}
report={'workbook_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':502,'cells':cells,'native_date_cells':dates,'tables':{k:len(v) for k,v in parsed.items()},'isolated_import':True,'idempotent':True,'original_purchase_stock_commitments_unchanged':True,'records_before':before,'records_after':before+502,'summary':s,'cases':expected['cases']}
(ROOT/'data/receipt_flow_import_rehearsal.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in report.items() if k!='summary'},ensure_ascii=False))
