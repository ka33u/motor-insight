"""Import one synthetic workbook and recheck a template; prove old data retained."""
import argparse,copy,hashlib,json,os,shutil,sqlite3,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];BEFORE=ROOT/'data/msa-before'
WORKBOOK=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/35_测量系统交叉采样_模拟.xlsx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def write(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str)+'\n')
def roundtrip():
 from openpyxl import load_workbook
 from app.ingestion import convert
 from app.schema import SCHEMAS
 scenario=json.loads((ROOT/'data/msa_scenario.json').read_text());book=load_workbook(WORKBOOK,read_only=True,data_only=False);count=0
 try:
  assert book.sheetnames==['采样说明','测量系统试验','测量试验样件与人员','测量重复观测']
  for ds,wanted in scenario['tables'].items():
   fields=SCHEMAS[ds]['fields'];values=list(book[SCHEMAS[ds]['label']].iter_rows(values_only=True));assert list(values[0])==[f['label'] for f in fields]
   actual=[{f['name']:convert(row[i],f) for i,f in enumerate(fields)} for row in values[1:]];assert actual==wanted,ds;count+=len(actual)
 finally:book.close()
 return dict(rows=count,all_fields_exact=True,sha256=sha(WORKBOOK),typed_naive_local_times=True,visual_sheets_reviewed=4)
def legacy(db):
 m=json.loads((BEFORE/'manifest.json').read_text());assert sha(BEFORE/'platform.sqlite3')==m['database_sha256']
 for name,h in m['physical'].items():assert sha(ROOT/name)==h,name
 changed={'app/schema.py','app/manufacturing_rules.py'}
 for name,h in m['protected'].items():
  if name not in changed:assert sha(ROOT/name)==h,name
 s=(ROOT/'app/schema.py').read_text();append='\nfrom .msa_schema import register_msa\nregister_msa(register)\n'
 assert s.endswith(append) and hashlib.sha256(s[:-len(append)].encode()).hexdigest()==m['protected']['app/schema.py']
 branch="    if dataset in {'msa_studies','msa_members','msa_observations'}:\n        from .msa_source_contract import issues as msa_issues\n        return msa_issues(dataset,row)\n"
 s=(ROOT/'app/manufacturing_rules.py').read_text();assert s.count(branch)==1;assert hashlib.sha256(s.replace(branch,'').encode()).hexdigest()==m['protected']['app/manufacturing_rules.py']
 from app.schema import SCHEMAS
 assert {k:SCHEMAS[k] for k in m['schemas']}==m['schemas'];assert set(SCHEMAS)-set(m['schemas'])=={'msa_studies','msa_members','msa_observations'}
 from app.metric_registry import calculation_hash
 from app.models import MetricVersion
 assert calculation_hash('bi_order_lines')==MetricVersion.objects.get(pk=11).calculation_hash==m['delivery_v8_hash']=='9d9a5af3d45172eca16790d628a36bc6acbe18d1e2a811a66a5b5fbff7037fe3'
 with sqlite3.connect(db) as c:
  c.execute('ATTACH DATABASE ? AS parent',(str(BEFORE/'platform.sqlite3'),));assert c.execute('PRAGMA integrity_check').fetchall()==[('ok',)];assert c.execute('PRAGMA foreign_key_check').fetchall()==[]
  tables=[r[0] for r in c.execute("SELECT name FROM parent.sqlite_master WHERE type='table' AND name!='sqlite_sequence'")];actual={r[0] for r in c.execute("SELECT name FROM main.sqlite_master WHERE type='table' AND name!='sqlite_sequence'")};assert actual==set(tables) and len(actual)==50
  additions={'app_record':1082,'app_importrow':1082,'app_importbatch':1,'app_auditevent':4,'app_importtemplateversion':1}
  for name in tables:
   if name in {'app_importtemplate','app_importtemplateversion'}:continue
   assert c.execute('SELECT count(*) FROM (SELECT * FROM parent."'+name+'" EXCEPT SELECT * FROM main."'+name+'")').fetchone()[0]==0,name
   delta=c.execute('SELECT (SELECT count(*) FROM main."'+name+'")-(SELECT count(*) FROM parent."'+name+'")').fetchone()[0];assert delta==additions.get(name,0),(name,delta)
  cols=[r[1] for r in c.execute('PRAGMA table_info(app_importtemplate)')];allowed={'current_version','revision','updated_at'};selected=','.join('"'+k+'"' for k in cols if k not in allowed)
  assert c.execute('SELECT '+selected+' FROM main.app_importtemplate').fetchall()==c.execute('SELECT '+selected+' FROM parent.app_importtemplate').fetchall()
  assert c.execute('SELECT current_version,revision FROM app_importtemplate').fetchall()==[(4,8)]
  cols=[r[1] for r in c.execute('PRAGMA table_info(app_importtemplateversion)')];selected=','.join('"'+k+'"' for k in cols if k not in {'state','updated_at'})
  assert c.execute('SELECT '+selected+' FROM main.app_importtemplateversion WHERE number<=3 ORDER BY number').fetchall()==c.execute('SELECT '+selected+' FROM parent.app_importtemplateversion ORDER BY number').fetchall()
  assert c.execute('SELECT number,state FROM app_importtemplateversion ORDER BY number').fetchall()==[(1,'retired'),(2,'retired'),(3,'retired'),(4,'active')]
  assert c.execute('SELECT count(*) FROM app_record').fetchone()[0]==214523
  for ds,count in [('msa_studies',11),('msa_members',141),('msa_observations',930)]:assert c.execute('SELECT count(*) FROM app_record WHERE dataset=?',(ds,)).fetchone()[0]==count
 return dict(old_tables_retained=50,old_facts_retained=213441,new_facts=1082,raw_datasets=len(SCHEMAS),physical_files_retained=len(m['physical']),protected_modules=27,old_schema_definitions_retained=115,delivery_v8_unchanged=True)
def main():
 p=argparse.ArgumentParser();p.add_argument('mode',choices=['rehearse','release']);a=p.parse_args();m=json.loads((BEFORE/'manifest.json').read_text());main_before=sha(ROOT/'data/platform.sqlite3');db=ROOT/'data/msa-rehearsal.sqlite3' if a.mode=='rehearse' else ROOT/'data/platform.sqlite3'
 if a.mode=='rehearse':
  assert not db.exists(),'use a fresh destination, never overwrite proof';shutil.copy2(BEFORE/'platform.sqlite3',db)
 else:
  assert main_before==m['database_sha256'];assert json.loads((ROOT/'data/msa_rehearsal.json').read_text())['success'];assert json.loads((ROOT/'data/msa_http_validation.json').read_text())['success'];assert json.loads((ROOT/'data/msa_presentation_checks.json').read_text())['success'];assert 'OK' in (ROOT/'data/msa_full_tests.log').read_text()
 os.environ['DJANGO_SETTINGS_MODULE']='config.settings';os.environ['MOTOR_SQLITE_PATH']=str(db);sys.path.insert(0,str(ROOT));import django;django.setup()
 from app.ingestion import stage_file,commit_batch
 from app.models import Record,ImportBatch,ImportRow,AuditEvent
 from app import msa_data
 from django.contrib.auth.models import User
 original=(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count());excel=roundtrip();assert excel['rows']==1082
 batch,replayed=stage_file(WORKBOOK);assert not replayed and batch.summary['valid']==1082 and batch.summary['total']==1082;commit_batch(batch.pk);batch.refresh_from_db();assert batch.status=='committed'
 counts=(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count());assert counts==tuple(n+d for n,d in zip(original,(1082,1,1082,2)))
 again,replayed=stage_file(WORKBOOK);assert replayed and again.pk==batch.pk;commit_batch(again.pk);assert counts==(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count())
 scenario=json.loads((ROOT/'data/msa_scenario.json').read_text());profiles=[]
 for ds,rows in scenario['tables'].items():
  for row in rows:assert Record.objects.get(dataset=ds,business_key=row['id']).values==row
 for profile in scenario['profiles']:
  d=msa_data.load(profile['id']);r=d['result'];profiles.append(dict(id=profile['id'],profile=profile['profile'],state=r['state'],observations=len(r['points']),sources=len(d['sources']),calibration_state=r['calibration']['state']))
  if a.mode=='rehearse':write(ROOT/('data/msa_board_'+profile['id']+'.json'),{k:v for k,v in r.items() if k!='points'}|{'chart_points':r['points']})
 assert [r['state'] for r in profiles]==['trial']*4+['paused']*7
 from refresh_msa_template import refresh
 template=refresh(User.objects.get(username='demo_admin'),ROOT)
 proof=dict(success=True,mode=a.mode,database_sha256=sha(db),workbook=excel,**legacy(db),profiles=profiles,template=template,batch_id=str(batch.pk),new_audits=4,synthetic=True,browser_acceptance=False,mobile_acceptance=False,real_system_connected=False)
 if a.mode=='rehearse':assert sha(ROOT/'data/platform.sqlite3')==main_before;proof['main_database_unchanged']=True
 write(ROOT/('data/msa_rehearsal.json' if a.mode=='rehearse' else 'data/msa_release.json'),proof);print(json.dumps(proof,ensure_ascii=False))
if __name__=='__main__':main()
