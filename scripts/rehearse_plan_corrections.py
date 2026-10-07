"""Read actual XLSX and rehearse versioned replacement in an isolated baseline."""
import os,sys,json,sqlite3,tempfile,zipfile,hashlib,runpy
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT),str(ROOT/'scripts')]
workspace=Path(tempfile.mkdtemp(prefix='motor-plan-review-'))
with zipfile.ZipFile(ROOT/'data/backups/motor-backup-20261003-234516.zip') as z:(workspace/'platform.sqlite3').write_bytes(z.read('data/platform.sqlite3'))
os.environ['MOTOR_SQLITE_PATH']=str(workspace/'platform.sqlite3');os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django
django.setup()
from django.test import override_settings
from openpyxl import load_workbook
from app.models import Record,ImportBatch,ImportDecision
from app.schema import SCHEMAS
from app.ingestion import convert,stage_file,commit_batch
from app.import_review import decide,snapshot
from app.manufacturing_rules import issues

file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/17_计划日期更正_模拟.xlsx'
expected=json.loads((ROOT/'data/plan_correction_scenario.json').read_text())['tables']['work_orders']
book=load_workbook(file,read_only=True,data_only=False);sheet=book['生产工单'];it=sheet.iter_rows(values_only=True);fields=SCHEMAS['work_orders']['fields']
assert list(next(it))==[f['label'] for f in fields]
parsed=[{f['name']:convert(v,f) for f,v in zip(fields,row)} for row in it];book.close();assert parsed==expected
before={r.business_key:snapshot(r) for r in Record.objects.filter(dataset='work_orders').select_related('source_row__batch')};count=Record.objects.count();decision_count=ImportDecision.objects.count()
old_file=Path(ImportBatch.objects.get(pk=before[expected[0]['id']]['source']['batch']).file_path)
with override_settings(BASE_DIR=workspace):
    batch,repeated=stage_file(file);assert not repeated and batch.summary=={'conflict':6,'total':6,'unknown_sheets':[]},batch.summary
    commit_batch(batch.pk);assert Record.objects.count()==count
    for row in batch.rows.order_by('row_number'):
        current=Record.objects.select_related('source_row__batch').get(dataset='work_orders',business_key=row.business_key);old=snapshot(current);assert old==before[row.business_key]
        token={k:old[k] for k in ['record_id','revision','hash']};token['candidate_hash']=row.record_hash
        decision,again=decide(batch.pk,row.pk,'isolated-rehearsal',{'action':'replace','reason':'按模拟三天窗口修正日期，保留原完工及承诺','expected':token})
        assert not again and decision.before==old and decision.after['revision']==old['revision']+1
    assert Record.objects.count()==count and ImportDecision.objects.count()==decision_count+6
    repeat,is_repeat=stage_file(file);assert is_repeat and repeat.pk==batch.pk;commit_batch(repeat.pk)
    assert Record.objects.count()==count and ImportDecision.objects.count()==decision_count+6
    for rec in Record.objects.filter(dataset='work_orders').select_related('source_row__batch'):
        old=before[rec.business_key]
        if rec.business_key in {r['id'] for r in expected}:
            assert rec.values==next(r for r in expected if r['id']==rec.business_key)
            assert [k for k,v in rec.values.items() if old['values'][k]!=v]==['planned_start']
        else:assert snapshot(rec)==old
    invalid,repeated=stage_file(old_file,force_recheck=True)
    bad={r.business_key for r in invalid.rows.filter(status='invalid') if any(x.get('code')=='WO_PLAN_ORDER' for x in r.issues)}
    assert bad=={r['id'] for r in expected} and invalid.summary['invalid']==6,invalid.summary
    commit_batch(invalid.pk)
    assert all(not issues('work_orders',r.values) for r in Record.objects.filter(dataset='work_orders'))

# Run generator in memory, without executing its file-writing main block.
generated=runpy.run_path(str(ROOT/'scripts/generate_scenario.py'),run_name='generator_rehearsal')
generated['validate']()
assert all(not issues('work_orders',r) for r in generated['DATA']['work_orders'])
report={'synthetic':True,'workbook':str(file),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':6,'cells':6*len(fields),'stage_conflicts':6,'approved_in_isolated_copy':6,'row_count_unchanged':count,'legacy_bad_rows_rejected':sorted(bad),'replay_idempotent':True,'generator_plans_checked':len(generated['DATA']['work_orders']),'production_database_written':False}
(ROOT/'data/plan_correction_rehearsal.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False,indent=2))
