"""Read every saved XLSX cell and rehearse atomic, idempotent import in a copy."""
import os,sys,json,sqlite3,tempfile,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));workspace=Path(tempfile.mkdtemp(prefix='motor-payables-import-'))
with sqlite3.connect(ROOT/'data/platform.sqlite3') as source,sqlite3.connect(workspace/'platform.sqlite3') as target:source.backup(target)
os.environ['MOTOR_SQLITE_PATH']=str(workspace/'platform.sqlite3');os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from django.test import override_settings
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert,stage_file,commit_batch
from app import analytics,payables
file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/19_供应商应付_模拟.xlsx';expected=json.loads((ROOT/'data/payables_scenario.json').read_text());parsed={};cells=0;dates=0
book=load_workbook(file,read_only=True,data_only=False)
for ds in expected['tables']:
    schema=SCHEMAS[ds];rows=book[schema['label']].iter_rows(values_only=True);assert list(next(rows))==[f['label'] for f in schema['fields']]
    parsed[ds]=[]
    for row in rows:
        parsed[ds].append({f['name']:convert(v,f) for f,v in zip(schema['fields'],row)})
        dates+=sum(f['type']=='date' and v is not None and not isinstance(v,str) for f,v in zip(schema['fields'],row))
    assert parsed[ds]==expected['tables'][ds],ds;cells+=len(parsed[ds])*len(schema['fields'])
book.close();before=Record.objects.count();added=sum(len(v) for v in parsed.values())
with override_settings(BASE_DIR=workspace):
    batch,repeated=stage_file(file);assert not repeated;assert batch.summary.get('valid')==added and not batch.summary.get('invalid') and not batch.summary.get('conflict'),batch.summary
    commit_batch(batch.pk);assert Record.objects.count()==before+added
    repeat,is_repeat=stage_file(file);assert is_repeat and repeat.pk==batch.pk;commit_batch(repeat.pk);assert Record.objects.count()==before+added
    for ds,rows in parsed.items():assert list(Record.objects.filter(dataset=ds,business_key__in=[r['id'] for r in rows]).order_by('business_key').values_list('values',flat=True))==sorted(rows,key=lambda r:r['id'])
    d=payables.current();assert not d.global_issues
    for index in [d.invoices,d.payments,d.plans,d.receipts]:assert not any(r['issues'] for r in index.values())
    summaries={tab:payables.summary(list(index.values()),tab) for tab,index in [('invoices',d.invoices),('payments',d.payments),('plans',d.plans),('receipts',d.receipts)]}
report={'workbook':str(file),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':added,'cells':cells,'native_date_cells':dates,'tables':{k:len(v) for k,v in parsed.items()},'isolated_import':True,'idempotent':True,'records_before':before,'records_after':before+added,'summaries':summaries}
(ROOT/'data/payables_import_rehearsal.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
