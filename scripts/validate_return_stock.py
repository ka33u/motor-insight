"""Check normally imported source cells and independently reconcile return/lot state examples."""
import hashlib,json,os,sys
from datetime import date,datetime
from decimal import Decimal
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert
from app import analytics,supply,return_stock

BOOK='47_退料库存状态_模拟.xlsx'


def verify():
    path=ROOT/'demo/workbooks'/BOOK;workbook=load_workbook(path,read_only=True,data_only=False);schemas={s['label']:s for s in SCHEMAS.values()};expected={};file_hash=hashlib.sha256(path.read_bytes()).hexdigest()
    try:
        assert len(workbook.sheetnames)==10
        for sheet in workbook:
            if sheet.title=='导入说明':continue
            schema=schemas[sheet.title];rows=list(sheet.iter_rows(values_only=True));assert list(rows[0])==[f['label'] for f in schema['fields']]
            for number,row in enumerate(rows[1:],2):
                for value,field in zip(row,schema['fields']):
                    if value is None:continue
                    if field['type'] in ('int','float'):assert type(value) in (int,float)
                    elif field['type'] in ('date','datetime'):assert isinstance(value,(date,datetime))
                    elif field['type']=='bool':assert type(value) is bool
                    else:assert isinstance(value,str)
                values={f['name']:convert(v,f) for v,f in zip(row,schema['fields'])};identity=schema['key'],values['id'];assert identity not in expected;expected[identity]=values
                record=Record.objects.select_related('source_row__batch').get(dataset=identity[0],business_key=identity[1]);source=record.source_row
                assert record.values==values==source.normalized
                assert (source.batch.filename,source.sheet,source.row_number)==(BOOK,sheet.title,number)
                assert source.batch.file_hash==file_hash
    finally:workbook.close()
    assert len(expected)==135
    facts=list(Record.objects.order_by('dataset','business_key').values_list('dataset','business_key','record_hash'))
    old=[r for r in facts if (r[0],r[1]) not in expected];digest=hashlib.sha256()
    for row in old:digest.update((json.dumps(row,ensure_ascii=False)+'\n').encode())
    assert len(facts)==234037 and len(old)==233902 and digest.hexdigest()=='ac6a3825a117a6d9a0144e91fd2c654ffc6d98d326ae339e46f161ebf1f523d2'
    raw=analytics._tables(analytics.revision());d=return_stock.ReturnStock(supply.SupplyData(raw))
    rows=d.cohort(return_stock.filters({'q':'MO-260928-Q'}));assert len(rows)==8
    expected_states=['待检','可用','隔离',None,'可用','待检','可用'];profiles=[]
    for n,state in enumerate(expected_states,1):
        chosen=[r for r in rows if r['work_order_id']==f'MO-260928-Q{n:03d}'];assert len(chosen)==(2 if n==7 else 1)
        for r in chosen:
            assert r['document_verified'] and r['current_state']==state
            assert r['arrival_state']==(None if n==5 else '待检')
        mid,lot,loc=chosen[0]['material_id'],chosen[0]['lot'],chosen[0]['location']
        openings=[r for r in raw['inventory_opening'] if (r['material_id'],r['lot'],r['location'])==(mid,lot,loc)]
        movements=[r for r in raw['inventory_movements'] if (r['material_id'],r['lot'],r['location'])==(mid,lot,loc) and r['occurred']<=analytics.AS_OF]
        balance=sum((Decimal(str(r['qty'])) for r in openings),Decimal(0))+sum((Decimal(str(r['qty_signed'])) for r in movements),Decimal(0))
        view=d.unique_lots(chosen)[0];assert Decimal(str(view['balance_qty']))==balance
        available=None if n==4 else balance if state=='可用' else Decimal(0)
        assert view['usable_state_qty']==available
        profiles.append(dict(work_order=chosen[0]['work_order_id'],returns=len(chosen),returned=sum(r['qty_signed'] for r in chosen),balance=float(balance),usable=None if available is None else float(available),arrival=chosen[0]['arrival_state'],current=state))
    u=d.units(rows)[0];assert (u['returned_qty'],u['target_lots'],u['target_balance_qty'],u['target_usable_qty'],u['unknown_lots'])==(22,7,21,None,1)
    normal=d.cohort(return_stock.filters({'material_id':'01.01.9101'}));u=d.units(normal)[0]
    assert (u['returned_qty'],u['target_balance_qty'],u['target_usable_qty'],u['target_lots'])==(20,19,9,6)
    matrix={r['label']:(r['returns'],r['document_verified']) for r in d.matrix(rows)}
    assert matrix=={'可用':(4,4),'待检':(2,2),'隔离':(1,1),'当前证据待核对':(1,1)}
    assert d.summary(rows)==dict(returns=8,target_lots=7,documents_verified=8,current_verified=7,arrival_known=7,attention=2)
    # Preserve the prior return scenario's material balances while adding isolated inputs.
    old_stock=d.stock.stock_index['01.01.9001'];assert (old_stock['balance_qty'],old_stock['usable_state_qty'])==(928.5,923.5)
    proof=dict(success=True,source_rows=135,sheets=10,records=234037,protected_old_facts=233902,source_cells_and_types_exact=True,
        workbook_sha256=file_hash,independent_target_lot_checks=7,profiles=profiles,return_rows=8,deduplicated_target_lots=7,
        returned_qty=22,target_balance_qty=21,target_usable_qty=None,known_material_returned=20,known_material_balance=19,known_material_usable=9,
        old_return_material_balance_preserved=True,browser_acceptance=False)
    (ROOT/'data/return_stock_validation.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n');return proof


if __name__=='__main__':print(json.dumps(verify(),ensure_ascii=False))
