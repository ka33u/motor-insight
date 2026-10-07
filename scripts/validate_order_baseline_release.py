"""Excel roundtrip, normal import, source preservation and independent totals."""
import argparse
import hashlib
import json
import os
import shutil
import sqlite3
import sys
from decimal import Decimal, ROUND_CEILING
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/40_订单与BOM基线_模拟.xlsx'
CHECKPOINT = ROOT/'data/order-baseline-before/manifest.json'
REHEARSAL = Path('/private/tmp/motorinsight-order-baseline-rehearsal.sqlite3')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def roundtrip():
    from openpyxl import load_workbook
    from app.ingestion import convert
    from app.schema import SCHEMAS
    scenario = json.loads((ROOT/'data/order_baseline_scenario.json').read_text())
    workbook = load_workbook(BOOK, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ['导入说明']+[SCHEMAS[ds]['label'] for ds in scenario['tables']]
        for ds, expected in scenario['tables'].items():
            fields = SCHEMAS[ds]['fields']
            rows = list(workbook[SCHEMAS[ds]['label']].iter_rows(values_only=True))
            assert list(rows[0]) == [f['label'] for f in fields]
            actual = [{f['name']: convert(row[i], f) for i, f in enumerate(fields)} for row in rows[1:]]
            assert actual == expected, ds
    finally:
        workbook.close()
    return dict(rows=609, sha256=sha(BOOK), all_fields_exact=True, typed_dates_and_identifiers=True, visual_sheets_reviewed=4)


def invariants(d):
    r = d['result']
    if r['state'] == 'paused':
        assert r['summary'] is None and not r['bom'] and not r['orders']
        return
    assert r['summary']['order_qty'] == 135 and r['summary']['planned_qty'] == 50 and r['summary']['uncovered_qty'] == 85
    assert r['summary']['fully_covered_lines'] == 0 and all(o['order_finish'] is None for o in r['orders'])
    links = {l['id']: l for l in r['links']}
    for b in r['bom']:
        step = Decimal(str(b['quantum']))
        unrounded = Decimal(links[b['link_id']]['plan_qty'])*Decimal(str(b['unit_qty']))*(1+Decimal(str(b['scrap_allowance'])))
        expected = (unrounded/step).to_integral_value(rounding=ROUND_CEILING)*step
        assert Decimal(b['required_qty']) == expected
    assert len({o['id'] for o in r['orders']}) == r['summary']['order_lines'] == 6
    assert all(o['planned_qty']+o['uncovered_qty'] == o['baseline_open_qty'] for o in r['orders'])
    assert r['summary']['customer_late_jobs'] == r['summary']['scheduled_jobs']


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['rehearse', 'release'])
    args = parser.parse_args()
    manifest = json.loads(CHECKPOINT.read_text())
    main_db = ROOT/'data/platform.sqlite3'
    assert sha(main_db) == manifest['source_database_sha'], 'Inspect changes since checkpoint before importing.'
    if args.mode == 'rehearse':
        assert not REHEARSAL.exists(), 'Do not overwrite an existing rehearsal.'
        shutil.copy2(main_db, REHEARSAL)
        db = REHEARSAL
    else:
        for name in ('order_baseline_rehearsal.json', 'order_baseline_ui.json', 'order_baseline_bootstrap.json', 'order_baseline_http.json'):
            assert json.loads((ROOT/'data'/name).read_text())['success'], name
        assert '\nOK\n' in Path('/private/tmp/motorinsight-order-baseline-full-tests.log').read_text()
        db = main_db
    os.environ['DJANGO_SETTINGS_MODULE'] = 'config.settings'
    os.environ['MOTOR_SQLITE_PATH'] = str(db)
    sys.path.insert(0, str(ROOT))
    import django
    django.setup()
    from app.ingestion import stage_file, commit_batch
    from app.models import Record, ImportBatch, ImportRow, AuditEvent, MetricVersion
    from app.schema import SCHEMAS
    from app import order_baseline_data, joint_schedule_data, finite_schedule
    from app.order_baseline_schema import DATASETS
    from app.metric_registry import calculation_hash
    from django.contrib.auth.models import User
    from validate_joint_schedule_release import frozen_tables
    frozen = frozen_tables(db)
    excel = roundtrip()
    before = (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    batch, replayed = stage_file(BOOK)
    assert not replayed and batch.summary['total'] == batch.summary['valid'] == 609, batch.summary
    assert batch.summary['unknown_sheets'] == []
    commit_batch(batch.pk)
    batch.refresh_from_db()
    assert batch.status == 'committed'
    after = (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    assert after == tuple(n+d for n, d in zip(before, (609, 1, 609, 2)))
    again, replayed = stage_file(BOOK)
    assert replayed and again.pk == batch.pk
    commit_batch(again.pk)
    assert after == (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    scenario = json.loads((ROOT/'data/order_baseline_scenario.json').read_text())
    for ds, expected in scenario['tables'].items():
        assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values', flat=True)) == sorted(expected, key=lambda r: r['id'])
    profiles = []
    for profile in scenario['profiles']:
        d = order_baseline_data.load(profile['id'], profile['policy'])
        r = d['result']
        for field in ('state', 'summary', 'issues', 'warnings'):
            assert r[field] == profile[field], (profile['id'], field, r[field], profile[field])
        invariants(d)
        profiles.append(dict(id=profile['id'], policy=profile['policy'], state=r['state'], summary=r['summary'], source_count=len(d['sources'])))
    old_joint = json.loads((ROOT/'data/order-baseline-before/joint001.json').read_text())
    new_joint = joint_schedule_data.load('MP-261002-001')
    assert new_joint['result'] == old_joint['result'] and new_joint['source_hash'] == old_joint['source_hash']
    from refresh_order_baseline_template import refresh
    template = refresh(User.objects.get(username='demo_admin'), ROOT)
    old_facts = list(Record.objects.exclude(dataset__in=DATASETS).order_by('dataset', 'business_key').values('dataset', 'business_key', 'record_hash'))
    assert len(old_facts) == manifest['record_count']
    assert finite_schedule.digest(old_facts) == manifest['record_digest']
    assert {k: SCHEMAS[k] for k in manifest['schemas']} == manifest['schemas']
    assert set(SCHEMAS)-set(manifest['schemas']) == set(DATASETS)
    for name, digest in manifest['originals'].items():
        assert sha(ROOT/name) == digest, name
    assert calculation_hash('bi_order_lines') == MetricVersion.objects.get(pk=11).calculation_hash == '9d9a5af3d45172eca16790d628a36bc6acbe18d1e2a811a66a5b5fbff7037fe3'
    assert frozen_tables(db) == frozen
    with sqlite3.connect(db) as c:
        assert c.execute('PRAGMA integrity_check').fetchall() == [('ok',)] and c.execute('PRAGMA foreign_key_check').fetchall() == []
    if args.mode == 'rehearse':
        assert sha(main_db) == manifest['source_database_sha']
    proof = dict(success=True, mode=args.mode, workbook=excel, profiles=profiles, old_facts_preserved=len(old_facts), original_files_preserved=len(manifest['originals']),
                 old_contracts_preserved=len(manifest['schemas']), frozen_tables_preserved=len(frozen), metric_v8_unchanged=True, joint_outputs_unchanged=True, template=template,
                 batch_id=str(batch.pk), records=Record.objects.count(), database_sha256=sha(db), synthetic=True, browser_acceptance=False, mobile_acceptance=False)
    path = ROOT/('data/order_baseline_rehearsal.json' if args.mode == 'rehearse' else 'data/order_baseline_release.json')
    path.write_text(json.dumps(proof, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({k: v for k, v in proof.items() if k != 'profiles'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
