"""Read-only XLSX, independent inspection counts, browser exports and preservation QA."""
import os,sys,json,csv,hashlib,sqlite3,zipfile,tempfile
from pathlib import Path
from collections import Counter,defaultdict
from statistics import median
from datetime import datetime
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django
django.setup()
from openpyxl import load_workbook
from django.contrib.auth.models import User
from app.models import Record,ImportBatch,IssueDisposition,AuditEvent,DataQualityScan,MetricVersion,AnalysisModel,Topic
from app.schema import SCHEMAS
from app.ingestion import convert,fingerprint
from app import process_quality as pq,data_health,metric_registry,analytics

baseline=ROOT/'data/backups/motor-backup-20261004-010831.zip'
delta={'app_record':2590,'app_importrow':2590,'app_importbatch':1,'app_issuedisposition':1,'app_dataqualityscan':1}
preserved={};source_count=0
with zipfile.ZipFile(baseline) as z,tempfile.TemporaryDirectory() as tmp:
    for entry in json.loads(z.read('manifest.json'))['files']:
        if entry['path'].startswith(('data/imports/','data/device_files/')):
            assert hashlib.sha256((ROOT/entry['path']).read_bytes()).hexdigest()==entry['sha256'];source_count+=1
    old_path=Path(tmp)/'before.sqlite3';old_path.write_bytes(z.read('data/platform.sqlite3'))
    with sqlite3.connect(old_path) as old,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
        tables=[r[0] for r in old.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'app_%'")]
        assert len(tables)==25
        for table in tables:
            before=list(old.execute(f'SELECT * FROM {table} ORDER BY id'))
            current=list(now.execute(f'SELECT * FROM {table} ORDER BY id'));byid={r[0]:r for r in current}
            assert all(byid.get(r[0])==r for r in before),(table,'prior row changed')
            if table!='app_auditevent':assert len(current)-len(before)==delta.get(table,0),(table,len(before),len(current))
            preserved[table]={'before':len(before),'after':len(current),'prior_rows_unchanged':True}

path=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/18_工序检验_模拟.xlsx'
file_hash=hashlib.sha256(path.read_bytes()).hexdigest()
assert file_hash==json.loads((ROOT/'data/process_quality_import_rehearsal.json').read_text())['sha256']
batch=ImportBatch.objects.get(file_hash=file_hash,status='committed')
assert Path(batch.file_path).read_bytes()==path.read_bytes() and batch.summary['committed']==2590
parsed={};cells=0;book=load_workbook(path,read_only=True,data_only=False)
for ds in pq.NEW_TABLES:
    fields=SCHEMAS[ds]['fields'];rows=book[SCHEMAS[ds]['label']].iter_rows();assert [c.value for c in next(rows)]==[f['label'] for f in fields]
    parsed[ds]=[]
    for n,row in enumerate(rows,2):
        values={}
        for c,f in zip(row,fields):
            assert c.data_type!='f'
            if f['name']=='id':assert isinstance(c.value,str) and c.data_type=='s'
            if f['type'] in ['date','datetime'] and c.value is not None:assert isinstance(c.value,datetime)
            values[f['name']]=convert(c.value,f);cells+=1
        rec=Record.objects.select_related('source_row').get(dataset=ds,business_key=values['id'])
        assert rec.values==rec.source_row.normalized==values and rec.revision==1
        assert rec.source_row.batch_id==batch.pk and rec.source_row.row_number==n and rec.source_row.sheet==SCHEMAS[ds]['label']
        parsed[ds].append(values)
book.close();assert cells==18717
assert {ds:len(v) for ds,v in parsed.items()}=={'process_specs':288,'process_check_plans':576,'process_checks':585,'process_readings':1141}
assert (Record.objects.count(),Record.objects.values('dataset').distinct().count(),ImportBatch.objects.exclude(status='superseded').count())==(110787,63,18)
assert (AnalysisModel.objects.count(),Topic.objects.count())==(48,17)
facts_hash=hashlib.sha256();data=defaultdict(list)
for rec in Record.objects.select_related('source_row').order_by('dataset','business_key').iterator(chunk_size=1000):
    assert rec.values==rec.source_row.normalized and rec.record_hash==rec.source_row.record_hash==fingerprint(rec.values)
    assert (rec.dataset,rec.business_key)==(rec.source_row.dataset,rec.source_row.business_key)
    facts_hash.update(json.dumps((rec.dataset,rec.business_key,rec.values,rec.revision,rec.source_row_id),ensure_ascii=False,sort_keys=True).encode())
    if rec.dataset in pq.TABLES:data[rec.dataset].append(rec.values)

# This calculation reads the workbook facts independently of ProcessQuality.
idx={ds:{r['id']:r for r in rows} for ds,rows in data.items()};checks=defaultdict(list);readings=defaultdict(list)
for c in parsed['process_checks']:
    if not c['voided'] and c['checked']<=analytics.AS_OF:checks[c['plan_id']].append(c)
for r in parsed['process_readings']:readings[r['check_id']].append(r)
states={};results={};first_pass=0;out_keys=set();missing_keys=set();inspection_count=0;comparable=defaultdict(list)
for p in parsed['process_check_plans']:
    op=idx['operations'][p['operation_id']];wo=idx['work_orders'][op['work_order_id']]
    branch='整机' if op['object_type']=='整机' else idx['batches'][op['object_id']]['kind']
    ordered=sorted(checks[p['id']],key=lambda c:c['checked']);assert len({c['checked'] for c in ordered})==len(ordered)
    results[p['id']]=[]
    for c in ordered:
        specifications={s['id']:s for s in parsed['process_specs'] if (s['product_id'],s['route_version'],s['process'],s['branch'],s['version'])==(wo['product_id'],wo['route_version'],op['process'],branch,p['spec_version']) and s['effective']<=c['checked'][:10] and (not s['expires'] or c['checked'][:10]<s['expires'])}
        assert len(specifications)==2 and op['started']<=c['checked']<=op['finished']
        values=readings[c['id']];assert all(r['spec_id'] in specifications for r in values)
        bad=any(r['unit']!=specifications[r['spec_id']]['unit'] for r in values)
        missing={s['id'] for s in specifications.values() if s['mandatory']}-{r['spec_id'] for r in values}
        outside=[r for r in values if (specifications[r['spec_id']]['lsl'] is not None and r['value']<specifications[r['spec_id']]['lsl']) or (specifications[r['spec_id']]['usl'] is not None and r['value']>specifications[r['spec_id']]['usl'])]
        state='资料待核对' if bad else '超限且漏项' if outside and missing else '有超限' if outside else '有漏项' if missing else '齐项且范围内'
        results[p['id']].append(state)
        if not bad:
            for mode in ['all']+(['first'] if c==ordered[0] else [])+(['latest'] if c==ordered[-1] else []):
                comparable[mode]+=values
    state=results[p['id']][-1] if ordered else '到期待检';states[p['id']]=state
    inspection_count+=bool(ordered);first_pass+=bool(ordered and results[p['id']][0]=='齐项且范围内')
    if state in ['有超限','超限且漏项']:out_keys.add(p['id'])
    if state in ['有漏项','超限且漏项']:missing_keys.add(p['id'])
counts=Counter(states.values());assert counts=={'资料待核对':15,'到期待检':33,'有超限':38,'超限且漏项':2,'齐项且范围内':462,'有漏项':26}
expected={'plans':576,'pending':33,'out':40,'missing':28,'pass':462,'attention':15,'future':0,'inspected':543,'first_pass':438,'unique_operations':288}
engine=pq.ProcessQuality(data,pq.filters({}));assert engine.summary(engine.selected())==expected
assert inspection_count==543 and first_pass==438
assert all(r['state']==states[r['id']] for r in engine.selected())
assert results['GJH26-00023']==['有超限','齐项且范围内']
parameter_summary={}
for mode in ['first','latest','all']:
    rows=[r for r in comparable[mode] if r['spec_id']=='CSGF-00031'];values=[r['value'] for r in rows]
    measured=pq.ProcessQuality(data,pq.filters({'tab':'parameters','mode':mode,'spec':'CSGF-00031'})).parameters()
    assert {r['id'] for r in measured['rows']}=={r['id'] for r in rows}
    assert measured['summary']['median']==median(values) and measured['summary']['samples']==len(values)
    parameter_summary[mode]=measured['summary']

def exported(name):
    path=ROOT/'outputs'/name;rows=list(csv.reader(path.open(encoding='utf-8-sig')))
    return json.loads(rows[0][4]),[dict(zip(rows[2],r)) for r in rows[3:]],hashlib.sha256(path.read_bytes()).hexdigest()
f,rows,out_hash=exported('工序检验_超限计划_浏览器导出.csv')
assert f==pq.filters({'stage':'out'}) and {r['应检计划号'] for r in rows}==out_keys
f,samples,sample_hash=exported('工序参数_区间样本_浏览器导出.csv')
assert f==pq.filters({'tab':'parameters','spec':'CSGF-00031','mode':'all','bin':'0'})
all_values=[r for r in comparable['all'] if r['spec_id']=='CSGF-00031'];low=min(r['value'] for r in all_values);upper=low+(max(r['value'] for r in all_values)-low)/8
assert {r['测量记录号'] for r in samples}=={r['id'] for r in all_values if low<=r['value']<upper} and len(samples)==2
assert all(low<=float(r['实测值'])<upper for r in samples)
follow=IssueDisposition.objects.get(key='process-quality:GJH26-00037')
assert (follow.version,follow.status,follow.owner)==(1,'待工艺核对','质量与工艺岗位（模拟）')
assert AuditEvent.objects.filter(action='process_quality.followup',object_id=follow.key).count()==1
scans=list(DataQualityScan.objects.order_by('created_at'));assert len(scans)==7
latest=scans[-1];assert latest.engine_hash==data_health.engine_hash() and data_health.digest(latest.snapshot)==latest.snapshot_hash
assert not latest.snapshot['issues'] and not latest.snapshot['unknown_datasets']
assert len(latest.snapshot['datasets'])==63 and sum(r['rows'] for r in latest.snapshot['datasets'].values())==110787
assert not data_health.project(latest,User.objects.get(username='demo_admin'))['stale']
assert all(data_health.digest(s.snapshot)==s.snapshot_hash for s in scans)
v=MetricVersion.objects.get(metric__key='DELIVERY_OTIF',version=4)
assert v.status=='published' and v.calculation_hash==metric_registry.calculation_hash('bi_order_lines')=='f6815a706b1f19c2e2be0af931c640d869b45b17bc0aa1f72d7fe2144356868f'
report={'synthetic':True,'baseline':str(baseline.relative_to(ROOT)),'workbook_sha256':file_hash,'batch':str(batch.pk),'workbook_rows':2590,'workbook_cells_checked':cells,'records':110787,'source_categories':63,'active_workbooks':18,'business_sha256':facts_hash.hexdigest(),'preserved_source_files':source_count,'preserved_tables':preserved,'summary':expected,'states':dict(counts),'parameter_comparison':parameter_summary,'browser_exports':{'out_rows':len(rows),'out_sha256':out_hash,'sample_rows':len(samples),'sample_sha256':sample_hash},'coordination_version':1,'scan':str(latest.pk),'basic_issues':0,'metric_v4_unchanged':True,'tests':683,'new_tests':37,'limits':'Synthetic specifications and explicitly registered pilot plans only. Range checks are not first-piece approval, whole-batch acceptance, shipping release, calibration verification or SPC capability.'}
if len(sys.argv)>1:
    backup=Path(sys.argv[1])
    with zipfile.ZipFile(backup) as z,tempfile.TemporaryDirectory() as tmp:
        manifest=json.loads(z.read('manifest.json'))
        for entry in manifest['files']:
            raw=z.read(entry['path']);assert len(raw)==entry['size'] and hashlib.sha256(raw).hexdigest()==entry['sha256']
        restored=Path(tmp)/'restore.sqlite3';restored.write_bytes(z.read('data/platform.sqlite3'))
        with sqlite3.connect(restored) as recovered,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
            assert recovered.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
            for table in tables:assert list(recovered.execute(f'SELECT * FROM {table} ORDER BY id'))==list(now.execute(f'SELECT * FROM {table} ORDER BY id'))
        report['backup']={'path':str(backup),'files':len(manifest['files']),'app_tables':len(tables),'sqlite_integrity':'ok','all_app_tables_match':True}
(ROOT/'data/process_quality_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in report.items() if k!='preserved_tables'},ensure_ascii=False,indent=2))
