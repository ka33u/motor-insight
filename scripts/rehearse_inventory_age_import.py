"""Compare every XLSX field and rehearse import in an isolated database copy."""
import os,sys,json,sqlite3,tempfile,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));workspace=Path(tempfile.mkdtemp(prefix='motor-inventory-age-import-'))
with sqlite3.connect(ROOT/'data/platform.sqlite3') as source,sqlite3.connect(workspace/'platform.sqlite3') as target:source.backup(target)
os.environ['MOTOR_SQLITE_PATH']=str(workspace/'platform.sqlite3');os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from django.test import override_settings
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert,stage_file,commit_batch
from app import inventory_age,supply,analytics
file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/24_库存批次资料_模拟.xlsx';expected=json.loads((ROOT/'data/inventory_age_scenario.json').read_text());parsed={};cells=0;dates=0
book=load_workbook(file,read_only=True,data_only=False)
for ds in expected['tables']:
 s=SCHEMAS[ds];rows=book[s['label']].iter_rows(values_only=True);assert list(next(rows))==[f['label'] for f in s['fields']];parsed[ds]=[]
 for row in rows:
  parsed[ds].append({f['name']:convert(v,f) for f,v in zip(s['fields'],row)});dates+=sum(f['type'] in ['date','datetime'] and v is not None and not isinstance(v,str) for f,v in zip(s['fields'],row))
 assert parsed[ds]==expected['tables'][ds],ds;cells+=len(parsed[ds])*len(s['fields'])
book.close();before=Record.objects.count();added=sum(map(len,parsed.values()));baseline=[supply.clean(r) for r in supply.SupplyData(analytics.tables()).lots]
with override_settings(BASE_DIR=workspace):
 batch,repeated=stage_file(file);assert not repeated;assert batch.summary.get('valid')==added and not batch.summary.get('invalid') and not batch.summary.get('conflict'),batch.summary
 commit_batch(batch.pk);assert Record.objects.count()==before+added
 repeat,is_repeat=stage_file(file);assert is_repeat;commit_batch(repeat.pk);assert Record.objects.count()==before+added
 for ds,rows in parsed.items():assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values',flat=True))==sorted(rows,key=lambda r:r['id'])
 d=inventory_age.InventoryAge();summary=inventory_age.summary(d.rows)
 assert len(d.rows)==173 and summary['overage']>0 and summary['near']>0 and summary['expired']>0 and summary['unknown']>0
 assert baseline==[supply.clean(r) for r in supply.SupplyData(analytics.tables()).lots]
 assert sum(r['date_version']==2 for r in d.rows)==6 and not d.global_issues
report={'workbook':str(file),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':added,'cells':cells,'native_date_cells':dates,'tables':{k:len(v) for k,v in parsed.items()},'isolated_import':True,'idempotent':True,'stock_unchanged':True,'records_before':before,'records_after':before+added,'summary':summary}
(ROOT/'data/inventory_age_import_rehearsal.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
