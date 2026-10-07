"""Read-only typed-cell comparison and disposable importer rehearsal."""
import os,sys,json,sqlite3,tempfile,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));workspace=Path(tempfile.mkdtemp(prefix='motor-certificate-import-'))
with sqlite3.connect(ROOT/'data/platform.sqlite3') as source,sqlite3.connect(workspace/'platform.sqlite3') as target:source.backup(target)
os.environ['MOTOR_SQLITE_PATH']=str(workspace/'platform.sqlite3');os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from django.test import override_settings
from django.contrib.auth.models import User
from openpyxl import load_workbook
from app.models import Record
from app.schema import SCHEMAS
from app.ingestion import convert,stage_file,commit_batch
from app import material_certificates as eng,supply,analytics,purchase_commitments,receipt_flow,incoming_quality
file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/28_材质证明_模拟.xlsx';expected=json.loads((ROOT/'data/material_certificate_scenario.json').read_text());parsed={};cells=dates=0
book=load_workbook(file,read_only=True,data_only=False)
for ds in eng.TABLES:
    s=SCHEMAS[ds];rows=book[s['label']].iter_rows(values_only=True);assert list(next(rows))==[f['label'] for f in s['fields']];parsed[ds]=[]
    for row in rows:
        assert not any(isinstance(v,str) and v.startswith('=') for v in row)
        parsed[ds].append({f['name']:convert(v,f) for f,v in zip(s['fields'],row)});dates+=sum(f['type'] in ['date','datetime'] and v is not None and not isinstance(v,str) for f,v in zip(s['fields'],row))
    assert parsed[ds]==expected['tables'][ds],ds;cells+=len(parsed[ds])*len(s['fields'])
book.close();before=Record.objects.count();assert not Record.objects.filter(dataset__in=eng.TABLES).exists()
def old_results():
    d=supply.SupplyData(analytics.tables());return {'purchase':[supply.clean(r) for r in d.po_rows],'lots':[supply.clean(r) for r in d.lots],'commitments':purchase_commitments.summary(purchase_commitments.Commitments().rows),'flow':receipt_flow.summary(receipt_flow.ReceiptFlow().rows),'characteristics':incoming_quality.summary(incoming_quality.IncomingQuality().rows)}
baseline=old_results();n=sum(map(len,parsed.values()));admin=User.objects.get(username='demo_admin')
with override_settings(BASE_DIR=workspace,DEVICE_FILE_ROOT=ROOT/'data/device_files'):
    batch,repeated=stage_file(file);assert not repeated and batch.summary['valid']==n and not batch.summary.get('invalid') and not batch.summary.get('conflict'),batch.summary
    commit_batch(batch.pk);assert Record.objects.count()==before+n
    again,repeated=stage_file(file);assert repeated;commit_batch(again.pk);assert Record.objects.count()==before+n
    for ds,rows in parsed.items():assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values',flat=True))==sorted(rows,key=lambda r:r['id'])
    data=eng.Certificates(resolve=eng.resolver(admin));summary=eng.summary(data.rows);proof=json.loads((ROOT/'data/material_certificate_scenario_build.json').read_text());assert summary==proof['summary'];assert not data.global_issues
    for key,case in proof['cases'].items():assert data.index[key]['state']==case['state'],key
    assert baseline==old_results()
report={'workbook_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':n,'cells':cells,'native_date_cells':dates,'tables':{k:len(v) for k,v in parsed.items()},'isolated_import':True,'idempotent':True,'records_before':before,'records_after':before+n,'original_results_unchanged':True,'summary':summary}
(ROOT/'data/material_certificate_import_rehearsal.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
