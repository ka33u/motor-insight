"""Normal import, exact old-row preservation, route/resources and receipt trials."""
import argparse,hashlib,json,os,shutil,sqlite3,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];BEFORE=ROOT/'data/finite-schedule-before';BOOK=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/36_有限资源试排_模拟.xlsx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def write(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str)+'\n')
def roundtrip():
 from openpyxl import load_workbook
 from app.ingestion import convert
 from app.schema import SCHEMAS
 wanted=json.loads((ROOT/'data/finite_schedule_scenario.json').read_text());book=load_workbook(BOOK,read_only=True,data_only=False)
 try:
  assert book.sheetnames==['导入说明']+[SCHEMAS[ds]['label'] for ds in wanted['tables']]
  for ds,rows in wanted['tables'].items():
   fields=SCHEMAS[ds]['fields'];actual=list(book[SCHEMAS[ds]['label']].iter_rows(values_only=True));assert list(actual[0])==[f['label'] for f in fields]
   converted=[{f['name']:convert(row[n],f) for n,f in enumerate(fields)} for row in actual[1:]];assert converted==rows,ds
 finally:book.close()
 return dict(rows=6006,sha256=sha(BOOK),all_fields_exact=True,typed_plant_wall_times=True,visual_sheets_reviewed=8)
def preserve(db):
 m=json.loads((BEFORE/'manifest.json').read_text());assert sha(BEFORE/'platform.sqlite3')==m['database_sha256']
 for n,h in m['physical'].items():assert sha(ROOT/n)==h,n
 for n,h in m['protected'].items():
  if n not in {'app/schema.py','app/manufacturing_rules.py'}:assert sha(ROOT/n)==h,n
 s=(ROOT/'app/schema.py').read_text();suffix='\nfrom .finite_schedule_schema import register_schedule\nregister_schedule(register)\n';assert s.endswith(suffix);assert sha(BEFORE/'app/schema.py')==hashlib.sha256(s[:-len(suffix)].encode()).hexdigest()
 branch="    if dataset in {'schedule_studies','schedule_jobs','schedule_tasks','schedule_edges','schedule_options','schedule_windows','schedule_blocks'}:\n        from .finite_schedule_contract import issues as schedule_issues\n        return schedule_issues(dataset,row)\n";s=(ROOT/'app/manufacturing_rules.py').read_text();assert s.count(branch)==1;assert hashlib.sha256(s.replace(branch,'').encode()).hexdigest()==m['protected']['app/manufacturing_rules.py']
 from app.schema import SCHEMAS
 from app.finite_schedule_schema import DATASETS
 from app.metric_registry import calculation_hash
 from app.models import MetricVersion
 assert {k:SCHEMAS[k] for k in m['schemas']}==m['schemas'];assert set(SCHEMAS)-set(m['schemas'])==set(DATASETS)
 assert calculation_hash('bi_order_lines')==m['delivery_v8_hash']==MetricVersion.objects.get(pk=11).calculation_hash
 with sqlite3.connect(db) as c:
  c.execute('ATTACH DATABASE ? AS parent',(str(BEFORE/'platform.sqlite3'),));assert c.execute('PRAGMA integrity_check').fetchall()==[('ok',)];assert c.execute('PRAGMA foreign_key_check').fetchall()==[]
  names={n for n, in c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name!='sqlite_sequence'")};assert names==set(m['tables']) and len(names)==52
  deltas={'app_record':6006,'app_importrow':6006,'app_importbatch':1,'app_auditevent':4,'app_importtemplateversion':1}
  for n in sorted(names):
   if n in {'app_importtemplate','app_importtemplateversion'}:continue
   assert c.execute('SELECT count(*) FROM (SELECT * FROM parent."'+n+'" EXCEPT SELECT * FROM main."'+n+'")').fetchone()[0]==0,n
   delta=c.execute('SELECT count(*) FROM main."'+n+'"').fetchone()[0]-m['tables'][n];assert delta==deltas.get(n,0),(n,delta)
  columns=[r[1] for r in c.execute('PRAGMA table_info(app_importtemplate)')];selected=','.join('"'+k+'"' for k in columns if k not in {'current_version','revision'});assert c.execute('SELECT '+selected+' FROM main.app_importtemplate').fetchall()==c.execute('SELECT '+selected+' FROM parent.app_importtemplate').fetchall();assert c.execute('SELECT current_version,revision FROM app_importtemplate').fetchall()==[(5,10)]
  columns=[r[1] for r in c.execute('PRAGMA table_info(app_importtemplateversion)')];selected=','.join('"'+k+'"' for k in columns if k!='state');assert c.execute('SELECT '+selected+' FROM main.app_importtemplateversion WHERE number<=4 ORDER BY number').fetchall()==c.execute('SELECT '+selected+' FROM parent.app_importtemplateversion ORDER BY number').fetchall();assert c.execute('SELECT number,state FROM app_importtemplateversion ORDER BY number').fetchall()==[(1,'retired'),(2,'retired'),(3,'retired'),(4,'retired'),(5,'active')]
 return dict(old_tables=52,old_facts=214523,new_facts=6006,raw_contracts=125,old_contracts_exact=118,physical_preserved=454,protected_exact=31,declared_additive_contract_modules=2,delivery_v8_unchanged=True)
def main():
 p=argparse.ArgumentParser();p.add_argument('mode',choices=['rehearse','release']);a=p.parse_args();before=sha(ROOT/'data/platform.sqlite3');db=ROOT/'data/finite-schedule-rehearsal.sqlite3' if a.mode=='rehearse' else ROOT/'data/platform.sqlite3'
 if a.mode=='rehearse':assert not db.exists();shutil.copy2(BEFORE/'platform.sqlite3',db)
 else:
  assert before==json.loads((BEFORE/'manifest.json').read_text())['database_sha256']
  for n in ('finite_schedule_rehearsal.json','finite_schedule_http.json','finite_schedule_presentation.json'):assert json.loads((ROOT/'data'/n).read_text())['success']
  assert '\nOK\n' in (ROOT/'data/finite_schedule_full_tests.log').read_text()
 os.environ['DJANGO_SETTINGS_MODULE']='config.settings';os.environ['MOTOR_SQLITE_PATH']=str(db);sys.path.insert(0,str(ROOT));import django;django.setup()
 from app.ingestion import stage_file,commit_batch
 from app.models import Record,ImportBatch,ImportRow,AuditEvent,AnalysisModelCard
 from app import finite_schedule_data,finite_schedule,model_card_results
 from django.contrib.auth.models import User
 from django.test import RequestFactory
 from django.urls import resolve
 excel=roundtrip();counts=(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count());batch,replayed=stage_file(BOOK);assert not replayed and batch.summary['total']==batch.summary['valid']==6006;assert batch.summary['unknown_sheets']==[];commit_batch(batch.pk);batch.refresh_from_db();assert batch.status=='committed'
 after=(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count());assert after==tuple(n+d for n,d in zip(counts,(6006,1,6006,2)));again,replayed=stage_file(BOOK);assert replayed and again.pk==batch.pk;commit_batch(again.pk);assert after==(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count())
 profiles=[];scenario=json.loads((ROOT/'data/finite_schedule_scenario.json').read_text())
 for ds,rows in scenario['tables'].items():
  for row in rows:assert Record.objects.get(dataset=ds,business_key=row['id']).values==row
 for profile in scenario['profiles']:
  d=finite_schedule_data.load(profile['id'],profile['policy']);r=d['result'];assert r['summary']==profile['summary'] and r['state']==profile['state'];profiles.append(dict(id=profile['id'],policy=profile['policy'],state=r['state'],summary=r['summary'],sources=len(d['sources'])))
  # Independently check every output reservation against source windows and blocks.
  assigned=[t for t in r['tasks'] if t['state']=='scheduled'];windows={w['id']:w for w in d['tables']['schedule_windows']};tasks={t['id']:t for t in r['tasks']}
  for t in assigned:
   w=windows[t['window_id']];assert w['started']<=t['started']<=t['process_started']<t['finished']<=min(w['finished'],r['study']['horizon_end'])
   for b in d['tables']['schedule_blocks']:
    if b['resource_id']==t['resource_id']:assert t['finished']<=b['started'] or t['started']>=b['finished']
   for e in d['tables']['schedule_edges']:
    if e['to_task_id']==t['id']:
     prev=tasks[e['from_task_id']];assert prev['state']=='scheduled';from datetime import timedelta
     assert finite_schedule.time(t['started'])>=finite_schedule.time(prev['finished'])+timedelta(seconds=finite_schedule.seconds(e['lag_minutes']))
  for resource in r['resources']:
   ordered=sorted([t for t in assigned if t['resource_id']==resource['id']],key=lambda t:t['started']);assert all(a['finished']<=b['started'] for a,b in zip(ordered,ordered[1:]));assert resource['busy_minutes']<=resource['available_minutes']
 from refresh_finite_schedule_template import refresh
 template=refresh(User.objects.get(username='demo_admin'),ROOT);preservation=preserve(db)
 cards=[];factory=RequestFactory()
 for card in AnalysisModelCard.objects.all():
  user=User.objects.get(pk=card.owner_id);url='/api/model-cards/'+str(card.pk)+'/run';request=factory.get(url);request.user=user;match=resolve(url);response=match.func(request,**match.kwargs);assert response.status_code==200,(card.code,response.status_code);value=json.loads(response.content);cards.append(dict(code=card.code,read_current=True))
 assert len(cards)==2
 proof=dict(success=True,mode=a.mode,workbook=excel,**preservation,profiles=profiles,template=template,old_private_cards=cards,batch_id=str(batch.pk),new_audits=4,database_sha256=sha(db),main_unchanged=a.mode=='rehearse',synthetic=True,whole_platform_complete=False,browser_acceptance=False,mobile_acceptance=False,real_system_connected=False)
 if a.mode=='rehearse':assert sha(ROOT/'data/platform.sqlite3')==before
 write(ROOT/('data/finite_schedule_rehearsal.json' if a.mode=='rehearse' else 'data/finite_schedule_release.json'),proof);print(json.dumps(proof,ensure_ascii=False))
if __name__=='__main__':main()
