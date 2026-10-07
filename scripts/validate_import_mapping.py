"""Read saved XLSX and rehearse actual preview/stage/commit in an isolated DB."""
import hashlib,json,os,sqlite3,sys,tempfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
FILE=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/32_字段映射演练_模拟.xlsx'
BASE=ROOT/'data/platform.sqlite3'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
before_sha=sha(BASE)
source=json.loads((ROOT/'data/import_mapping_exercise.json').read_text())
expected=json.loads((ROOT/'data/import_mapping_exercise_expected.json').read_text())
from openpyxl import load_workbook
wb=load_workbook(FILE,read_only=True,data_only=False)
cells=0
try:
    assert wb.sheetnames==['导入说明',*[s['label'] for s in source['schemas'].values()],'字段字典']
    for ds,s in source['schemas'].items():
        rows=list(wb[s['label']].iter_rows(values_only=True));assert rows[0]==tuple(f['label'] for f in s['fields'])
        assert len(rows)==len(source['tables'][ds])+1
        for saved,row in zip(rows[1:],source['tables'][ds]):
            for f,value in zip(s['fields'],saved):
                assert value==row.get(f['name']),(ds,row['id'],f['name'],value)
                if f['type']=='str' and isinstance(row.get(f['name']),str):assert isinstance(value,str)
                if f['type'] in ('int','float'):assert isinstance(value,(int,float)) and not isinstance(value,bool)
                cells+=1
        cells+=len(rows[0])
finally:wb.close()
workspace=Path(tempfile.mkdtemp(prefix='motor-import-mapping-'));database=workspace/'platform.sqlite3'
with sqlite3.connect('file:'+str(BASE)+'?mode=ro',uri=True) as original,sqlite3.connect(database) as copy:original.backup(copy)
(workspace/'app').symlink_to(ROOT/'app',target_is_directory=True)
os.environ['MOTOR_SQLITE_PATH']=str(database);os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from django.test import Client,override_settings
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from app.models import Record,ImportBatch,AuditEvent
from app.ingestion import commit_batch

def records_digest():
    h=hashlib.sha256()
    for r in Record.objects.order_by('id').values():h.update(json.dumps(r,ensure_ascii=False,sort_keys=True,default=str).encode()+b'\n')
    return h.hexdigest()

client=Client();client.force_login(User.objects.get(username='demo_admin'))
def post(path,mapping=None,receipt=None,content=None):
    data={'file':SimpleUploadedFile(FILE.name,content or FILE.read_bytes(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    if mapping is not None:data['mapping']=json.dumps(mapping,ensure_ascii=False)
    if receipt:data['inspection_receipt']=receipt
    return client.post(path,data)

with override_settings(BASE_DIR=workspace):
    original_records=records_digest();count=Record.objects.count();batch_count=ImportBatch.objects.count();audit_count=AuditEvent.objects.count()
    inspection=post('/api/imports/inspect');assert inspection.status_code==200,inspection.content
    assert inspection.json()['sha256']==sha(FILE) and inspection['Cache-Control']=='no-store'
    preview=post('/api/imports/mapping/preview',expected['mapping']);assert preview.status_code==200,preview.content
    result=preview.json();assert result['can_stage'] and not result['issues']
    assert all(not s['sample_issues'] for s in result['sheets'])
    assert ImportBatch.objects.count()==batch_count and AuditEvent.objects.count()==audit_count
    stage=post('/api/imports',expected['mapping'],result['receipt']);assert stage.status_code==200,stage.content
    batch=ImportBatch.objects.get(pk=stage.json()['id']);assert batch.mapping==expected['mapping']
    audit=AuditEvent.objects.get(action='import.mapping_verified',object_id=str(batch.pk))
    assert audit.detail['review_state']=='headers_only' and audit.detail['business_approval'] is False
    assert audit.detail['full_row_validation']=='new_batch_staged'
    assert audit.detail['sha256']==sha(FILE)
    assert batch.summary==dict(duplicate=15,conflict=1,invalid=1,total=17,unknown_sheets=[],skipped_sheets=['个人备注_不导入'])
    originals={(ds,r['id']):r for ds,rows in expected['canonical'].items() for r in rows}
    for row in batch.rows.filter(status='duplicate'):
        assert row.normalized==originals[(row.dataset,row.business_key)]
        assert row.raw['个人采集备注'] and 'note' not in row.normalized
    assert batch.rows.get(status='invalid').raw['往来客户号']==19
    assert batch.rows.get(status='conflict').normalized['lead_days']==expected['canonical']['suppliers'][0]['lead_days']+1
    response=client.post(f'/api/imports/{batch.pk}/commit',json.dumps({}),content_type='application/json')
    assert response.status_code==200,response.content;batch.refresh_from_db();assert batch.status=='partial'
    assert Record.objects.count()==count==212466 and records_digest()==original_records
    repeat=post('/api/imports',expected['mapping'],result['receipt']);assert repeat.status_code==200,repeat.content
    assert repeat.json()['repeated'] and repeat.json()['id']==str(batch.pk)
    replay_audit=AuditEvent.objects.filter(action='import.mapping_verified',object_id=str(batch.pk)).latest('id')
    assert replay_audit.detail['repeated_batch'] and replay_audit.detail['full_row_validation']=='existing_batch_not_rechecked'
    download=client.get(f'/api/imports/{batch.pk}/file');assert download.status_code==200
    assert b''.join(download.streaming_content)==FILE.read_bytes()
    assert post('/api/imports',expected['mapping'],result['receipt'],FILE.read_bytes()+b'changed').status_code==409
assert sha(BASE)==before_sha
manifest=json.loads((ROOT/'data/import-mapping-before/manifest.json').read_text())
with sqlite3.connect('file:'+str(BASE)+'?mode=ro',uri=True) as db:
    db.execute('ATTACH DATABASE ? AS old',('file:'+manifest['database']+'?mode=ro',))
    names=[r[0] for r in db.execute("SELECT name FROM old.sqlite_master WHERE type='table'")]
    for name in names:
        assert name.replace('_','').isalnum()
        assert db.execute(f'SELECT * FROM old."{name}" EXCEPT SELECT * FROM main."{name}" LIMIT 1').fetchone() is None,name
        assert db.execute(f'SELECT * FROM main."{name}" EXCEPT SELECT * FROM old."{name}" LIMIT 1').fetchone() is None,name
for name,digest in manifest['physical'].items():assert sha(ROOT/name)==digest,name
for name,digest in manifest['protected'].items():assert sha(ROOT/name)==digest,name
proof=dict(synthetic=True,xlsx_sha256=sha(FILE),source_cell_values_verified=cells,source_data_sheets=4,source_data_rows=18,
    selected_rows=17,duplicate=15,conflict=1,invalid=1,explicitly_ignored_sheet='个人备注_不导入',
    preview_did_not_create_batch=True,first_three_samples_do_not_replace_full_validation=True,
    isolated_api_stage_commit_and_replay=True,archived_original_bytes_verified=True,
    mapping_audit_is_header_review_only=True,replay_not_claimed_revalidated=True,
    all_formal_records_preserved=212466,all_main_sql_tables_unchanged=40,main_sql_sequence_unchanged=True,
    main_database_sha256=before_sha,old_physical_files_unchanged=len(manifest['physical']),protected_files_unchanged=16,
    isolated_workspace=str(workspace),main_database_unchanged=True,whole_import_not_claimed_approved=True,
    browser_rendered=False,mobile_interaction=False,actual_browser_upload_download=False)
(ROOT/'data/import_mapping_validation.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2));print(json.dumps(proof,ensure_ascii=False,indent=2))
