"""Normal XLSX import on a full parent copy; never seed main business facts."""
import hashlib,json,os,shutil,sqlite3,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
MAIN=ROOT/'data/platform.sqlite3'
REHEARSAL=ROOT/'data/spc-rehearsal'
WORKBOOK=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/34_过程稳定性采样_模拟.xlsx'
PARENT='cc17088ff8da71770975a599da397faee8ae074f6efd85dc83a825f7ea0d4866'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def compare(actual,expected):
    assert actual.keys()==expected.keys(),(actual.keys(),expected.keys())
    for key,value in expected.items():
        assert actual[key]==value,(key,actual[key],value)

def main():
    assert sha(MAIN)==PARENT
    REHEARSAL.mkdir(exist_ok=True);db=REHEARSAL/'platform.sqlite3'
    if db.exists():raise RuntimeError('既有SPC演练库保留，请先核对而非覆盖')
    shutil.copy2(MAIN,db);os.environ['MOTOR_SQLITE_PATH']=str(db);os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings');sys.path.insert(0,str(ROOT))
    import django;django.setup()
    from django.test import override_settings,RequestFactory
    from django.urls import resolve
    from django.contrib.auth import get_user_model
    from openpyxl import load_workbook
    from app.ingestion import stage_file,commit_batch,convert
    from app.schema import SCHEMAS
    from app.models import Record,ImportBatch,ImportRow,AuditEvent
    scenario=json.loads((ROOT/'data/spc_scenario.json').read_text());book=load_workbook(WORKBOOK,read_only=True,data_only=False)
    try:
        for dataset,expected in scenario['tables'].items():
            schema=SCHEMAS[dataset];sheet=book[schema['label']];rows=list(sheet.iter_rows(values_only=True))
            assert list(rows[0])==[f['label'] for f in schema['fields']]
            assert len(rows)-1==len(expected)
            for row,wanted in zip(rows[1:],expected):compare({f['name']:convert(value,f) for f,value in zip(schema['fields'],row)},wanted)
    finally:book.close()
    before=(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count())
    with override_settings(BASE_DIR=REHEARSAL):
        batch,replayed=stage_file(WORKBOOK);assert not replayed and batch.summary==dict(valid=750,total=750,unknown_sheets=[]),batch.summary
        for row in batch.rows.all():compare(row.normalized,next(r for r in scenario['tables'][row.dataset] if r['id']==row.business_key))
        commit_batch(batch.pk);batch.refresh_from_db();assert batch.summary==dict(committed=750,total=750,unknown_sheets=[])
        counts=(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count())
        again,replayed=stage_file(WORKBOOK);assert replayed and again.pk==batch.pk;commit_batch(again.pk)
        assert counts==(Record.objects.count(),ImportBatch.objects.count(),ImportRow.objects.count(),AuditEvent.objects.count())
    assert counts==tuple(n+v for n,v in zip(before,(750,1,750,2)))
    factory=RequestFactory();users=list(get_user_model().objects.all());profiles=[]
    def get(url,user,fields=None):
        request=factory.get(url,fields or {});request.user=user;match=resolve(url);reply=match.func(request,**match.kwargs)
        assert reply.status_code==200,(url,reply.status_code,reply.content[:300]);return json.loads(reply.content)
    reads_before=AuditEvent.objects.count()
    for user in users:
        directory=get('/api/spc',user);assert len(directory['rows'])==10
        for s in directory['rows']:
            d=get('/api/spc/'+s['id'],user);assert len(d['chart_points'])==d['total'] and d['cpk'] is None
            point=d['chart_points'][0]['id'];detail=get('/api/spc/'+s['id']+'/points/'+point,user,dict(receipt=d['receipt']))
            assert any(x['dataset']=='spc_observations' and x['key']==point for x in detail['sources'])
            refs=[];p=1
            while True:
                e=get('/api/spc/'+s['id']+'/sources',user,dict(receipt=d['receipt'],page=p));refs+=e['rows']
                if p*e['size']>=e['total']:break
                p+=1
            assert len(refs)==len({(r['dataset'],r['key']) for r in refs})==e['total']
            if user==users[0]:profiles.append(dict(id=s['id'],state=d['state'],signals=d['signal_counts'],coverage=d['coverage'],issues=d['issues'],limits=d['limits'],sources=e['total']))
    assert AuditEvent.objects.count()==reads_before
    assert [r['state'] for r in profiles]==['trial']*4+['paused']*6
    assert profiles[1]['signals']==dict(baseline=dict(i=1,mr=0),monitor=dict(i=40,mr=3))
    assert profiles[2]['signals']['monitor']==dict(i=1,mr=2)
    assert profiles[3]['coverage']['missing_sequences']==[54]
    with sqlite3.connect(db) as c:
        c.execute('ATTACH DATABASE ? AS parent',(str(ROOT/'data/spc-before/platform.sqlite3'),))
        for name, in c.execute("SELECT name FROM parent.sqlite_master WHERE type='table'").fetchall():
            if name=='sqlite_sequence':continue
            assert c.execute('SELECT count(*) FROM (SELECT * FROM parent."'+name+'" EXCEPT SELECT * FROM main."'+name+'")').fetchone()[0]==0,name
        assert c.execute('PRAGMA integrity_check').fetchall()==[('ok',)]
    assert sha(MAIN)==PARENT
    out=dict(success=True,synthetic=True,main_unchanged=True,workbook_sha256=sha(WORKBOOK),
             workbook_rows=750,all_fields_roundtrip=True,normal_import=dict(batch_id=str(batch.pk),counts=counts,replay_unchanged=True),
             six_roles=len(users),read_only_no_audits=True,old_tables_preserved=True,profiles=profiles)
    (ROOT/'data/spc_rehearsal_validation.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(out,ensure_ascii=False))

if __name__=='__main__':main()
