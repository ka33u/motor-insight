"""Normal Excel import, conservation/window checks and old-fact preservation."""
import argparse
import hashlib
import json
import os
import shutil
import sqlite3
import sys
from collections import defaultdict
from datetime import timedelta
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/39_物料人机联立试排_模拟.xlsx'
CHECKPOINT = ROOT/'data/joint-before/manifest.json'
REHEARSAL = Path('/private/tmp/motorinsight-joint-rehearsal.sqlite3')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def roundtrip():
    from openpyxl import load_workbook
    from app.ingestion import convert
    from app.schema import SCHEMAS
    wanted = json.loads((ROOT/'data/joint_schedule_scenario.json').read_text())
    book = load_workbook(BOOK, read_only=True, data_only=False)
    try:
        assert book.sheetnames == ['导入说明']+[SCHEMAS[ds]['label'] for ds in wanted['tables']]
        for ds, rows in wanted['tables'].items():
            fields = SCHEMAS[ds]['fields']
            actual = list(book[SCHEMAS[ds]['label']].iter_rows(values_only=True))
            assert list(actual[0]) == [f['label'] for f in fields]
            converted = [{f['name']: convert(row[n], f) for n, f in enumerate(fields)} for row in actual[1:]]
            assert converted == rows, ds
    finally:
        book.close()
    return dict(rows=1569, sha256=sha(BOOK), all_fields_exact=True, typed_plant_wall_times=True, visual_sheets_reviewed=5)


def invariants(d):
    from app import finite_schedule as resource
    r = d['result']
    if r['state'] == 'paused':
        assert r['summary'] is None and not r['tasks'] and not r['reservations']
        return
    parent = d['parent']
    raw = parent['base']['tables']
    tasks = {t['id']: t for t in r['tasks']}
    assert len(tasks) == len(raw['schedule_tasks'])
    assigned = [t for t in r['tasks'] if t['state'] == 'scheduled']
    windows = {w['id']: w for w in raw['schedule_windows']}
    person_windows = {w['id']: w for w in parent['tables']['crew_windows']}
    credentials = {q['id']: q for q in r['credentials']}
    for t in assigned:
        for w in (windows[t['window_id']], person_windows[t['worker_window_id']]):
            assert w['started'] <= t['started'] <= t['process_started'] < t['finished'] <= min(w['finished'], r['resource_study']['horizon_end'])
        q = credentials[t['credential_id']]
        assert q['eligible'] and q['employee_id'] == t['employee_id'] and q['process'] == t['process']
        assert q['effective_from'] <= t['started'] < t['finished'] <= q['effective_until']
        for ds, side, identity in [('schedule_blocks', raw, 'resource_id'), ('crew_blocks', parent['tables'], 'employee_id')]:
            for block in side[ds]:
                if block[identity] == t[identity]:
                    assert t['finished'] <= block['started'] or t['started'] >= block['finished']
        for edge in raw['schedule_edges']:
            if edge['to_task_id'] == t['id']:
                previous = tasks[edge['from_task_id']]
                assert previous['state'] == 'scheduled'
                assert resource.time(t['started']) >= resource.time(previous['finished'])+timedelta(seconds=resource.seconds(edge['lag_minutes']))
    for side, identity in [('resources', 'resource_id'), ('workers', 'employee_id')]:
        for v in r[side]:
            ordered = sorted([t for t in assigned if t[identity] == v['id']], key=lambda t: t['started'])
            assert all(a['finished'] <= b['started'] for a, b in zip(ordered, ordered[1:]))
            assert v['busy_minutes'] <= v['available_minutes']+1e-8
            assert abs(v['busy_minutes']-sum((resource.time(t['finished'])-resource.time(t['started'])).total_seconds()/60 for t in ordered)) < 1e-8
    by_lot, by_need = defaultdict(Decimal), defaultdict(Decimal)
    lots = {l['id']: l for l in r['lots']}
    demands = {n['id']: n for n in r['demands']}
    for allocation in r['reservations']:
        lot, need, task = lots[allocation['supply_id']], demands[allocation['demand_id']], tasks[allocation['task_id']]
        assert lot['status'] == '可预留' and task['state'] == 'scheduled'
        assert allocation['reserved_at'] == task['started'] >= lot['available_from']
        assert allocation['material_id'] == lot['material_id'] == need['material_id']
        assert allocation['unit'] == lot['unit'] == need['unit']
        assert allocation['job_id'] == need['job_id'] == task['job_id']
        assert need['route_id'] == task['route_id'] and Decimal(allocation['qty']) > 0
        by_lot[lot['id']] += Decimal(allocation['qty'])
        by_need[need['id']] += Decimal(allocation['qty'])
    for lot in r['lots']:
        assert Decimal(lot['reserved_qty']) == by_lot[lot['id']]
        assert Decimal(lot['usable_qty']) == Decimal(lot['reserved_qty'])+Decimal(lot['remaining_qty'])
        assert Decimal(lot['qty']) == Decimal(lot['usable_qty'])+Decimal(lot['excluded_qty'])
        assert Decimal(lot['remaining_qty']) >= 0
    for need in r['demands']:
        assert by_need[need['id']] == Decimal(need['reserved_qty'])
        assert need['reserved_qty'] in ('0', need['required_qty'])
        if need['reservation']:
            for task in assigned:
                if task['job_id'] == need['job_id'] and task['route_id'] == need['route_id']:
                    assert task['started'] >= need['reservation']['reserved_at']
    assert r['summary']['qty'] == sum(j['qty'] for j in raw['schedule_jobs'])


def frozen_tables(db):
    """Hash every table that this increment is not allowed to write."""
    allowed = {'app_record', 'app_importrow', 'app_importbatch', 'app_auditevent', 'app_importtemplate', 'app_importtemplateversion', 'sqlite_sequence'}
    result = {}
    with sqlite3.connect(db) as c:
        names = [v[0] for v in c.execute("SELECT name FROM sqlite_master WHERE type='table'") if v[0] not in allowed]
        for name in sorted(names):
            digest = hashlib.sha256()
            for row in c.execute('SELECT * FROM "'+name+'" ORDER BY rowid'):
                digest.update(json.dumps(row, ensure_ascii=False, separators=(',', ':'), default=str).encode())
                digest.update(b'\n')
            result[name] = digest.hexdigest()
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['rehearse', 'release'])
    args = parser.parse_args()
    manifest = json.loads(CHECKPOINT.read_text())
    main_db = ROOT/'data/platform.sqlite3'
    assert sha(main_db) == manifest['source_database_sha'], 'Main database changed since checkpoint; inspect before proceeding.'
    if args.mode == 'rehearse':
        assert not REHEARSAL.exists(), 'Use the existing proof or inspect the prior rehearsal; never overwrite it.'
        shutil.copy2(main_db, REHEARSAL)
        db = REHEARSAL
    else:
        for name in ('joint_schedule_rehearsal.json', 'joint_schedule_ui.json', 'joint_schedule_bootstrap.json'):
            assert json.loads((ROOT/'data'/name).read_text())['success'], name
        assert '\nOK\n' in Path('/private/tmp/motorinsight-joint-full-tests.log').read_text()
        db = main_db
    os.environ['DJANGO_SETTINGS_MODULE'] = 'config.settings'
    os.environ['MOTOR_SQLITE_PATH'] = str(db)
    sys.path.insert(0, str(ROOT))
    import django
    django.setup()
    from app.ingestion import stage_file, commit_batch
    from app.models import Record, ImportBatch, ImportRow, AuditEvent, MetricVersion
    from app.schema import SCHEMAS
    from app import joint_schedule_data, crew_schedule_data, finite_schedule
    from app.joint_schedule_schema import DATASETS
    from app.metric_registry import calculation_hash
    from django.contrib.auth.models import User
    frozen = frozen_tables(db)
    excel = roundtrip()
    before = (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    batch, replayed = stage_file(BOOK)
    assert not replayed and batch.summary['total'] == batch.summary['valid'] == 1569, batch.summary
    assert batch.summary['unknown_sheets'] == []
    commit_batch(batch.pk)
    batch.refresh_from_db()
    assert batch.status == 'committed'
    after = (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    assert after == tuple(n+d for n, d in zip(before, (1569, 1, 1569, 2)))
    again, replayed = stage_file(BOOK)
    assert replayed and again.pk == batch.pk
    commit_batch(again.pk)
    assert after == (Record.objects.count(), ImportBatch.objects.count(), ImportRow.objects.count(), AuditEvent.objects.count())
    scenario = json.loads((ROOT/'data/joint_schedule_scenario.json').read_text())
    for ds, rows in scenario['tables'].items():
        assert list(Record.objects.filter(dataset=ds).order_by('business_key').values_list('values', flat=True)) == sorted(rows, key=lambda r: r['id']), ds
    profiles = []
    for profile in scenario['profiles']:
        d = joint_schedule_data.load(profile['id'], profile['policy'])
        r = d['result']
        assert (r['state'], r['summary'], r['issues']) == (profile['state'], profile['summary'], profile['issues'])
        invariants(d)
        profiles.append(dict(id=profile['id'], policy=profile['policy'], state=r['state'], summary=r['summary'], source_count=len(d['sources'])))
    for key in ('CR-261002-001', 'CR-261002-004'):
        old = json.loads((ROOT/f'data/joint-before/{key}.json').read_text())
        current = crew_schedule_data.load(key)
        assert current['result'] == old['result'] and current['source_hash'] == old['source_hash']
    from refresh_joint_template import refresh
    template = refresh(User.objects.get(username='demo_admin'), ROOT)
    old_facts = list(Record.objects.exclude(dataset__in=DATASETS).order_by('dataset', 'business_key').values('dataset', 'business_key', 'record_hash'))
    assert len(old_facts) == manifest['record_count']
    assert finite_schedule.digest(old_facts) == manifest['record_digest']
    assert {k: SCHEMAS[k] for k in manifest['schemas']} == manifest['schemas']
    assert set(SCHEMAS)-set(manifest['schemas']) == set(DATASETS)
    for name, digest in manifest['originals'].items():
        assert sha(ROOT/name) == digest, name
    assert calculation_hash('bi_order_lines') == MetricVersion.objects.get(pk=11).calculation_hash == '9d9a5af3d45172eca16790d628a36bc6acbe18d1e2a811a66a5b5fbff7037fe3'
    assert frozen_tables(db) == frozen
    with sqlite3.connect(db) as c:
        assert c.execute('PRAGMA integrity_check').fetchall() == [('ok',)]
        assert c.execute('PRAGMA foreign_key_check').fetchall() == []
    if args.mode == 'rehearse':
        assert sha(main_db) == manifest['source_database_sha']
    proof = dict(success=True, mode=args.mode, workbook=excel, profiles=profiles, old_facts_preserved=len(old_facts), original_files_preserved=len(manifest['originals']),
                 old_contracts_preserved=len(manifest['schemas']), frozen_tables_preserved=len(frozen), metric_v8_unchanged=True, crew_outputs_unchanged=True,
                 template=template, batch_id=str(batch.pk), records=Record.objects.count(), database_sha256=sha(db), synthetic=True, browser_acceptance=False, mobile_acceptance=False)
    path = ROOT/('data/joint_schedule_rehearsal.json' if args.mode == 'rehearse' else 'data/joint_schedule_release.json')
    path.write_text(json.dumps(proof, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({k: v for k, v in proof.items() if k != 'profiles'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
