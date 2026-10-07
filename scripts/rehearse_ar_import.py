"""Read the actual workbook, round-trip every cell, rehearse import in a DB copy."""
import os,sys,json,sqlite3,tempfile,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
workspace=Path(tempfile.mkdtemp(prefix='motor-ar-rehearsal-'))
with sqlite3.connect(ROOT/'data/platform.sqlite3') as origin,sqlite3.connect(workspace/'platform.sqlite3') as dest:origin.backup(dest)
os.environ['MOTOR_SQLITE_PATH']=str(workspace/'platform.sqlite3');os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from django.test import override_settings
from openpyxl import load_workbook
from app.schema import SCHEMAS
from app.models import Record
from app.ingestion import convert,stage_file,commit_batch
from app import receivables,analytics

file=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/13_历史应收_模拟.xlsx'
expected=json.loads((ROOT/'data/ar_history_scenario.json').read_text())
book=load_workbook(file,read_only=True,data_only=False);parsed={};cells=0
for ds in ['ar_opening','ar_events']:
    schema=SCHEMAS[ds];sheet=book[schema['label']];rows=sheet.iter_rows(values_only=True)
    assert list(next(rows))==[f['label'] for f in schema['fields']]
    parsed[ds]=[{f['name']:convert(v,f) for f,v in zip(schema['fields'],row)} for row in rows]
    assert parsed[ds]==expected['tables'][ds],ds
    cells+=len(parsed[ds])*len(schema['fields'])
book.close()
before=Record.objects.count()
with override_settings(BASE_DIR=workspace):
    batch,repeated=stage_file(file);assert not repeated
    assert batch.summary.get('valid')==102 and not batch.summary.get('invalid') and not batch.summary.get('conflict'),batch.summary
    batch=commit_batch(batch.pk)
    assert Record.objects.count()==before+102
    again,repeated=stage_file(file);assert repeated and again.pk==batch.pk
    commit_batch(again.pk);assert Record.objects.count()==before+102
    for ds,rows in parsed.items():assert {r['id']:r for r in Record.objects.filter(dataset=ds).values_list('values',flat=True)}=={r['id']:r for r in rows}
    d=receivables.current();assert not d.global_issues
    assert all(not r['issues'] for r in d.rows),[(r['key'],r['issues']) for r in d.rows if r['issues']]
    report=dict(workbook=str(file),file_sha256=hashlib.sha256(file.read_bytes()).hexdigest(),roundtrip_cells=cells,roundtrip_rows=102,isolated_import=True,idempotent=True,records_before=before,records_after=Record.objects.count(),summary=receivables.summary(d.rows),by_kind=receivables.breakdown(d.rows)['by_kind'],aging=receivables.breakdown(d.rows)['aging'])
(ROOT/'data/ar_import_rehearsal.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
