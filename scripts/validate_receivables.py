"""Reconcile imported AR against raw facts and a preserved pre-import snapshot."""
import os,sys,json,csv,hashlib,sqlite3,zipfile,tempfile
from pathlib import Path
from collections import defaultdict,Counter
from datetime import date
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django;django.setup()
from openpyxl import load_workbook
from app import receivables,metric_registry
from app.models import Record,IssueDisposition,AuditEvent,MetricVersion
from app.schema import SCHEMAS
from app.ingestion import convert

baseline=ROOT/'data/backups/motor-backup-20261003-173402.zip'
preserved={}
with zipfile.ZipFile(baseline) as z,tempfile.TemporaryDirectory() as temp:
    manifest=json.loads(z.read('manifest.json'))
    for item in manifest['files']:
        if item['path'].startswith(('data/imports/','data/device_files/')):
            assert hashlib.sha256((ROOT/item['path']).read_bytes()).hexdigest()==item['sha256']
    beforefile=Path(temp)/'baseline.sqlite3';beforefile.write_bytes(z.read('data/platform.sqlite3'))
    with sqlite3.connect(beforefile) as old,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
        tables=[r[0] for r in old.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'app_%'")]
        additions={'app_record':102,'app_importbatch':1,'app_importrow':102,'app_issuedisposition':1}
        for table in tables:
            cols=[r[1] for r in old.execute(f'PRAGMA table_info({table})')];index=cols.index('id')
            before=list(old.execute(f'SELECT * FROM {table} ORDER BY id'));after=list(now.execute(f'SELECT * FROM {table} ORDER BY id'))
            current={r[index]:r for r in after}
            assert all(current.get(r[index])==r for r in before),table
            if table!='app_auditevent':assert len(after)-len(before)==additions.get(table,0),(table,len(before),len(after))
            preserved[table]={'before':len(before),'after':len(after),'all_prior_rows_unchanged':True}

raw=defaultdict(list)
for ds,value in Record.objects.values_list('dataset','values').iterator():raw[ds].append(value)
assert sum(map(len,raw.values()))==106940 and len(raw)==51
assert len(raw['ar_opening'])==42 and len(raw['ar_events'])==60
sourcefile=ROOT/'outputs/01a0f580-2480-7820-bc97-8ca293421c39/13_历史应收_模拟.xlsx'
book=load_workbook(sourcefile,read_only=True,data_only=False);cells=0
for ds in ['ar_opening','ar_events']:
    schema=SCHEMAS[ds];rows=book[schema['label']].iter_rows(values_only=True)
    assert list(next(rows))==[f['label'] for f in schema['fields']]
    parsed=[{f['name']:convert(v,f) for f,v in zip(schema['fields'],row)} for row in rows]
    assert {r['id']:r for r in parsed}=={r['id']:r for r in raw[ds]}
    cells+=len(parsed)*len(schema['fields'])
book.close();assert cells==960

cutoff=date(2026,10,1);expected={};events=defaultdict(list);payments=defaultdict(list)
for e in raw['ar_events']:events[e['opening_id']].append(e)
for p in raw['payments']:payments[p['invoice_id']].append(p)
def age(balance,due):
    days=(cutoff-date.fromisoformat(due)).days
    if balance==0:return 'settled'
    if days<0:return 'not_due'
    if days==0:return 'due_today'
    return next(key for cap,key in [(30,'d1_30'),(60,'d31_60'),(90,'d61_90'),(180,'d91_180'),(100000,'d181_plus')] if days<=cap)
for kind,ds in [('opening','ar_opening'),('invoice','invoices')]:
    for r in raw[ds]:
        basis=r['opening_balance_cents'] if kind=='opening' else r['net_cents']+r['tax_cents']
        valid=[e for e in events[r['id']] if e['status']=='已过账' and e['occurred']<=cutoff.isoformat()] if kind=='opening' else []
        receipt=-sum(e['delta_cents'] for e in valid if e['kind']=='收款核销') if kind=='opening' else sum(p['amount_cents'] for p in payments[r['id']] if p['paid']<=cutoff.isoformat())
        reversal=sum(e['delta_cents'] for e in valid if e['kind']=='核销冲回');credit=-sum(e['delta_cents'] for e in valid if e['kind']=='贷项冲减')
        balance=basis-receipt+reversal-credit;assert balance>=0
        key=kind+':'+r['id'];expected[key]=dict(kind=kind,id=r['id'],customer_id=r['customer_id'],basis_cents=basis,receipt_cents=receipt,reversal_cents=reversal,credit_cents=credit,balance_cents=balance,bucket=age(balance,r['due']))
d=receivables.current();assert not d.global_issues and len(d.rows)==264
for r in d.rows:
    assert not r['issues'] and all(r[k]==v for k,v in expected[r['key']].items()),r['key']
balance=sum(r['balance_cents'] for r in expected.values())
overdue=sum(r['balance_cents'] for r in expected.values() if r['bucket'] not in ['not_due','due_today','settled'])
assert balance==707499477 and overdue==186629600
s=receivables.summary(d.rows);assert s['balance_cents']==balance and s['overdue_cents']==overdue
assert abs(s['overdue_ratio']-overdue/balance*100)<1e-10
breakdown=receivables.breakdown(d.rows)
for group in breakdown['aging']:
    match=[r for r in expected.values() if r['bucket']==group['key']]
    assert group['rows']==len(match) and group['value']==sum(r['balance_cents'] for r in match)
for group in breakdown['customers']:
    match=[r for r in expected.values() if r['customer_id']==group['id']]
    assert group['rows']==len(match) and group['balance_cents']==sum(r['balance_cents'] for r in match)
for group in breakdown['by_kind']:
    match=[r for r in expected.values() if r['kind']==group['kind']]
    for field in ['basis_cents','receipt_cents','reversal_cents','credit_cents','balance_cents']:
        assert group[field]==sum(r[field] for r in match)

source_keys={(r['dataset'],r['key']) for row in d.rows for r in row['sources']};files={}
for ds,key in source_keys:
    r=Record.objects.select_related('source_row__batch').get(dataset=ds,business_key=key)
    assert r.values==r.source_row.normalized and r.record_hash==r.source_row.record_hash
    b=r.source_row.batch;files[str(b.pk)]={'path':b.file_path,'sha256':b.file_hash,'filename':b.filename}
for f in files.values():assert hashlib.sha256(Path(f['path']).read_bytes()).hexdigest()==f['sha256']
exportfile=ROOT/'outputs/应收31至60天_浏览器导出.csv'
exported=list(csv.reader(exportfile.open(encoding='utf-8-sig')));assert exported[0][2]=='2026-10-01'
assert json.loads(exported[0][4])['bucket']=='d31_60'
selected={r['id']:r for r in expected.values() if r['bucket']=='d31_60'}
assert len(exported[4:])==5 and {r[0] for r in exported[4:]}==set(selected)
for r in exported[4:]:
    e=selected[r[0]]
    for index,key in [(9,'basis_cents'),(10,'receipt_cents'),(11,'reversal_cents'),(13,'credit_cents'),(14,'balance_cents')]:assert int(r[index])==e[key]
assert sum(int(r[14]) for r in exported[4:])==33140646
note=IssueDisposition.objects.get(key='ar:opening:QCYS-260901-00004')
assert note.version==1 and note.status=='待客户确认'
assert AuditEvent.objects.filter(action='receivable.followup').count()==1
v=MetricVersion.objects.select_related('metric').get(metric__key='DELIVERY_OTIF',version=3)
assert v.status=='published' and v.calculation_hash==metric_registry.calculation_hash(v.metric.dataset)
record_hash=hashlib.sha256()
for r in Record.objects.order_by('dataset','business_key').values_list('dataset','business_key','values','revision','source_row_id').iterator():record_hash.update(json.dumps(r,ensure_ascii=False,sort_keys=True).encode())
report=dict(synthetic=True,baseline=str(baseline.relative_to(ROOT)),business_rows=106940,source_datasets=51,source_workbooks=13,business_sha256=record_hash.hexdigest(),new_excel_rows=102,new_excel_cells=cells,summary=s,aging=breakdown['aging'],by_kind=breakdown['by_kind'],source_records=len(source_keys),source_files=len(files),preserved_tables=preserved,browser_csv_rows=5,browser_csv_sha256=hashlib.sha256(exportfile.read_bytes()).hexdigest(),tests=370,new_tests=27,limits='Only imported synthetic CNY AR at date precision; no U8/MES connection, unbilled AR, prepayment, bank reconciliation, multi-currency, bad-debt or closed-period ledger.')
if len(sys.argv)>1:
    backup=Path(sys.argv[1])
    with zipfile.ZipFile(backup) as z,tempfile.TemporaryDirectory() as temp:
        manifest=json.loads(z.read('manifest.json'))
        for f in manifest['files']:
            content=z.read(f['path']);assert len(content)==f['size'] and hashlib.sha256(content).hexdigest()==f['sha256']
        p=Path(temp)/'restored.sqlite3';p.write_bytes(z.read('data/platform.sqlite3'))
        with sqlite3.connect(p) as restored,sqlite3.connect(ROOT/'data/platform.sqlite3') as now:
            assert restored.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
            for t in tables:assert list(restored.execute(f'SELECT * FROM {t} ORDER BY id'))==list(now.execute(f'SELECT * FROM {t} ORDER BY id')),t
        report['backup']={'path':str(backup),'manifest_files':len(manifest['files']),'sqlite_integrity':'ok','all_app_tables_match':True}
(ROOT/'data/receivables_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str))
print(json.dumps({k:v for k,v in report.items() if k!='preserved_tables'},ensure_ascii=False,default=str))
