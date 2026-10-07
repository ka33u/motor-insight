"""Explicit XLSX column mapping; preview never changes business records."""
import hashlib
import json
import zipfile
from pathlib import Path
from django.core import signing
from openpyxl import load_workbook
from .schema import SCHEMAS
from . import access

SALT='motor.import.mapping.v1'
MAX_AGE=15*60
META_SHEETS={'导入说明','字段字典'}
NOTICE='这是表头与映射预检，尚未逐行入库。编号需原生文本，金额、单位和日期不自动换算；完整行级校验及冲突审核仍在导入批次办理。'


class MappingStale(ValueError):
    pass


def digest(value):
    return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()


def rules_hash():
    root=Path(__file__).resolve().parent
    return digest({name:hashlib.sha256((root/name).read_bytes()).hexdigest()
        for name in ['import_mapping.py','ingestion.py','manufacturing_rules.py','import_templates.py']})


def parse_mapping(raw):
    if not isinstance(raw,str) or len(raw)>200000:
        raise ValueError('映射JSON不存在或过长')
    def pairs(items):
        result={}
        for key,value in items:
            if key in result:raise ValueError('映射含重复键：'+key)
            result[key]=value
        return result
    value=json.loads(raw,object_pairs_hook=pairs)
    validate_shape(value)
    return value


def validate_shape(mapping):
    if not isinstance(mapping,dict) or len(mapping)>128:raise ValueError('映射必须为不超过128页的对象')
    for sheet,conf in mapping.items():
        if not isinstance(sheet,str) or not sheet or len(sheet)>31 or not isinstance(conf,dict):raise ValueError('工作表映射格式无效')
        if set(conf)-{'dataset','fields','ignore','skip'}:raise ValueError(sheet+'含未知映射设置')
        if conf.get('skip') is True:
            if conf!={'skip':True}:raise ValueError(sheet+'忽略整页时不能同时映射字段')
            continue
        if 'skip' in conf:raise ValueError(sheet+'忽略整页必须明确为true')
        ds=conf.get('dataset')
        if ds is not None and (not isinstance(ds,str) or ds not in SCHEMAS):raise ValueError(sheet+'目标业务表不存在')
        fields=conf.get('fields',{});ignored=conf.get('ignore',[])
        if not isinstance(fields,dict) or len(fields)>512 or any(not isinstance(k,str) or not k or not isinstance(v,str) or not v for k,v in fields.items()):raise ValueError(sheet+'字段映射必须为文本对象')
        if not isinstance(ignored,list) or len(ignored)>512 or any(not isinstance(x,str) or not x for x in ignored) or len(set(ignored))!=len(ignored):raise ValueError(sheet+'忽略列设置无效')
        if set(ignored)&set(fields):raise ValueError(sheet+'同一列不能同时映射和忽略')


def headers_of(sheet):
    if sheet.max_column and sheet.max_column>512:raise ValueError(sheet.title+'超过512列，请先裁剪空白和无关列')
    first=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())
    headers=[str(v).strip() if v is not None else '' for v in first]
    nonempty=[x for x in headers if x]
    if len(set(nonempty))!=len(nonempty):raise ValueError(sheet.title+'有重复表头')
    if any(len(x)>500 for x in headers):raise ValueError(sheet.title+'表头过长')
    return headers


def columns_for(schema,headers,conf):
    """Resolve each canonical target once, including implicit standard aliases."""
    names={f['name']:f for f in schema['fields']}
    labels={f['label']:f['name'] for f in schema['fields']}
    aliases={**labels,**{n:n for n in names}}
    fields=conf.get('fields',{});ignored=set(conf.get('ignore',[]))
    if set(fields)-set(headers) or ignored-set(headers):raise ValueError('映射或忽略列在实际工作表中不存在')
    columns={};unmatched=[]
    for pos,h in enumerate(headers):
        if not h or h in ignored:continue
        target=fields.get(h,h)
        if h in fields and target not in aliases:raise ValueError('目标字段不存在：'+target)
        target=aliases.get(target)
        if target is None:unmatched.append(h);continue
        if target in columns:raise ValueError('多列映射到同一字段：'+names[target]['label']+'；请明确选择并忽略其余列')
        columns[target]=pos
    return columns,unmatched


def check_zip(path):
    path=Path(path)
    if path.suffix.lower()!='.xlsx':raise ValueError('目前只接受 .xlsx 文件')
    if path.stat().st_size>50*1024*1024:raise ValueError('文件超过50MB')
    try:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist())>350*1024*1024:raise ValueError('解压后的工作簿超过允许大小')
            if any(i.filename.lower().endswith('vbaproject.bin') for i in z.infolist()):raise ValueError('不接受含宏工作簿')
    except zipfile.BadZipFile as ex:raise ValueError('不是可读取的XLSX工作簿') from ex


def inspect(path,user,mapping=None,strict=False,filename=''):
    check_zip(path)
    if mapping is not None:validate_shape(mapping)
    names={s['label']:s for s in SCHEMAS.values()}
    source=hashlib.sha256(Path(path).read_bytes()).hexdigest()
    try:book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    except (ValueError,KeyError,OSError) as ex:raise ValueError('工作簿结构无法读取') from ex
    try:
        if len(book.sheetnames)>128:raise ValueError('工作簿超过128个工作表，请分批导入')
        if mapping is not None and set(mapping)-set(book.sheetnames):raise ValueError('映射含工作簿中不存在的工作表')
        sheets=[];errors=[];selected=0
        from .ingestion import convert,json_value
        for sheet in book:
            meta=sheet.title in META_SHEETS
            conf=(mapping or {}).get(sheet.title,{})
            schema=SCHEMAS.get(conf.get('dataset')) if conf.get('dataset') else names.get(sheet.title)
            headers=headers_of(sheet)
            samples=[dict(row=i,values=[json_value(v) for v in values]) for i,values in enumerate(sheet.iter_rows(min_row=2,max_row=4,max_col=len(headers) or 1,values_only=True),2) if any(v is not None for v in values)] if not meta else []
            entry=dict(name=sheet.title,headers=headers,samples=samples,declared_last_row=sheet.max_row,
                dataset=schema['key'] if schema and not conf.get('skip') and not meta else None,
                ignored_sheet=bool(meta or conf.get('skip')),metadata_sheet=meta,bindings={},issues=[],sample_issues=[],unmatched=[])
            if not entry['ignored_sheet']:
                if strict and sheet.title not in (mapping or {}):entry['issues'].append('请明确选择业务表或忽略整页')
                if not schema:entry['issues'].append('尚未选择目标业务表')
                else:
                    selected+=1
                    # Suggestions require one exact alias; ambiguous aliases stay blank.
                    for f in schema['fields']:
                        candidates=[h for h in headers if h in {f['name'],f['label']}]
                        if len(candidates)==1:entry['bindings'][f['name']]=candidates[0]
                    try:
                        columns,unmatched=columns_for(schema,headers,conf)
                        entry['bindings']={name:headers[pos] for name,pos in columns.items()}
                        entry['unmatched']=unmatched
                        if strict and unmatched:entry['issues'].append('存在未映射且未明确忽略的列：'+'、'.join(unmatched))
                        if strict:
                            missing=[f['label'] for f in schema['fields'] if f['required'] and f['name'] not in columns]
                            if missing:entry['issues'].append('必填列未映射：'+'、'.join(missing))
                        for sample in samples:
                            for f in schema['fields']:
                                pos=columns.get(f['name']);v=sample['values'][pos] if pos is not None and pos<len(sample['values']) else None
                                try:convert(v,f)
                                except (ValueError,TypeError,OverflowError) as ex:entry['sample_issues'].append(dict(row=sample['row'],field=f['name'],label=f['label'],message=str(ex)))
                    except ValueError as ex:entry['issues'].append(str(ex))
            errors.extend(sheet.title+'：'+x for x in entry['issues']);sheets.append(entry)
        if strict and not selected:errors.append('至少选择一张业务表')
        result=dict(filename=filename,sha256=source,schema_hash=digest(SCHEMAS),rules_hash=rules_hash(),mapping=mapping,
            sheets=sheets,schemas=[s for key,s in SCHEMAS.items() if access.allowed(user,key)],issues=errors,
            can_stage=bool(strict and not errors),sample_rows_per_sheet=3,notice=NOTICE)
        if result['can_stage']:
            result['receipt']=signing.dumps(dict(user_id=user.pk,role=access.role(user),sha256=source,filename=filename,
                schema_hash=result['schema_hash'],rules_hash=result['rules_hash'],mapping_hash=digest(mapping)),salt=SALT,compress=True)
            result['receipt_valid_seconds']=MAX_AGE
        return result
    finally:book.close()


def verify_receipt(receipt,path,filename,mapping,user):
    try:payload=signing.loads(receipt,salt=SALT,max_age=MAX_AGE)
    except (signing.BadSignature,ValueError,TypeError) as ex:raise MappingStale('映射预检依据已失效，请重新核对工作簿') from ex
    expected=dict(user_id=user.pk,role=access.role(user),sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest(),
        filename=filename,schema_hash=digest(SCHEMAS),rules_hash=rules_hash(),mapping_hash=digest(mapping))
    if payload!=expected:raise MappingStale('账号、文件、字段定义或映射已变化，请重新预检')
    return payload
