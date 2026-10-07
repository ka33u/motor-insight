"""XLSX -> staged raw rows -> validation -> atomic commit.

Business imports never read scenario.json. Repeat ingestion is idempotent.
Different content under an existing key is quarantined. Only the explicit
version-checked review workflow can approve replacement with retained history.
"""
import hashlib,json,math,zipfile
from datetime import date,datetime
from pathlib import Path
from collections import Counter
from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from .models import ImportBatch,ImportRow,Record,AuditEvent
from .schema import SCHEMAS
from .manufacturing_rules import issues as manufacturing_issues
from .import_mapping import validate_shape,columns_for,headers_of

def fingerprint(obj):return hashlib.sha256(json.dumps(obj,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
def json_value(v):return v.isoformat() if isinstance(v,(date,datetime)) else v

def convert(value,field):
    if value is None or value=='':
        if field['required']:raise ValueError('必填值缺失')
        return None
    if isinstance(value,str) and value.startswith('='):raise ValueError('来源列不可使用公式；需导出经核验的值')
    kind=field['type']
    if kind=='str':
        # Keep leading zeroes. Reject numeric IDs rather than guess lost digits.
        if not isinstance(value,str):raise ValueError('需文本值，避免编号前导零丢失')
        value=value.strip()
        if not value and field['required']:raise ValueError('必填文本不能为空白')
        return value or None
    if kind in ['int','float']:
        if isinstance(value,bool):raise ValueError('布尔值不能代替数值')
        n=float(value)
        if not math.isfinite(n):raise ValueError('数值必须有限')
        if kind=='int' and not n.is_integer():raise ValueError('必须为整数')
        return int(n) if kind=='int' else n
    if kind=='bool':
        if isinstance(value,bool):return value
        if value in [0,1,'0','1','是','否','true','false','TRUE','FALSE']:return value in [1,'1','是','true','TRUE']
        raise ValueError('需是/否或布尔值')
    if kind in ['date','datetime']:
        if isinstance(value,date) and not isinstance(value,datetime):value=datetime.combine(value,datetime.min.time())
        elif not isinstance(value,datetime):value=datetime.fromisoformat(str(value))
        return value.date().isoformat() if kind=='date' else value.isoformat(timespec='seconds')
    raise ValueError('未知字段类型')

def stage_file(path,filename=None,mapping=None,*,force_recheck=False):
    path=Path(path);mapping=mapping or {}
    if path.suffix.lower()!='.xlsx':raise ValueError('目前只接受 .xlsx 文件')
    if path.stat().st_size>50*1024*1024:raise ValueError('文件超过50MB')
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist())>350*1024*1024:raise ValueError('解压后的工作簿超过允许大小')
        if any(i.filename.lower().endswith('vbaproject.bin') for i in z.infolist()):raise ValueError('不接受含宏工作簿')
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    previous=ImportBatch.objects.filter(file_hash=digest,status__in=['committed','partial']).order_by('-created_at').first()
    if previous and previous.mapping==mapping and not force_recheck:return previous,True
    batch=ImportBatch.objects.create(filename=(filename or path.name)[:255],file_hash=digest,file_path='',mapping=mapping)
    archive=settings.BASE_DIR/'data/imports';archive.mkdir(parents=True,exist_ok=True)
    dest=archive/f'{batch.id}.xlsx';dest.write_bytes(path.read_bytes());batch.file_path=str(dest);batch.save(update_fields=['file_path'])
    book=load_workbook(dest,read_only=True,data_only=False)
    names={s['label']:s for s in SCHEMAS.values()}; staged=[];seen={};unknown=[];skipped=[]
    existing={(r['dataset'],r['business_key']):r['record_hash'] for r in Record.objects.values('dataset','business_key','record_hash')}
    try:
        validate_shape(mapping)
        if set(mapping)-set(book.sheetnames):raise ValueError('映射含工作簿中不存在的工作表')
        for sheet in book:
            if sheet.title in ['导入说明','字段字典']:continue
            conf=mapping.get(sheet.title,{})
            if conf.get('skip'):
                skipped.append(sheet.title);continue
            schema=SCHEMAS.get(conf.get('dataset')) if conf.get('dataset') else names.get(sheet.title)
            if not schema:unknown.append(sheet.title);continue
            iterator=sheet.iter_rows(values_only=True);headers=next(iterator,None)
            if not headers:continue
            headers=headers_of(sheet)
            columns,_=columns_for(schema,headers,conf)
            for row_number,values in enumerate(iterator,2):
                if not any(v is not None for v in values):continue
                raw={h:json_value(values[i] if i<len(values) else None) for i,h in enumerate(headers) if h}
                data={};errors=[]
                for f in schema['fields']:
                    pos=columns.get(f['name']);val=values[pos] if pos is not None and pos<len(values) else None
                    try:data[f['name']]=convert(val,f)
                    except (ValueError,TypeError,OverflowError) as ex:errors.append({'field':f['name'],'label':f['label'],'message':str(ex)})
                key=str(data.get(schema['primary_key']) or '')
                rh=fingerprint(data);identity=(schema['key'],key);status='invalid' if errors else 'valid'
                # Validate amounts before deduplicating. An invalid row must never
                # satisfy another row's reference through a duplicate copy.
                for name in ['qty','planned_qty','input_qty','good_qty','scrap_qty','rework_qty','amount_cents','unit_price_cents']:
                    if name in data and data[name] is not None and data[name]<0:
                        status='invalid';errors.append({'field':name,'message':'该字段不允许负数'})
                errors.extend(manufacturing_issues(schema['key'],data))
                if errors:status='invalid'
                if not errors:
                    prior=seen.get(identity,existing.get(identity))
                    if prior==rh:status='duplicate'
                    elif prior is not None:status='conflict';errors.append({'field':schema['primary_key'],'message':'相同主键内容不同，需审核更正，原记录未覆盖'})
                    else:seen[identity]=rh
                staged.append(ImportRow(batch=batch,sheet=sheet.title,row_number=row_number,dataset=schema['key'],business_key=key,raw=raw,normalized=data,record_hash=rh,status=status,issues=errors))
        # Iteration removes references to rows invalidated in this batch.
        while True:
            keys=set(existing)|{(r.dataset,r.business_key) for r in staged if r.status=='valid'}
            changed=False
            for r in staged:
                if r.status not in ['valid','duplicate']:continue
                for f in SCHEMAS[r.dataset]['fields']:
                    val=r.normalized.get(f['name'])
                    if f['reference'] and val is not None and (f['reference'],val) not in keys:
                        r.issues.append({'field':f['name'],'label':f['label'],'message':f'关联记录不存在：{f["reference"]}/{val}'});r.status='invalid';changed=True
            if not changed:break
        if not staged:raise ValueError('没有识别到可导入的数据行，请检查工作表名称或字段映射')
        with transaction.atomic():
            ImportRow.objects.bulk_create(staged,batch_size=700)
            counts=dict(Counter(r.status for r in staged));counts['total']=len(staged);counts['unknown_sheets']=unknown
            if skipped:counts['skipped_sheets']=skipped
            batch.status='staged';batch.summary=counts;batch.save(update_fields=['status','summary'])
            AuditEvent.objects.create(action='import.stage',object_type='ImportBatch',object_id=str(batch.id),detail={'filename':batch.filename,'counts':counts})
    except Exception as ex:
        batch.status='failed';batch.summary={'error':str(ex)};batch.save(update_fields=['status','summary']);raise
    finally:book.close()
    return batch,False

def commit_batch(batch_id):
    with transaction.atomic():
        batch=ImportBatch.objects.select_for_update().get(pk=batch_id)
        if batch.status in ['committed','partial']:return batch
        if batch.status!='staged':raise ValueError('该批次尚未完成校验')
        pending=list(batch.rows.filter(status='valid'))
        # Re-run current business rules: a preview may predate a rule update.
        for r in pending:
            errors=manufacturing_issues(r.dataset,r.normalized)
            if errors:r.status='invalid';r.issues=errors
        current={(r['dataset'],r['business_key']):r['record_hash'] for r in Record.objects.values('dataset','business_key','record_hash')}
        # References may have changed since staging. Recheck the dependency set
        # before writing, including cascades through invalid staged parents.
        while True:
            keys=set(current)|{(r.dataset,r.business_key) for r in pending if r.status=='valid'}
            modified=False
            for r in pending:
                if r.status!='valid':continue
                for field in SCHEMAS[r.dataset]['fields']:
                    v=r.normalized.get(field['name'])
                    if field['reference'] and v is not None and (field['reference'],v) not in keys:
                        r.status='invalid';r.issues=[{'field':field['name'],'message':'提交时关联记录已不存在，请重新校验'}];modified=True
            if not modified:break
        to_create=[];changed=[]
        for r in pending:
            if r.status!='valid':changed.append(r);continue
            existing=current.get((r.dataset,r.business_key))
            if existing:
                r.status='duplicate' if existing==r.record_hash else 'conflict'
                if r.status=='conflict':r.issues=[{'message':'校验后发生并发导入，主键内容冲突，请重新审核'}]
            else:
                to_create.append(Record(dataset=r.dataset,business_key=r.business_key,values=r.normalized,record_hash=r.record_hash,source_row=r));r.status='committed'
                current[(r.dataset,r.business_key)]=r.record_hash
            changed.append(r)
        Record.objects.bulk_create(to_create,batch_size=600)
        ImportRow.objects.bulk_update(changed,['status','issues'],batch_size=600)
        counts=dict(Counter(batch.rows.values_list('status',flat=True)));counts['total']=sum(counts.values());counts['unknown_sheets']=batch.summary.get('unknown_sheets',[])
        if batch.summary.get('skipped_sheets'):counts['skipped_sheets']=batch.summary['skipped_sheets']
        batch.summary=counts;batch.status='partial' if any(counts.get(k) for k in ['invalid','conflict']) else 'committed';batch.committed_at=timezone.now();batch.save()
        AuditEvent.objects.create(action='import.commit',object_type='ImportBatch',object_id=str(batch.id),detail=counts)
        return batch
